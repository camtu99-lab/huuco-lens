import base64
import csv
import html
import io
import json
import math
import random
import re
import time
import unicodedata
from datetime import datetime

import streamlit as st
import streamlit.components.v1 as components
from google import genai
from google.genai import errors, types
from PIL import Image
from collections import Counter

from rdkit import Chem, RDLogger
from rdkit.Chem import Descriptors, Draw, rdMolDescriptors
from rdkit.Chem.EnumerateStereoisomers import EnumerateStereoisomers, StereoEnumerationOptions

RDLogger.DisableLog("rdApp.*")

MODEL = st.secrets.get("GEMINI_MODEL", "gemini-2.5-flash")
MAX_PER_SESSION = int(st.secrets.get("MAX_PER_SESSION", 15))
BENCH_MAX = int(st.secrets.get("BENCH_MAX", 30))

APP_NAME = st.secrets.get("APP_NAME", "MolSnap")
APP_ICON = st.secrets.get("APP_ICON", "🧪")

st.set_page_config(page_title=APP_NAME, page_icon=APP_ICON, layout="wide", initial_sidebar_state="auto")

RULES = """Bạn là giáo viên Hóa hữu cơ Việt Nam kiêm chuyên gia nhận dạng công thức. Quy ước trong SGK Việt Nam: công thức khai triển/thu gọn (CH3-CH2-OH, CH3COOH), công thức khung (mỗi đầu mút/đỉnh là C, H gắn vào C không vẽ), vòng benzen vẽ lục giác với 3 liên kết đôi xen kẽ hoặc có vòng tròn bên trong, C6H5- hoặc Ph là phenyl, nhóm chức như -OH, -CHO, -COOH, -COO-, -CO-, -NH2, -NO2, -C≡N, -X. Đếm cẩn thận số C, vị trí nhánh và nhóm thế trên vòng (o-, m-, p-). Danh pháp dùng cách đọc tiếng Việt theo SGK (metan, etanol, propan-2-ol, axit etanoic, etanal, benzen, phenol, anilin, toluen, etyl axetat...), kèm tên thường gọi nếu có."""

FMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn, các khóa:
{"smiles":"SMILES chuẩn (không có H tường minh)","do_tin_cay":số 0-1,"ten_danh_phap":"tên thay thế/IUPAC đọc kiểu SGK Việt Nam","ten_thuong_goi":"tên thông thường hoặc rỗng","loai_hop_chat":"ví dụ: ancol no đơn chức","nhom_chuc":["..."],"trang_thai_nhiet_do":"thể ở điều kiện thường, nhận xét nhiệt độ sôi/tan","tinh_tan":"tính tan trong nước và trong dung môi thường gặp (etanol, ete...): nêu rõ tan nhiều, ít tan hay không tan; chỉ ghi độ tan bằng số khi chắc chắn, không chắc thì nói rõ, không bịa số","tinh_chat_vat_ly":["..."],"tinh_chat_hoa_hoc":["phản ứng đặc trưng kèm phương trình ngắn, điều kiện"],"dieu_che_ung_dung":["..."],"dong_phan":"nhận xét ngắn nếu có ý nghĩa","luu_y":["điều chưa chắc chắn trong ảnh, nét vẽ mơ hồ, cảnh báo an toàn"]}
Nếu ảnh không phải công thức hóa học hoặc không đọc được, đặt "smiles":"" và giải thích trong luu_y."""

SUB = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")


@st.cache_resource
def get_client():
    return genai.Client(api_key=st.secrets["GEMINI_API_KEY"])


def to_jpeg(file_bytes: bytes) -> bytes:
    img = Image.open(io.BytesIO(file_bytes)).convert("RGB")
    img.thumbnail((1600, 1600))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def ask_ai(contents, temperature=None):
    """Gọi Gemini, tự thử lại khi máy chủ quá tải (503) hoặc hết lượt tạm thời (429),
    rồi chuyển sang mô hình dự phòng."""
    models = [MODEL]
    fb = st.secrets.get("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite")
    if fb and fb != MODEL:
        models.append(fb)
    last = None
    for m in models:
        for i in range(3):
            try:
                resp = get_client().models.generate_content(
                    model=m,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        **({} if temperature is None else {"temperature": temperature}),
                    ),
                )
                text = (resp.text or "").replace("```json", "").replace("```", "").strip()
                return json.loads(text[text.index("{"): text.rindex("}") + 1])
            except errors.ServerError as e:
                last = e
                time.sleep(2 * (i + 1))
            except errors.ClientError as e:
                code = getattr(e, "code", None)
                if code == 429:
                    last = e
                    time.sleep(3 * (i + 1))
                elif code == 404:
                    last = e
                    break  # mô hình này không còn dùng được, chuyển sang mô hình tiếp theo
                else:
                    raise
    raise last


def analyze_image(file_bytes: bytes):
    prompt = (
        RULES
        + "\n\nĐây là ảnh chụp công thức cấu tạo viết tay. Hãy nhận dạng đúng một chất hữu cơ trong ảnh. "
        + "Chép ĐÚNG như hình vẽ, KHÔNG tự sửa lỗi hóa trị và không thêm hoặc bớt nguyên tử. "
        + "Nếu hình vẽ sai (ví dụ nguyên tử C có 5 liên kết) vẫn ghi SMILES đúng theo hình và nêu rõ lỗi trong luu_y.\n"
        + FMT
    )
    return ask_ai([types.Part.from_bytes(data=to_jpeg(file_bytes), mime_type="image/jpeg"), prompt], temperature=0)


def analyze_smiles(smiles: str):
    prompt = (
        RULES
        + f"\n\nChất có SMILES: {smiles}\nGiữ nguyên SMILES này trong khóa \"smiles\", đặt do_tin_cay=1 và cung cấp tên cùng tính chất.\n"
        + FMT
    )
    return ask_ai([prompt])


def chem_numbers(smiles: str):
    """Tính trực tiếp bằng RDKit, không phụ thuộc AI."""
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        return None, None
    formula = rdMolDescriptors.CalcMolFormula(mol)
    pretty = "".join(c.translate(SUB) if c.isdigit() else c for c in formula)
    molh = Chem.AddHs(mol)
    cnt = {}
    for a in molh.GetAtoms():
        cnt[a.GetSymbol()] = cnt.get(a.GetSymbol(), 0) + 1
    dbe = None
    if cnt.get("C"):
        x = sum(cnt.get(k, 0) for k in ("F", "Cl", "Br", "I"))
        dbe = (2 * cnt["C"] + 2 + cnt.get("N", 0) + cnt.get("P", 0) - cnt.get("H", 0) - x) / 2
    return mol, {"formula": pretty, "mass": Descriptors.MolWt(mol), "dbe": dbe}


def bullets(title, items):
    if items:
        st.markdown(f"**{title}**")
        for t in [items] if isinstance(items, str) else items:
            st.markdown(f"- {t}")


def render(r: dict, from_image: bool):
    smiles = (r.get("smiles") or "").strip()
    if not smiles:
        st.warning("Chưa nhận dạng được công thức trong ảnh.")
    mol, num = chem_numbers(smiles) if smiles else (None, None)
    if smiles and mol is None:
        st.error("Mã SMILES không hợp lệ. Hãy sửa ở ô bên dưới rồi bấm Phân tích lại.")
    c1, c2 = st.columns(2)
    with c1:
        if mol is not None:
            st.image(Draw.MolToImage(mol, size=(420, 300)))
    with c2:
        if smiles:
            st.markdown("**SMILES**")
            st.code(smiles, language=None)
        if r.get("ten_danh_phap"):
            st.markdown(f"**Tên (danh pháp):** {r['ten_danh_phap']}")
        if r.get("ten_thuong_goi"):
            st.markdown(f"**Tên thường gọi:** {r['ten_thuong_goi']}")
        if r.get("loai_hop_chat"):
            st.markdown(f"**Loại hợp chất:** {r['loai_hop_chat']}")
        if num:
            st.markdown(f"**Công thức phân tử:** {num['formula']}")
            st.markdown(f"**Khối lượng mol:** {num['mass']:.2f} g/mol")
            if num["dbe"] is not None:
                k = num["dbe"]
                st.markdown(f"**Độ bất bão hòa (k):** {int(k) if k == int(k) else k}")
        if r.get("tinh_tan"):
            _tt = r["tinh_tan"]
            st.markdown("**Tính tan:** " + (" ".join(str(x) for x in _tt) if isinstance(_tt, list) else str(_tt)))
    if from_image and isinstance(r.get("do_tin_cay"), (int, float)):
        p = max(0.0, min(1.0, float(r["do_tin_cay"])))
        st.progress(p, text=f"Độ tin cậy nhận dạng: {round(p * 100)}%")
    if r.get("nhom_chuc"):
        st.markdown("**Nhóm chức:** " + " · ".join(r["nhom_chuc"]))
    bullets("Trạng thái, nhiệt độ", r.get("trang_thai_nhiet_do"))
    bullets("Tính chất vật lí", r.get("tinh_chat_vat_ly"))
    bullets("Tính chất hóa học", r.get("tinh_chat_hoa_hoc"))
    bullets("Điều chế và ứng dụng", r.get("dieu_che_ung_dung"))
    bullets("Đồng phân", r.get("dong_phan"))
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))



# ---------------- Phản ứng hóa học ----------------
RULES_RX = """Bạn là giáo viên Hóa học Việt Nam. Hãy đọc phương trình phản ứng (viết tay hoặc gõ) và phân loại theo cách của SGK Việt Nam (Hóa 8-12). Các loại thường gặp: phản ứng thế (thế halogen, thế nitro, thế brom vào phenol...), cộng (cộng H2, Br2, HX, H2O), tách (tách H2O, tách HX, tách H2), trùng hợp, trùng ngưng, oxi hóa hoàn toàn (đốt cháy), oxi hóa không hoàn toàn, oxi hóa - khử, hóa hợp, phân hủy, trao đổi, trung hòa, este hóa, thủy phân, xà phòng hóa, tráng bạc, lên men, cracking, đề hiđro hóa, hiđro hóa. Một phản ứng có thể thuộc nhiều loại cùng lúc (ví dụ: cộng đồng thời là oxi hóa - khử); hãy nêu loại chính trước. Đọc kỹ hệ số và điều kiện (nhiệt độ, xúc tác, áp suất). Với muối/bazơ vô cơ dùng SMILES dạng ion, ví dụ NaOH là [Na+].[OH-], HCl là Cl, H2O là O, O2 là O=O, H2 là [H][H]."""

RFMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn:
{"phuong_trinh":"phương trình đọc được, viết dạng chữ, dùng -> giữa hai vế",
"chat_tham_gia":[{"he_so":số nguyên,"ten":"tên chất","cong_thuc":"công thức phân tử","smiles":"SMILES"}],
"san_pham":[{"he_so":số nguyên,"ten":"","cong_thuc":"","smiles":""}],
"dieu_kien":"nhiệt độ, xúc tác... hoặc rỗng",
"loai_phan_ung":["loại chính","loại phụ nếu có"],
"giai_thich":"2-4 câu: dựa vào sự thay đổi nào của cấu tạo/nhóm chức/số oxi hóa mà kết luận loại phản ứng",
"phuong_trinh_can_bang":"phương trình đã cân bằng đúng nếu ảnh chưa cân bằng hoặc sai, ngược lại để rỗng",
"luu_y":["điều chưa chắc chắn khi đọc ảnh, ví dụ hệ số hoặc chất bị mờ"]}
Nếu không phải phương trình phản ứng hoặc không đọc được, trả về chat_tham_gia và san_pham rỗng và giải thích trong luu_y."""


def analyze_reaction(img_bytes, text):
    head = RULES_RX + "\n\n"
    if img_bytes:
        prompt = head + "Đây là ảnh chụp một phương trình phản ứng hóa học.\n" + RFMT
        return ask_ai([types.Part.from_bytes(data=to_jpeg(img_bytes), mime_type="image/jpeg"), prompt])
    return ask_ai([head + f"Phương trình: {text}\n" + RFMT])


def atom_balance(species):
    """Đếm nguyên tử và điện tích mỗi vế bằng RDKit. Trả về None nếu có SMILES lỗi."""
    tot, charge = {}, 0
    for coef, smi in species:
        mol = Chem.MolFromSmiles(smi or "")
        if mol is None:
            return None
        for a in Chem.AddHs(mol).GetAtoms():
            tot[a.GetSymbol()] = tot.get(a.GetSymbol(), 0) + coef
            charge += coef * a.GetFormalCharge()
    return tot, charge


def to_int(v):
    try:
        return max(1, int(v))
    except Exception:
        return 1


def render_reaction(r: dict):
    left, right = r.get("chat_tham_gia") or [], r.get("san_pham") or []
    if r.get("phuong_trinh"):
        st.markdown("**Phương trình đọc được**")
        st.code(r["phuong_trinh"], language=None)
    types_ = r.get("loai_phan_ung") or []
    if types_:
        st.success("**Loại phản ứng:** " + " · ".join(types_))
    if r.get("giai_thich"):
        st.markdown("**Vì sao:** " + r["giai_thich"])
    if r.get("dieu_kien"):
        st.markdown("**Điều kiện:** " + r["dieu_kien"])

    rows, mols, legends = [], [], []
    for role, lst in (("Tham gia", left), ("Sản phẩm", right)):
        for sp in lst:
            m = Chem.MolFromSmiles(sp.get("smiles") or "")
            mass = f"{Descriptors.MolWt(m):.2f}" if m is not None else "—"
            rows.append({"Vai trò": role, "Hệ số": to_int(sp.get("he_so")), "Tên": sp.get("ten", ""),
                         "Công thức": sp.get("cong_thuc", ""), "SMILES": sp.get("smiles", ""), "M (g/mol)": mass})
            if m is not None and m.GetNumHeavyAtoms() > 1:
                mols.append(m)
                legends.append(sp.get("cong_thuc") or "")
    if rows:
        st.dataframe(rows, hide_index=True, width="stretch")
    if mols:
        try:
            st.image(Draw.MolsToGridImage(mols, molsPerRow=3, subImgSize=(260, 190), legends=legends))
        except Exception:
            pass

    if left and right:
        bl = atom_balance([(to_int(s.get("he_so")), s.get("smiles")) for s in left])
        br = atom_balance([(to_int(s.get("he_so")), s.get("smiles")) for s in right])
        if bl is None or br is None:
            st.info("Không kiểm tra được cân bằng vì có chất có SMILES không hợp lệ.")
        elif bl == br:
            st.success("Phương trình **đã cân bằng** (RDKit đếm nguyên tử và điện tích hai vế bằng nhau).")
        else:
            diff = [f"{k}: {bl[0].get(k, 0)} ≠ {br[0].get(k, 0)}" for k in sorted(set(bl[0]) | set(br[0]))
                    if bl[0].get(k, 0) != br[0].get(k, 0)]
            if bl[1] != br[1]:
                diff.append(f"điện tích: {bl[1]} ≠ {br[1]}")
            st.warning("Phương trình **chưa cân bằng** theo hệ số đọc được. Lệch: " + "; ".join(diff))
            if r.get("phuong_trinh_can_bang"):
                st.markdown("**Gợi ý cân bằng:**")
                st.code(r["phuong_trinh_can_bang"], language=None)
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))


def reaction_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    src = st.radio("Nguồn phương trình", ["🖼️ Tải ảnh", "📷 Chụp ảnh", "⌨️ Gõ phương trình"], horizontal=True, key="rx_src")
    img, txt = None, None
    if src.startswith("🖼️"):
        f = st.file_uploader("Chọn ảnh phương trình", type=["jpg", "jpeg", "png", "webp"], key="rx_up")
        img = f.getvalue() if f else None
    elif src.startswith("📷"):
        f = st.camera_input("Chụp phương trình", key="rx_cam")
        img = f.getvalue() if f else None
    else:
        txt = st.text_input("Phương trình", placeholder="Ví dụ: CH2=CH2 + Br2 -> CH2Br-CH2Br", key="rx_txt").strip()
    if st.button("Nhận dạng và phân loại phản ứng", type="primary", disabled=not (img or txt), key="rx_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang phân tích…"):
                try:
                    st.session_state.rx = analyze_reaction(img, txt)
                    add_history(((st.session_state.rx or {}).get('phuong_trinh') or txt or 'Phương trình')[:60], 'Phản ứng')
                    st.session_state.n += 1
                except Exception as e:
                    st.session_state.rx = None
                    show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)
    if st.session_state.get("rx"):
        st.divider()
        render_reaction(st.session_state.rx)



# ---------------- Đồng phân ----------------
RULES_ISO = """Bạn là giáo viên Hóa hữu cơ Việt Nam. Xác định đồng phân theo SGK: đồng phân cấu tạo (mạch cacbon, vị trí nhóm chức hoặc liên kết bội, nhóm chức) và đồng phân lập thể (hình học cis-trans, quang học). Tên gọi dùng danh pháp tiếng Việt theo SGK (butan-1-ol, đimetyl ete, 2-metylpropan-2-ol...)."""

PRETTY = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")


def formula_of(m):
    return rdMolDescriptors.CalcMolFormula(m)


def pretty(f):
    return "".join(c.translate(PRETTY) if c.isdigit() else c for c in f)


def noiso(m):
    return Chem.MolToSmiles(m, isomericSmiles=False)


def stereo_count(m):
    try:
        opts = StereoEnumerationOptions(onlyUnassigned=True, unique=True, tryEmbedding=False, maxIsomers=32)
        return max(1, len(list(EnumerateStereoisomers(m, options=opts))))
    except Exception:
        return 1


def mirror(smi):
    return smi.replace("@@", "\x00").replace("@", "@@").replace("\x00", "@")


def relation(m1, m2):
    f1, f2 = formula_of(m1), formula_of(m2)
    if f1 != f2:
        return "khac_cong_thuc"
    if noiso(m1) != noiso(m2):
        return "cau_tao"
    i1, i2 = Chem.MolToSmiles(m1), Chem.MolToSmiles(m2)
    if i1 == i2:
        return "cung_chat"
    mir = Chem.MolFromSmiles(mirror(i1))
    if mir is not None and Chem.MolToSmiles(mir) == i2:
        return "doi_quang"
    if not Chem.FindMolChiralCenters(m1) and not Chem.FindMolChiralCenters(m2):
        return "hinh_hoc"
    return "lap_the"


# ---- Sinh đồng phân cấu tạo bằng thuật toán (không dùng AI) ----
VALENCE = {"C": 4, "N": 3, "O": 2, "S": 2, "F": 1, "Cl": 1, "Br": 1, "I": 1}
_BT = {1: Chem.BondType.SINGLE, 2: Chem.BondType.DOUBLE, 3: Chem.BondType.TRIPLE}
ISO_TIME = float(st.secrets.get("ISO_TIME", 15))


class TooBig(Exception):
    pass


def _canon(atoms, edges, strict=False):
    rw = Chem.RWMol()
    for e in atoms:
        rw.AddAtom(Chem.Atom(e))
    for i, j, o in edges:
        rw.AddBond(i, j, _BT[o])
    m = rw.GetMol()
    Chem.SanitizeMol(m)
    if strict:
        for b in m.GetBonds():
            if b.GetBondType() == Chem.BondType.TRIPLE and b.IsInRing():
                raise ValueError("liên kết ba trong vòng")
    return Chem.MolToSmiles(m, isomericSmiles=False)


def enumerate_isomers(counts, nH, time_limit=15, max_graphs=200000, strict=False):
    """Liệt kê TẤT CẢ đồng phân cấu tạo của công thức (mọi đồ thị liên thông đúng hóa trị, không trùng nhau).
    strict=True: chỉ mạch hở hoặc 1 vòng, không có O–O, không liên kết ba trong vòng."""
    els = sorted(counts)
    N = sum(counts.values())
    sigma = sum(VALENCE[e] * counts[e] for e in els)
    if sigma - nH < 0 or (sigma - nH) % 2:
        return []
    B = (sigma - nH) // 2
    if B < N - 1:
        return []
    t0, memo, total = time.time(), {}, [0]

    def build(sub):
        if sub in memo:
            return memo[sub]
        n = sum(sub)
        res = {}
        if n == 1:
            e = els[[i for i, c in enumerate(sub) if c][0]]
            res[_canon([e], [])] = ([e], [])
            memo[sub] = res
            return res
        upper = B - (N - n)
        for i, c in enumerate(sub):
            if not c:
                continue
            prev = list(sub)
            prev[i] -= 1
            e, vn = els[i], VALENCE[els[i]]
            for atoms, edges in list(build(tuple(prev)).values()):
                T = sum(o for _, _, o in edges)
                if T + 1 > upper:
                    continue
                used = [0] * len(atoms)
                for a, b, o in edges:
                    used[a] += o
                    used[b] += o
                free = [VALENCE[atoms[k]] - used[k] for k in range(len(atoms))]
                k = len(atoms)

                def rec(j, rem_new, added, att):
                    if time.time() - t0 > time_limit:
                        raise TooBig("time")
                    if j == k:
                        if att and T + added <= upper and (n < N or T + added == B):
                            na, ne = atoms + [e], edges + [(a, k, o) for a, o in att]
                            if strict:
                                if len(ne) - len(na) + 1 > 1:
                                    return
                                if e == "O" and any(atoms[a] == "O" for a, _ in att):
                                    return
                            try:
                                key = _canon(na, ne, strict)
                            except Exception:
                                return
                            if key not in res:
                                res[key] = (na, ne)
                                total[0] += 1
                                if total[0] > max_graphs:
                                    raise TooBig("size")
                        return
                    rec(j + 1, rem_new, added, att)
                    for o in (1, 2, 3):
                        if o <= rem_new and o <= free[j] and T + added + o <= upper:
                            rec(j + 1, rem_new - o, added + o, att + [(j, o)])

                rec(0, vn, 0, [])
        memo[sub] = res
        return res

    return sorted(build(tuple(counts[e] for e in els)).keys())


def parse_formula(s):
    s = (s or "").translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")).replace(" ", "")
    toks = re.findall(r"([A-Z][a-z]?)(\d*)", s)
    if not toks or "".join(a + b for a, b in toks) != s:
        raise ValueError("không phải công thức phân tử")
    c = {}
    for a, b in toks:
        c[a] = c.get(a, 0) + (int(b) if b else 1)
    nH = c.pop("H", 0)
    bad = [e for e in c if e not in VALENCE]
    if bad or not c:
        raise ValueError("nguyên tố chưa hỗ trợ: " + ", ".join(bad or ["?"]))
    return c, nH


def is_simple_formula(s):
    s = (s or "").translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")).replace(" ", "")
    toks = re.findall(r"([A-Z][a-z]?)(\d*)", s)
    if not toks or "".join(a + b for a, b in toks) != s:
        return False
    syms = [a for a, _ in toks]
    return len(syms) == len(set(syms)) and (any(b for _, b in toks) or "H" in syms)


def hill_key(counts, nH):
    parts = []
    if "C" in counts:
        parts.append(("C", counts["C"]))
    if nH:
        parts.append(("H", nH))
    parts += [(e, counts[e]) for e in sorted(counts) if e != "C"]
    return "".join(f"{e}{n if n > 1 else ''}" for e, n in parts)


def mol_counts(m):
    c = atom_counts(m)
    nH = c.pop("H", 0)
    return c, nH


# Phân loại theo chương trình phổ thông (dùng riêng cho tab đồng phân)
ISO_GROUPS = [
    ("ancol bậc I", "[CX4;H2,H3][OX2H]"), ("ancol bậc II", "[CX4;H1]([#6])([#6])[OX2H]"),
    ("ancol bậc III", "[CX4;H0]([#6])([#6])([#6])[OX2H]"), ("phenol", "c[OX2H]"),
    ("ete", "[#6;!$([#6]=O)][OX2][#6;!$([#6]=O)]"), ("anđehit", "[CX3;H1,H2;!$(C[OX2,NX3])]=O"),
    ("xeton", "[#6][CX3](=O)[#6]"), ("axit cacboxylic", "[CX3](=O)[OX2H1]"), ("este", "[CX3](=O)[OX2][#6]"),
    ("amin bậc I", "[NX3;H2;!$(NC=[O,N,S])][#6]"), ("amin bậc II", "[NX3;H1;!$(NC=[O,N,S])]([#6])[#6]"),
    ("amin bậc III", "[NX3;H0;!$(NC=[O,N,S])]([#6])([#6])[#6]"), ("amit", "[NX3][CX3]=O"),
    ("nitro", "[N+](=O)[O-]"), ("nitrin", "[CX2]#[NX1]"), ("anken (C=C)", "[CX3]=[CX3]"), ("ankin (C≡C)", "[CX2]#[CX2]"),
    ("hợp chất thơm", "a"), ("dẫn xuất halogen", "[#6][F,Cl,Br,I]"),
]
_ISO_GP = [(n, Chem.MolFromSmarts(p)) for n, p in ISO_GROUPS]
UNSTABLE = [
    ("enol (OH trên C=C)", "[OX2H][CX3]=[CX3]"), ("ynol (OH trên C≡C)", "[OX2H][CX2]#[CX2]"),
    ("hai nhóm OH trên cùng một C", "[CX4]([OX2H])[OX2H]"), ("hemiaxetal", "[CX4]([OX2H])[OX2][#6]"),
    ("OH và halogen trên cùng một C", "[CX4]([OX2H])[F,Cl,Br,I]"), ("OH và N trên cùng một C", "[CX4]([OX2H])[NX3]"),
]
_UNST = [(n, Chem.MolFromSmarts(p)) for n, p in UNSTABLE]


_PAIR = Chem.MolFromSmarts("[CX3](=O)[OX2,NX3]")
_NITRO = Chem.MolFromSmarts("[N+](=O)[O-]")


def func_units(m):
    """Số nhóm chức chứa dị tố: axit/este/amit và nitro mỗi nhóm chỉ tính 1 dù có nhiều nguyên tử."""
    n_het = sum(1 for a in m.GetAtoms() if a.GetSymbol() in ("O", "N", "S", "F", "Cl", "Br", "I"))
    return n_het - len(m.GetSubstructMatches(_PAIR)) - 2 * len(m.GetSubstructMatches(_NITRO))


def is_mono_open(m):
    """Đơn chức (tối đa 1 nhóm chức chứa dị tố) và không có vòng no/không thơm (vòng benzen vẫn được giữ)."""
    return func_units(m) <= 1 and not any(a.IsInRing() and not a.GetIsAromatic() for a in m.GetAtoms())


def iso_labels(m):
    lab = [n for n, p in _ISO_GP if m.HasSubstructMatch(p)]
    has_ring = any(not a.GetIsAromatic() and a.IsInRing() for a in m.GetAtoms())
    if not lab:
        if all(a.GetSymbol() == "C" for a in m.GetAtoms()):
            return ["xicloankan"] if has_ring else ["ankan"]
        lab = ["khác"]
    if has_ring:
        lab.append("có vòng không thơm")
    return lab


@st.cache_data(show_spinner=False, max_entries=64)
def iso_enumerate(fkey, strict):
    counts, nH = parse_formula(fkey)
    try:
        return {"ok": True, "smiles": enumerate_isomers(counts, nH, time_limit=ISO_TIME, strict=strict)}
    except TooBig as e:
        return {"ok": False, "why": str(e)}


def stereo_split(m):
    """(tổng, hình học, quang học): số dạng lập thể; hình học = cis/trans bỏ qua quang học; quang học = bỏ qua cis/trans."""
    try:
        opts = StereoEnumerationOptions(onlyUnassigned=True, unique=True, tryEmbedding=False, maxIsomers=64)
        isos = list(EnumerateStereoisomers(m, options=opts))
    except Exception:
        return 1, 1, 1
    geo, opt = set(), set()
    for x in isos:
        a = Chem.Mol(x)  # hình học: giữ cis/trans của nối đôi và cấu hình tương đối trên vòng; bỏ tâm bất đối ngoài vòng và gộp cặp đối quang
        for at in a.GetAtoms():
            if not at.IsInRing():
                at.SetChiralTag(Chem.ChiralType.CHI_UNSPECIFIED)
        s1 = Chem.MolToSmiles(a)
        mir = Chem.MolFromSmiles(mirror(s1))
        geo.add(min(s1, Chem.MolToSmiles(mir)) if mir is not None else s1)
        b = Chem.Mol(x)  # quang học: giữ mọi tâm bất đối, bỏ cis/trans của nối đôi
        for bd in b.GetBonds():
            if bd.GetBondType() == Chem.BondType.DOUBLE:
                bd.SetStereo(Chem.BondStereo.STEREONONE)
                bd.SetBondDir(Chem.BondDir.NONE)
        opt.add(Chem.MolToSmiles(b))
    return max(1, len(isos)), max(1, len(geo)), max(1, len(opt))


# Từ vựng nhóm mà AI được phép dùng; mỗi từ mở rộng thành các nhãn do bộ máy phân loại
GROUP_VOCAB = {
    "ankan": ["ankan"], "xicloankan": ["xicloankan"], "anken": ["anken (C=C)"], "ankin": ["ankin (C≡C)"],
    "thơm": ["hợp chất thơm"], "ancol": ["ancol bậc I", "ancol bậc II", "ancol bậc III"],
    "ancol bậc I": ["ancol bậc I"], "ancol bậc II": ["ancol bậc II"], "ancol bậc III": ["ancol bậc III"],
    "phenol": ["phenol"], "ete": ["ete"], "anđehit": ["anđehit"], "xeton": ["xeton"],
    "axit cacboxylic": ["axit cacboxylic"], "este": ["este"],
    "amin": ["amin bậc I", "amin bậc II", "amin bậc III"], "amin bậc I": ["amin bậc I"],
    "amin bậc II": ["amin bậc II"], "amin bậc III": ["amin bậc III"], "amit": ["amit"], "nitro": ["nitro"],
    "nitrin": ["nitrin"], "dẫn xuất halogen": ["dẫn xuất halogen"],
}
PRESETS = {
    "mono": {"ben": True, "chuc": "don", "mach": "ho_hoac_thom", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
    "pho": {"ben": True, "chuc": "bat_ky", "mach": "bat_ky", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
    "math": {"ben": False, "chuc": "bat_ky", "mach": "bat_ky", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
}


def norm_params(raw):
    """Kiểm tra tham số AI trích ra: chỉ nhận giá trị trong danh mục. Giá trị AI để trống thì dùng mặc định và ghi rõ."""
    raw = raw or {}
    warn, defaults = [], []
    names = []
    for g in raw.get("nhom") or []:
        key = str(g).strip().lower().replace("đ", "đ")
        hit = next((k for k in GROUP_VOCAB if k.lower() == key), None)
        if hit:
            names.append(hit)
        else:
            warn.append(f"Bỏ qua nhóm không có trong danh mục: {g}")
    p = {"ben": True, "nhom_ten": names, "nhom": [GROUP_VOCAB[n] for n in names]}
    kieu = raw.get("kieu_nhom")
    p["kieu"] = kieu if kieu in ("bat_ky", "tat_ca") else "bat_ky"
    for k, ok, dflt in (("chuc", ("don", "da", "bat_ky"), "don"), ("mach", ("ho", "vong", "bat_ky"), "ho_hoac_thom"),
                        ("lap_the", ("khong", "hinh_hoc", "quang_hoc", "tat_ca"), "khong")):
        v = raw.get(k)
        if v in ok:
            p[k] = v
        else:
            p[k] = dflt
            defaults.append(k)
    if "xicloankan" in names and "mach" in defaults:
        p["mach"] = "bat_ky"
    p["defaults"] = defaults
    return p, warn


def pass_filters(m, labels, p):
    u = func_units(m)
    if (p["chuc"] == "don" and u > 1) or (p["chuc"] == "da" and u < 2):
        return False
    rings = m.GetRingInfo().NumRings()
    nonarom = any(a.IsInRing() and not a.GetIsAromatic() for a in m.GetAtoms())
    if (p["mach"] == "ho" and rings) or (p["mach"] == "vong" and not rings) or (p["mach"] == "ho_hoac_thom" and nonarom):
        return False
    if p["nhom"]:
        hit = [bool(set(labels) & set(g)) for g in p["nhom"]]
        return all(hit) if p["kieu"] == "tat_ca" else any(hit)
    return True


PARSE_PROMPT = """Bạn là bộ PHÂN TÍCH CÂU HỎI Hóa học. Nhiệm vụ DUY NHẤT: trích tham số từ câu hỏi về đồng phân. TUYỆT ĐỐI không đếm, không liệt kê chất, không trả lời câu hỏi.
Từ vựng nhóm (chỉ dùng đúng các chuỗi này): ankan, xicloankan, anken, ankin, thơm, ancol, ancol bậc I, ancol bậc II, ancol bậc III, phenol, ete, anđehit, xeton, axit cacboxylic, este, amin, amin bậc I, amin bậc II, amin bậc III, amit, nitro, nitrin, dẫn xuất halogen.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"cong_thuc":"công thức phân tử viết liền, ví dụ C4H8O2, hoặc null",
"chat_goc":"tên hoặc SMILES của chất nếu câu hỏi nhắc một chất cụ thể thay vì công thức, hoặc null",
"nhom":["nhóm mà đồng phân cần tìm phải thuộc, lấy từ từ vựng; để [] nếu câu hỏi không nêu nhóm"],
"kieu_nhom":"bat_ky nếu chỉ cần thuộc ít nhất một nhóm, tat_ca nếu phải có đồng thời mọi nhóm, null nếu không nêu",
"chuc":"don nếu câu hỏi nói đơn chức, da nếu nói đa chức hoặc tạp chức, bat_ky nếu nói rõ tính mọi loại, null nếu không nhắc",
"mach":"ho nếu mạch hở, vong nếu mạch vòng, bat_ky nếu nói rõ tính cả hai, null nếu không nhắc",
"lap_the":"khong nếu nói không tính đồng phân lập thể, hinh_hoc nếu tính thêm đồng phân hình học (cis-trans), quang_hoc nếu tính thêm đồng phân quang học, tat_ca nếu tính mọi đồng phân lập thể, null nếu không nhắc",
"can_hoi_lai":"một câu hỏi làm rõ nếu thiếu công thức hoặc chất, hoặc câu hỏi mơ hồ; nếu không thì để rỗng",
"tom_tat":"câu hỏi được hiểu lại ngắn gọn bằng tiếng Việt"}
Không đoán những gì câu hỏi không nói: để null hoặc []. Ví dụ: 'Có bao nhiêu đồng phân este của C4H8O2?' thì nhom=["este"], cong_thuc="C4H8O2", các mục còn lại null."""


def parse_question(q):
    raw = ask_ai([PARSE_PROMPT + "\n\nCâu hỏi: " + q], temperature=0)
    p, warn = norm_params(raw)
    return p, warn, raw


def resolve_formula(img, txt):
    """Trả về (counts, nH, SMILES chất gốc hoặc None, lỗi hoặc None)."""
    if img:
        smi, _ = transcribe_student(img)
        m, err = parse_student(smi)
        if m is None:
            return None, None, None, "Hình vẽ sai nên không xác định được công thức: " + (err or "không đọc được hình")
    elif is_simple_formula(txt):
        try:
            c, h = parse_formula(txt)
            return c, h, None, None
        except ValueError as e:
            return None, None, None, f"Công thức không hợp lệ ({e})."
    else:
        m, err = parse_student(txt) if smiles_like(txt) else (None, None)
        if m is None:
            smi = to_smiles(None, txt)  # tên chất hoặc công thức thu gọn: nhờ AI đổi sang SMILES
            m, err = parse_student(smi)
            if m is None:
                return None, None, None, "Không hiểu chất đã nhập. Hãy gõ công thức phân tử (ví dụ C4H10O), SMILES hoặc tên chất."
    c, h = mol_counts(m)
    if any(e not in VALENCE for e in c):
        return None, None, None, "Chất có nguyên tố chưa được hỗ trợ."
    return c, h, Chem.MolToSmiles(m, isomericSmiles=False), None


def ai_names(smiles_list):
    """Nhờ AI đặt tên. Cấu trúc do thuật toán sinh nên chỉ phần tên có thể sai."""
    lines = "\n".join(f"{i + 1}. {s_}" for i, s_ in enumerate(smiles_list))
    r = ask_ai([RULES_ISO + "\n\nĐặt tên theo danh pháp SGK cho từng chất sau (theo đúng thứ tự). "
                'Chỉ trả JSON {"ten":["tên 1","tên 2",...]} có đúng ' + str(len(smiles_list)) + " phần tử.\n" + lines], temperature=0)
    names = r.get("ten") or []
    return {s_: (names[i] if i < len(names) else "") for i, s_ in enumerate(smiles_list)}


_SMI_RE = re.compile(r"[A-Za-z0-9@+\-\[\]()=#$/\\%.:]+")


def smiles_like(t):
    """Chuỗi giống SMILES thật: không có khoảng trắng hoặc chữ có dấu. Tránh câu tiếng Việt như 'Có bao nhiêu...' bị RDKit đọc nhầm."""
    t = (t or "").strip()
    return bool(t) and _SMI_RE.fullmatch(t) is not None and Chem.MolFromSmiles(t) is not None


def to_smiles(img, text):
    if smiles_like(text):
        return text
    prompt = RULES + '\n\nChuyển chất sau thành SMILES chuẩn. Chỉ trả JSON {"smiles":"","ten":""}, để rỗng nếu không đọc được.'
    if img:
        r = ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"), prompt])
    elif text:
        r = ask_ai([prompt + "\nChất: " + text])
    else:
        return ""
    return (r.get("smiles") or "").strip()


REL = {
    "cung_chat": ("warning", "Hai công thức biểu diễn **cùng một chất**, không phải đồng phân."),
    "khac_cong_thuc": ("error", "**Không phải đồng phân** vì khác công thức phân tử."),
    "cau_tao": ("success", "Hai chất là **đồng phân cấu tạo** của nhau."),
    "hinh_hoc": ("success", "Hai chất là **đồng phân hình học (cis–trans)** của nhau."),
    "doi_quang": ("success", "Hai chất là **đồng phân quang học (đối quang)** của nhau."),
    "lap_the": ("success", "Hai chất là **đồng phân lập thể** (không phải đối quang) của nhau."),
}


def compare_ui_result(s1, s2):
    m1, m2 = Chem.MolFromSmiles(s1), Chem.MolFromSmiles(s2)
    if m1 is None or m2 is None:
        st.error("Không đọc được một trong hai chất. Hãy kiểm tra lại ảnh hoặc nhập SMILES.")
        return
    c1, c2 = st.columns(2)
    for col, m, s_, tag in ((c1, m1, s1, "A"), (c2, m2, s2, "B")):
        with col:
            st.markdown(f"**Chất {tag}**")
            st.image(Draw.MolToImage(m, size=(320, 240)))
            st.markdown(f"{pretty(formula_of(m))} · M = {Descriptors.MolWt(m):.2f} g/mol")
            st.code(s_, language=None)
    rel = relation(m1, m2)
    kind, msg = REL[rel]
    getattr(st, kind)(msg)
    if rel == "khac_cong_thuc":
        st.markdown(f"{pretty(formula_of(m1))} ≠ {pretty(formula_of(m2))}")
        return
    if rel == "cung_chat":
        return
    try:
        r = ask_ai([RULES_ISO + f'\n\nChất A: {s1}\nChất B: {s2}\nRDKit đã xác nhận quan hệ: {msg}\n'
                    'Hãy đặt tên hai chất và giải thích ngắn gọn theo SGK. Chỉ trả JSON '
                    '{"ten_A":"","ten_B":"","loai_dong_phan":"ví dụ: đồng phân nhóm chức","giai_thich":"2-3 câu: điểm khác nhau về cấu tạo hoặc không gian"}'])
        st.markdown(f"**Tên:** A = {r.get('ten_A', '')}; B = {r.get('ten_B', '')}")
        if r.get("loai_dong_phan"):
            st.markdown("**Loại đồng phân:** " + r["loai_dong_phan"])
        if r.get("giai_thich"):
            st.markdown(r["giai_thich"])
    except Exception as e:
        st.caption(f"Không lấy được phần giải thích: {type(e).__name__}")


def build_iso_result(fkey, orig, p):
    """Bộ máy hóa học: sinh toàn bộ đồng phân, kiểm tra và đếm theo tham số p. AI không tham gia."""
    out = iso_enumerate(fkey, p["ben"])
    if not out["ok"]:
        return {"ok": False, "why": out["why"], "fkey": fkey}
    kept, excluded, hidden = [], {}, 0
    for smi in out["smiles"]:
        m = Chem.MolFromSmiles(smi)
        if p["ben"]:
            bad = next((n for n, pt in _UNST if m.HasSubstructMatch(pt)), None)
            if bad:
                excluded[bad] = excluded.get(bad, 0) + 1
                continue
        labels = iso_labels(m)
        if not pass_filters(m, labels, p):
            hidden += 1
            continue
        tot, geo, opt = stereo_split(m)
        kept.append({"smiles": smi, "labels": labels, "stereo": tot, "geo": geo, "opt": opt})
    return {"ok": True, "fkey": fkey, "params": p, "kept": kept, "excluded": excluded, "hidden": hidden,
            "total_math": len(out["smiles"]), "orig": orig}


LB_CHUC = {"don": "đơn chức", "da": "đa chức hoặc tạp chức", "bat_ky": "mọi loại chức"}
LB_MACH = {"ho": "mạch hở", "vong": "mạch vòng", "ho_hoac_thom": "mạch hở hoặc vòng benzen", "bat_ky": "mọi loại mạch"}
LB_LT = {"khong": "không tính đồng phân lập thể", "hinh_hoc": "kể cả đồng phân hình học (cis–trans)",
         "quang_hoc": "kể cả đồng phân quang học", "tat_ca": "kể cả mọi đồng phân lập thể"}


def answer_of(r):
    k, lt = r["kept"], r["params"]["lap_the"]
    return {"hinh_hoc": sum(x["geo"] for x in k), "quang_hoc": sum(x["opt"] for x in k),
            "tat_ca": sum(x["stereo"] for x in k)}.get(lt, len(k))


def render_iso(r, names):
    kept, p = r["kept"], r["params"]
    nh = " hoặc ".join(p["nhom_ten"])
    st.success(f"**{answer_of(r)} đồng phân** của {pretty(r['fkey'])}" + (f" ({nh})" if nh else ""))
    hid = r["hidden"] + (sum(r["excluded"].values()) if p["ben"] else 0)
    st.caption(f"Cấu tạo {len(kept)} · kể cả hình học {sum(k['geo'] for k in kept)} · quang học {sum(k['opt'] for k in kept)} · "
               f"lập thể {sum(k['stereo'] for k in kept)} · đã ẩn hoặc loại {hid}")
    comp = {}
    for k in kept:
        for lb in k["labels"]:
            d = comp.setdefault(lb, [0, 0])
            d[0] += 1
            d[1] += k["stereo"]
    with st.expander("Thống kê theo nhóm"):
        st.dataframe([{"Nhóm": n, "Cấu tạo": v[0], "Kể cả lập thể": v[1]} for n, v in sorted(comp.items())], hide_index=True, width="stretch")
        if r["params"]["ben"] and r["excluded"]:
            st.caption("Đã loại cấu trúc kém bền: " + ", ".join(f"{n} ({v})" for n, v in r["excluded"].items()))
    pick = st.multiselect("Lọc theo nhóm", sorted(comp), key="iso_filter")
    shown = [k for k in kept if not pick or any(lb in pick for lb in k["labels"])]
    if pick:
        st.caption(f"{len(shown)} đồng phân cấu tạo thuộc nhóm đã chọn.")
    orig = r.get("orig")
    hit = next((i for i, k in enumerate(kept, 1) if k["smiles"] == orig), None) if orig else None
    if hit:
        st.info(f"Chất bạn nhập là đồng phân số {hit}.")
    rows = [{"STT": i, "Nhóm": " + ".join(k["labels"]),
             "Lập thể": "không" if k["stereo"] == 1 else f"{k['stereo']} dạng",
             **({"Tên (AI)": names.get(k["smiles"], "")} if names else {}), "SMILES": k["smiles"]} for i, k in enumerate(shown, 1)]
    st.dataframe(rows[:200], hide_index=True, width="stretch")
    if len(rows) > 200:
        st.caption(f"Hiện 200 dòng đầu trong {len(rows)}. Hãy lọc theo nhóm.")
    if rows:
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
        st.download_button("⬇️ Tải CSV", buf.getvalue().encode("utf-8-sig"), file_name=f"dong_phan_{r['fkey']}.csv", mime="text/csv", key="iso_dl")
    ms = [Chem.MolFromSmiles(k["smiles"]) for k in shown[:48]]
    if ms:
        try:
            st.image(Draw.MolsToGridImage(ms, molsPerRow=4, subImgSize=(230, 180), legends=[str(i) for i in range(1, len(ms) + 1)]))
        except Exception:
            pass


def render_understood(u, r):
    """Một dòng cho biết máy hiểu câu hỏi thế nào (chỉ khi AI đọc câu hỏi)."""
    if not u:
        return
    p = r["params"]
    dflt = set(p.get("defaults", []))
    nhom = (" và " if p["kieu"] == "tat_ca" else " hoặc ").join(p["nhom_ten"]) or "mọi nhóm"
    bits = [f"nhóm: {nhom}", LB_CHUC[p["chuc"]] + ("*" if "chuc" in dflt else ""), LB_MACH[p["mach"]] + ("*" if "mach" in dflt else ""),
            LB_LT[p["lap_the"]] + ("*" if "lap_the" in dflt else "")]
    st.caption("Máy hiểu: " + " · ".join(bits) + ("  (*: mặc định vì câu hỏi không nêu)" if dflt else ""))
    if u.get("chat_goc") and r.get("orig"):
        st.caption(f"Chất được nhắc tới: {u['chat_goc']}")
        try:
            st.image(Draw.MolToImage(Chem.MolFromSmiles(r["orig"]), size=(220, 150)))
        except Exception:
            pass
    for w in u.get("warn") or []:
        st.caption(w)


def isomer_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    t1, t2 = st.tabs(["Liệt kê đồng phân", "So sánh hai chất"])
    with t1:
        q = st.text_input("Câu hỏi hoặc công thức", key="iso_q",
                          placeholder="Ví dụ: số đồng phân este của C4H8O2   ·   C4H10O").strip()
        with st.expander("Tùy chọn"):
            scope = st.radio("Phạm vi khi chỉ nhập công thức", ["Đơn chức, mạch hở (đề SGK)", "Phổ thông (cả đa chức và 1 vòng)",
                                                               "Toán học (mọi cấu trúc)"], key="iso_scope")
            want_names = st.checkbox("Gọi tên bằng AI (có thể sai)", key="iso_names")
            f = st.file_uploader("Hoặc dùng ảnh công thức", type=["jpg", "jpeg", "png", "webp"], key="iso_up")
        img = f.getvalue() if f else None
        if st.button("Đếm đồng phân", type="primary", disabled=not (q or img), key="iso_go"):
            direct = bool(img) or is_simple_formula(q) or smiles_like(q)
            ai_used = bool(img or want_names or not direct)
            if ai_used and st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
            else:
                with st.spinner("Đang tính…"):
                    try:
                        st.session_state.iso = None
                        counts = nH = orig = err = und = None
                        if direct:
                            counts, nH, orig, err = resolve_formula(img, q)
                            p = dict(PRESETS["mono" if scope.startswith("Đơn") else ("pho" if scope.startswith("Phổ") else "math")])
                        else:
                            p, warn, raw = parse_question(q)
                            ftxt, gc = raw.get("cong_thuc"), raw.get("chat_goc")
                            und = {"warn": warn, "chat_goc": gc}
                            if ftxt and is_simple_formula(str(ftxt)):
                                counts, nH = parse_formula(str(ftxt))
                            elif gc:
                                counts, nH, orig, err = resolve_formula(None, str(gc))
                            else:
                                err = raw.get("can_hoi_lai") or "Câu hỏi chưa có công thức hoặc tên chất."
                        if err:
                            show_top_error(top, err)
                        else:
                            n_heavy = sum(counts.values())
                            fkey = hill_key(counts, nH)
                            if n_heavy > 12:
                                show_top_error(top, f"{pretty(fkey)} có {n_heavy} nguyên tử khác H, quá lớn để liệt kê chính xác (tối đa 12).")
                            else:
                                res = build_iso_result(fkey, orig, p)
                                if not res["ok"]:
                                    show_top_error(top, f"{pretty(fkey)} có quá nhiều đồng phân để sinh trong {ISO_TIME:.0f} giây. Hãy nêu rõ nhóm chức.")
                                else:
                                    nm = {}
                                    if want_names and 0 < len(res["kept"]) <= 60:
                                        nm = ai_names([k["smiles"] for k in res["kept"]])
                                    st.session_state.iso, st.session_state.iso_nm, st.session_state.iso_und = res, nm, und
                                    add_history(pretty(res['fkey']), 'Đồng phân')
                                    if ai_used:
                                        st.session_state.n += 1
                    except Exception as e:
                        st.session_state.iso = None
                        show_top_error(top, "Có lỗi, hãy thử lại.", e)
        res = st.session_state.get("iso")
        if res:
            st.divider()
            render_understood(st.session_state.get("iso_und"), res)
            if res["kept"]:
                render_iso(res, st.session_state.get("iso_nm") or {})
            else:
                st.warning(f"Không có đồng phân nào của {pretty(res['fkey'])} thỏa điều kiện (đã ẩn {res['hidden']}).")
    with t2:
        st.caption("Mỗi chất nhập bằng ảnh, SMILES hoặc tên.")
        ca, cb = st.columns(2)
        inp = []
        for col, tag in ((ca, "A"), (cb, "B")):
            with col:
                st.markdown(f"**Chất {tag}**")
                f = st.file_uploader(f"Ảnh chất {tag}", type=["jpg", "jpeg", "png", "webp"], key=f"cmp_up{tag}")
                t = st.text_input(f"Hoặc SMILES/tên chất {tag}", key=f"cmp_t{tag}").strip()
                inp.append((f.getvalue() if f else None, t))
        ok = all(i or t for i, t in inp)
        if st.button("So sánh", type="primary", disabled=not ok, key="cmp_go"):
            if st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
            else:
                with st.spinner("Đang phân tích…"):
                    try:
                        smi = [to_smiles(i, t) for i, t in inp]
                        st.session_state.n += 1
                        st.session_state.cmp = smi
                    except Exception as e:
                        st.session_state.cmp = None
                        show_top_error(top, "Có lỗi khi đọc chất, hãy thử lại.", e)
        if st.session_state.get("cmp"):
            st.divider()
            compare_ui_result(*st.session_state.cmp)



# ---------------- Dụng cụ thí nghiệm ----------------
RULES_EQ = """Bạn là giáo viên Hóa học Việt Nam hướng dẫn thực hành thí nghiệm ở trường phổ thông. Hãy nhận biết dụng cụ thí nghiệm theo tên gọi trong SGK Việt Nam, ví dụ: ống nghiệm, cốc thủy tinh, bình tam giác, bình cầu (đáy tròn, đáy bằng, có nhánh), phễu lọc, phễu chiết, đèn cồn, kẹp ống nghiệm, giá sắt, kiềng đun, lưới amiăng, ống đong, pipet, buret, bình định mức, cối và chày sứ, chén sứ, đũa thủy tinh, thìa lấy hóa chất, nhiệt kế, ống sinh hàn, bình Kíp, ống dẫn khí, nút cao su, giấy lọc, cân, bình tia, giá để ống nghiệm. Chỉ nêu những gì thấy rõ trong ảnh; nếu không chắc thì giảm độ tin cậy và nói rõ."""

EFMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn:
{"loai_anh":"dụng cụ riêng lẻ | bộ lắp thí nghiệm | không phải dụng cụ thí nghiệm",
"dung_cu":[{"ten":"tên theo SGK","ten_khac":"tên gọi khác hoặc tên tiếng Anh, có thể rỗng","vi_tri":"vị trí trong ảnh, ví dụ: bên trái, ở giữa","chat_lieu":"thủy tinh | sứ | kim loại | nhựa | cao su | ...","cong_dung":"1-2 câu","cach_dung":["bước hoặc lưu ý sử dụng"],"luu_y_an_toan":["điều cần chú ý để an toàn"],"do_tin_cay":số 0-1}],
"thi_nghiem":null hoặc {"ten":"thí nghiệm có thể đang lắp, ví dụ: điều chế khí oxi từ KMnO4","mo_ta":"1-2 câu","nhan_xet_lap_rap":["điểm đúng hoặc điểm có thể sai/nguy hiểm trong cách lắp, chỉ nêu khi nhìn thấy rõ"]},
"luu_y":["điều chưa chắc chắn, ví dụ ảnh mờ, bị che khuất"]}
Nếu ảnh có nhiều dụng cụ thì liệt kê từng dụng cụ nhận ra được. Chỉ đặt thi_nghiem khi ảnh là một bộ lắp ráp rõ ràng. Nếu không phải dụng cụ thí nghiệm, để dung_cu rỗng và giải thích trong luu_y."""


def analyze_equipment(img, text):
    head = RULES_EQ + "\n\n"
    if img:
        return ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"),
                       head + "Đây là ảnh dụng cụ thí nghiệm hóa học. Hãy nhận biết.\n" + EFMT])
    return ask_ai([head + "Hãy giải thích dụng cụ có tên: " + text + ".\n" + EFMT])


def render_equipment(r):
    kind = r.get("loai_anh")
    tools = r.get("dung_cu") or []
    if kind:
        st.markdown(f"**Loại ảnh:** {kind}")
    if not tools:
        st.warning("Chưa nhận ra dụng cụ thí nghiệm nào trong ảnh.")
    for t in tools:
        with st.container(border=True):
            st.markdown(f"### {t.get('ten', 'Dụng cụ')}")
            if t.get("ten_khac"):
                st.caption(f"Tên khác: {t['ten_khac']}")
            meta = [x for x in (t.get("vi_tri"), t.get("chat_lieu")) if x]
            if meta:
                st.markdown("**Vị trí / chất liệu:** " + " · ".join(meta))
            conf = t.get("do_tin_cay")
            if isinstance(conf, (int, float)):
                p = max(0.0, min(1.0, float(conf)))
                st.progress(p, text=f"Độ tin cậy: {round(p * 100)}%")
                if p < 0.6:
                    st.warning("Độ tin cậy thấp, hãy hỏi giáo viên hoặc chụp lại rõ hơn.")
            if t.get("cong_dung"):
                st.markdown("**Công dụng:** " + t["cong_dung"])
            bullets("Cách dùng", t.get("cach_dung"))
            if t.get("luu_y_an_toan"):
                st.warning("**An toàn:** " + " • ".join([t["luu_y_an_toan"]] if isinstance(t["luu_y_an_toan"], str) else t["luu_y_an_toan"]))
    tn = r.get("thi_nghiem")
    if isinstance(tn, dict) and (tn.get("ten") or tn.get("mo_ta")):
        st.divider()
        st.markdown("### Bộ thí nghiệm trong ảnh")
        if tn.get("ten"):
            st.markdown("**Có thể là:** " + tn["ten"])
        if tn.get("mo_ta"):
            st.markdown(tn["mo_ta"])
        bullets("Nhận xét về cách lắp", tn.get("nhan_xet_lap_rap"))
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))
    st.caption("Kết quả do AI nhận biết qua ảnh nên có thể nhầm, nhất là khi ảnh mờ hoặc dụng cụ bị che. "
               "Khi làm thí nghiệm thật, hãy làm theo hướng dẫn của giáo viên và quy định an toàn của phòng thực hành.")


def equipment_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    _m = st.radio("Chức năng", ["Nhận biết dụng cụ", "Mô phỏng thí nghiệm"], horizontal=True, key="eq_mode")
    if _m.startswith("Mô phỏng"):
        _sim_ui(top)
        return
    src = st.radio("Nguồn", ["🖼️ Tải ảnh", "📷 Chụp ảnh", "⌨️ Gõ tên dụng cụ"], horizontal=True, key="eq_src")
    img, txt = None, None
    if src.startswith("📷"):
        f = st.camera_input("Chụp dụng cụ hoặc bộ thí nghiệm", key="eq_cam")
        img = f.getvalue() if f else None
    elif src.startswith("🖼️"):
        f = st.file_uploader("Ảnh dụng cụ hoặc bộ thí nghiệm", type=["jpg", "jpeg", "png", "webp"], key="eq_up")
        img = f.getvalue() if f else None
    else:
        txt = st.text_input("Tên dụng cụ", placeholder="Ví dụ: ống sinh hàn", key="eq_txt").strip()
    if st.button("Nhận biết dụng cụ", type="primary", disabled=not (img or txt), key="eq_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang phân tích…"):
                try:
                    st.session_state.eq = analyze_equipment(img, txt)
                    add_history(txt or ', '.join(str(t.get('ten', '')) for t in ((st.session_state.eq or {}).get('dung_cu') or [])[:2]) or 'Dụng cụ', 'Dụng cụ')
                    st.session_state.n += 1
                except Exception as e:
                    st.session_state.eq = None
                    show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)
    if st.session_state.get("eq"):
        st.divider()
        render_equipment(st.session_state.eq)



# ---------------- Chẩn đoán lỗi & chấm bài ----------------
GROUPS = [
    ("ancol", "[CX4][OX2H]"), ("phenol", "c[OX2H]"), ("ete", "[CX4][OX2][CX4]"),
    ("anđehit", "[CX3;H1,H2;!$(C[OX2,NX3])]=O"), ("xeton", "[#6][CX3](=O)[#6]"),
    ("axit cacboxylic", "[CX3](=O)[OX2H1]"), ("este", "[CX3](=O)[OX2][#6]"),
    ("amin", "[NX3;!$(NC=[O,N,S])][#6]"), ("nitro", "[N+](=O)[O-]"),
    ("anken (C=C)", "[CX3]=[CX3]"), ("ankin (C≡C)", "[CX2]#[CX2]"),
    ("vòng benzen", "c1ccccc1"), ("dẫn xuất halogen", "[#6][F,Cl,Br,I]"),
]
_GP = [(n, Chem.MolFromSmarts(p)) for n, p in GROUPS]


def groups_of(m):
    return [n for n, p in _GP if m.HasSubstructMatch(p)]


def atom_counts(m):
    c = {}
    for a in Chem.AddHs(m).GetAtoms():
        c[a.GetSymbol()] = c.get(a.GetSymbol(), 0) + 1
    return c


def diff_text(ms, mt):
    a, b = atom_counts(ms), atom_counts(mt)
    order = ["C", "H"] + sorted(k for k in set(a) | set(b) if k not in ("C", "H"))
    out = []
    for k in order:
        d = a.get(k, 0) - b.get(k, 0)
        if d > 0:
            out.append(f"thừa {d} nguyên tử {k}")
        elif d < 0:
            out.append(f"thiếu {-d} nguyên tử {k}")
    return ", ".join(out)


def parse_student(smi):
    """Trả về (mol, thông báo lỗi bằng tiếng Việt)."""
    smi = (smi or "").strip()
    if not smi:
        return None, "Chưa đọc được công thức em vẽ."
    m = Chem.MolFromSmiles(smi, sanitize=False)
    if m is None:
        return None, "Công thức bị sai cú pháp (thiếu hoặc thừa ngoặc, vòng chưa đóng hoặc ký hiệu lạ)."
    try:
        Chem.SanitizeMol(m)
    except Chem.AtomValenceException as e:
        mm = re.search(r"atom # (\d+) (\w+), (\d+)", str(e))
        if mm:
            return None, f"Nguyên tử {mm.group(2)} thứ {int(mm.group(1)) + 1} có hóa trị {mm.group(3)}, vượt quá hóa trị cho phép."
        return None, "Có nguyên tử vượt quá hóa trị cho phép."
    except Chem.KekulizeException:
        return None, "Vòng thơm vẽ sai: không sắp xếp được các nối đôi xen kẽ trong vòng."
    except Exception:
        return None, "Công thức có lỗi hóa trị."
    return m, None


def draw_student(smi, size=(320, 220)):
    """Vẽ lại đúng hình học sinh đã vẽ, kể cả khi sai hóa trị; khoanh đỏ nguyên tử sai nếu có."""
    m = Chem.MolFromSmiles((smi or "").strip(), sanitize=False)
    if m is None:
        return None
    bad = []
    try:
        Chem.SanitizeMol(m)
    except Chem.AtomValenceException as e:
        mm = re.search(r"atom # (\d+)", str(e))
        if mm:
            bad = [int(mm.group(1))]
    except Exception:
        pass
    try:
        m.UpdatePropertyCache(strict=False)
        Chem.FastFindRings(m)
        return Draw.MolToImage(m, size=size, highlightAtoms=bad, highlightColor=(1.0, 0.55, 0.55))
    except Exception:
        return None


def grade_structure(smi, target_smi):
    mt = Chem.MolFromSmiles(target_smi)
    ms, err = parse_student(smi)
    tips = ["Nhớ hóa trị: C có 4 liên kết, N có 3, O có 2, H và halogen có 1."]
    if ms is None:
        return {"kq": "loi", "nx": [err], "goi_y": tips}
    fs, ft = formula_of(ms), formula_of(mt)
    if noiso(ms) == noiso(mt):
        return {"kq": "dung", "nx": ["Chính xác. Công thức cấu tạo em vẽ đúng."], "goi_y": []}
    gs, gt = set(groups_of(ms)), set(groups_of(mt))
    if fs == ft:
        nx = [f"Công thức phân tử đúng ({pretty(fs)}) nhưng cấu tạo chưa đúng yêu cầu: em đã vẽ một đồng phân khác."]
        if gs != gt:
            nx.append("Nhóm chức khác với chất cần tìm. Hình em vẽ có: " + (", ".join(sorted(gs)) or "không có nhóm chức đặc trưng")
                      + ". Đề cần: " + (", ".join(sorted(gt)) or "không có nhóm chức đặc trưng") + ".")
            nx.append("Loại đồng phân em đã vẽ: đồng phân nhóm chức.")
        else:
            nx.append("Nhóm chức đúng nhưng mạch cacbon hoặc vị trí nhóm thế/liên kết bội chưa đúng.")
        return {"kq": "dong_phan", "nx": nx, "goi_y": ["Đếm lại số cacbon ở mạch chính và vị trí gắn nhóm thế."]}
    nx = [f"Công thức phân tử của hình em vẽ là {pretty(fs)}. So với chất cần tìm, em {diff_text(ms, mt)}."]
    if gs != gt:
        nx.append("Nhóm chức chưa khớp: hình em vẽ có " + (", ".join(sorted(gs)) or "không có nhóm chức đặc trưng")
                  + "; đề cần " + (", ".join(sorted(gt)) or "không có nhóm chức đặc trưng") + ".")
    return {"kq": "sai_cong_thuc", "nx": nx, "goi_y": ["Đếm lại số nguyên tử C và số H gắn vào mỗi C."]}


def norm_name(x):
    x = unicodedata.normalize("NFD", (x or "").lower().replace("đ", "d"))
    x = "".join(c for c in x if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", x)


BANK = [
    ("metan", "C", []), ("etan", "CC", []), ("propan", "CCC", []), ("butan", "CCCC", ["n-butan"]),
    ("2-metylpropan", "CC(C)C", ["isobutan"]), ("pentan", "CCCCC", ["n-pentan"]),
    ("2-metylbutan", "CCC(C)C", ["isopentan"]), ("2,2-đimetylpropan", "CC(C)(C)C", ["neopentan"]),
    ("eten", "C=C", ["etilen"]), ("propen", "C=CC", ["propilen"]), ("but-1-en", "C=CCC", ["1-buten"]),
    ("but-2-en", "CC=CC", ["2-buten"]), ("etin", "C#C", ["axetilen"]), ("propin", "C#CC", []),
    ("benzen", "c1ccccc1", []), ("toluen", "Cc1ccccc1", ["metylbenzen"]), ("stiren", "C=Cc1ccccc1", ["vinylbenzen"]),
    ("phenol", "Oc1ccccc1", ["hiđroxibenzen"]), ("anilin", "Nc1ccccc1", ["phenylamin"]),
    ("metanol", "CO", ["ancol metylic"]), ("etanol", "CCO", ["ancol etylic"]),
    ("propan-1-ol", "CCCO", ["ancol propylic"]), ("propan-2-ol", "CC(C)O", ["ancol isopropylic"]),
    ("butan-1-ol", "CCCCO", []), ("etylen glicol", "OCCO", ["etan-1,2-điol"]),
    ("glixerol", "OCC(O)CO", ["glixerin", "propan-1,2,3-triol"]),
    ("metanal", "C=O", ["fomanđehit", "anđehit fomic"]), ("etanal", "CC=O", ["anđehit axetic", "axetanđehit"]),
    ("propanon", "CC(C)=O", ["axeton"]), ("axit metanoic", "OC=O", ["axit fomic"]),
    ("axit etanoic", "CC(=O)O", ["axit axetic"]), ("axit propanoic", "CCC(=O)O", ["axit propionic"]),
    ("axit benzoic", "OC(=O)c1ccccc1", []), ("metyl axetat", "COC(C)=O", ["metyl etanoat"]),
    ("etyl axetat", "CCOC(C)=O", ["etyl etanoat"]), ("metylamin", "CN", ["metanamin"]), ("etylamin", "CCN", ["etanamin"]),
    ("glyxin", "NCC(=O)O", ["axit aminoaxetic"]), ("alanin", "CC(N)C(=O)O", ["axit 2-aminopropanoic"]),
    ("đimetyl ete", "COC", []), ("đietyl ete", "CCOCC", []), ("etyl clorua", "CCCl", ["cloetan"]),
    ("clorofom", "ClC(Cl)Cl", ["triclometan"]),
]

GRAPH_PROMPT = RULES + """

Đây là ảnh BÀI LÀM của học sinh và sẽ được chấm điểm. Hãy CHÉP LẠI TRUNG THỰC cấu trúc học sinh đã vẽ dưới dạng ĐỒ THỊ nguyên tử - liên kết. TUYỆT ĐỐI không sửa bài, không đoán học sinh định vẽ chất nào và không tự cân lại hóa trị. Nếu hình vẽ sai (ví dụ nguyên tử C nối với 5 nét) thì vẫn chép đúng 5 nét đó.

Quy tắc:
- Mỗi nguyên tử trong hình là một đối tượng {"id":số thứ tự bắt đầu từ 1,"el":"C" hoặc "O", "N", "Cl", "Br"...,"h":số H}.
- Công thức khung (đường gấp khúc): mỗi đầu mút và mỗi chỗ gấp khúc là một nguyên tử C với "h":null (H ngầm định, không vẽ).
- Công thức thu gọn hoặc khai triển: "h" là SỐ H được ghi kèm hoặc vẽ nối trực tiếp vào nguyên tử đó. Ví dụ CH3 thì h=3; nếu có 4 nét nối từ C tới 4 chữ H thì h=4. Đếm đúng từng nét, không làm tròn theo hóa trị. Nguyên tử không ghi và không vẽ H nào thì "h":null.
- KHÔNG tạo đối tượng riêng cho các H đã tính vào "h".
- Mỗi liên kết giữa hai nguyên tử (không kể H) là [id1,id2,bậc], bậc 1, 2 hoặc 3 tính theo số nét song song.
- Vòng benzen vẽ có vòng tròn bên trong: ghi 6 nguyên tử C và các liên kết xen kẽ bậc 1 và 2.
Chỉ trả về JSON: {"atoms":[...],"bonds":[[1,2,1],...],"ghi_chu":"điều chưa rõ hoặc mơ hồ trong hình, có thể rỗng"}"""


def graph_to_smiles(g):
    """Dựng phân tử từ đồ thị AI chép lại, giữ nguyên số H đã vẽ để lỗi hóa trị không bị che mất."""
    atoms, bonds = g.get("atoms") or [], g.get("bonds") or []
    if not atoms:
        return ""
    rw, idx = Chem.RWMol(), {}
    for a in atoms:
        el = str(a.get("el", "C")).strip()
        at = Chem.Atom(el[:1].upper() + el[1:].lower())
        h = a.get("h")
        if isinstance(h, (int, float)) and not isinstance(h, bool):
            at.SetNumExplicitHs(int(h))
            at.SetNoImplicit(True)
        idx[a["id"]] = rw.AddAtom(at)
    kinds = {1: Chem.BondType.SINGLE, 2: Chem.BondType.DOUBLE, 3: Chem.BondType.TRIPLE}
    for b in bonds:
        i, j, o = b[0], b[1], int(b[2])
        if i != j and rw.GetBondBetweenAtoms(idx[i], idx[j]) is None:
            rw.AddBond(idx[i], idx[j], kinds[o])
    m = rw.GetMol()
    m.UpdatePropertyCache(strict=False)
    Chem.FastFindRings(m)
    return Chem.MolToSmiles(m)


def transcribe_student(img):
    r = ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"), GRAPH_PROMPT], temperature=0)
    try:
        smi = graph_to_smiles(r)
    except Exception:
        raise ValueError("Không dựng được cấu trúc từ ảnh, hãy chụp rõ hơn.")
    return smi, r.get("ghi_chu", "")


def analyze_image_strict(img):
    """Đọc ảnh trung thực (đồ thị nguyên tử) rồi mới gọi tên và tính chất nếu công thức hợp lệ."""
    smi, note = transcribe_student(img)
    res = {"smiles": smi}
    if smi and parse_student(smi)[0] is not None:
        res = analyze_smiles(smi)
        res["smiles"] = smi
        res.pop("do_tin_cay", None)
    if note:
        old = res.get("luu_y") or []
        res["luu_y"] = ([old] if isinstance(old, str) else list(old)) + [note]
    return res


RANDOM_Q = "🎲 Ngẫu nhiên"


def new_question(kind):
    st.session_state.pq = {"kind": kind, "idx": random.randrange(len(BANK)), "result": None}
    st.session_state["pr_pick"] = RANDOM_Q


def practice_ui():
    st.session_state.setdefault("pq", None)
    st.session_state.setdefault("ph", [])
    top = st.container()  # kết quả chấm và lý do sai hiện ở đầu tab
    c1, c2 = st.columns([3, 2])
    with c1:
        kind = st.radio("Dạng bài", ["✍️ Vẽ công thức cấu tạo", "🏷️ Gọi tên chất"], horizontal=True, key="pr_kind")
    with c2:
        who = st.text_input("Tên học sinh (không bắt buộc)", key="pr_name")
    if st.session_state.pq is None or st.session_state.pq["kind"] != kind:
        new_question(kind)
    names = [b[0] for b in BANK]
    pick = st.selectbox("Chọn đề cụ thể (hoặc để ngẫu nhiên)", [RANDOM_Q] + names, key="pr_pick")
    if pick != RANDOM_Q and names.index(pick) != st.session_state.pq["idx"]:
        st.session_state.pq = {"kind": kind, "idx": names.index(pick), "result": None}
    q = st.session_state.pq
    ten, smi, alias = BANK[q["idx"]]
    mt = Chem.MolFromSmiles(smi)

    with st.container(border=True):
        if kind.startswith("✍️"):
            st.markdown(f"#### Viết công thức cấu tạo của: **{ten}**")
            if st.checkbox("Gợi ý: hiện công thức phân tử", key=f"pr_hint_{q['idx']}"):
                st.markdown(f"Công thức phân tử: **{pretty(formula_of(mt))}**")
            how = st.radio("Nộp bài bằng", ["🖼️ Ảnh bài làm", "📷 Chụp ảnh", "⌨️ Nhập SMILES"], horizontal=True, key="pr_how")
            img, ans = None, ""
            if how.startswith("🖼️"):
                f = st.file_uploader("Ảnh bài làm", type=["jpg", "jpeg", "png", "webp"], key="pr_up")
                img = f.getvalue() if f else None
            elif how.startswith("📷"):
                f = st.camera_input("Chụp bài làm", key="pr_cam")
                img = f.getvalue() if f else None
            else:
                ans = st.text_input("SMILES bài làm", key="pr_smi").strip()
            go = st.button("Nộp bài", type="primary", disabled=not (img or ans), key="pr_go")
        else:
            st.markdown("#### Gọi tên chất sau")
            st.image(Draw.MolToImage(mt, size=(360, 240)))
            ans = st.text_input("Tên chất (theo SGK)", key="pr_name_ans").strip()
            img = None
            go = st.button("Nộp bài", type="primary", disabled=not ans, key="pr_go2")

    if go and q["result"] is None:
        if img and st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            try:
                note, given = "", ans
                if kind.startswith("✍️"):
                    if img:
                        with st.spinner("Đang đọc bài làm…"):
                            given, note = transcribe_student(img)
                        st.session_state.n += 1
                    res = grade_structure(given, smi)
                else:
                    ok = norm_name(ans) in {norm_name(x) for x in [ten] + alias}
                    res = {"kq": "dung" if ok else "sai",
                           "nx": ["Chính xác." if ok else "Chưa đúng. Hãy xác định mạch chính, nhóm chức và vị trí nhóm thế rồi gọi tên lại."],
                           "goi_y": []}
                res.update({"given": given, "note": note})
                q["result"] = res
                add_history(ten, 'Luyện tập')
                st.session_state.ph.append({
                    "Thời gian": datetime.now().strftime("%H:%M:%S"), "Học sinh": who, "Dạng bài": "vẽ CTCT" if kind.startswith("✍️") else "gọi tên",
                    "Đề": ten, "Bài làm": given, "Kết quả": res["kq"]})
            except Exception as e:
                show_top_error(top, "Có lỗi khi chấm bài, hãy thử lại.", e)

    with top:
        res = q["result"]
        if res:
            head = {"dung": "Đúng rồi!", "dong_phan": "CHƯA ĐÚNG: đã vẽ một đồng phân khác", "sai_cong_thuc": "CHƯA ĐÚNG: sai công thức phân tử",
                    "loi": "CHƯA ĐÚNG: công thức có lỗi", "sai": "CHƯA ĐÚNG"}[res["kq"]]
            if res["kq"] == "dung":
                st.success("**" + head + "**")
                for t in res["nx"]:
                    st.markdown("- " + t)
            else:
                st.markdown(errbanner(head, *res["nx"]), unsafe_allow_html=True)
            st.markdown("**Đề:** " + (f"viết công thức cấu tạo của **{ten}**" if kind.startswith("✍️") else "gọi tên chất trong hình"))
            for t in res["goi_y"]:
                st.markdown("- 💡 " + t)
            if res.get("note"):
                st.caption("Ghi chú khi đọc ảnh: " + res["note"])
            if kind.startswith("✍️") and res.get("given"):
                im = draw_student(res["given"])
                if im is not None:
                    st.markdown("**Hình em vẽ (máy đọc lại đúng như hình, giữ nguyên lỗi" + (", nguyên tử sai khoanh đỏ" if res["kq"] == "loi" else "") + "):**")
                    st.image(im)
            with st.expander("Xem đáp án"):
                st.image(Draw.MolToImage(mt, size=(300, 200)))
                st.markdown(f"**{ten}** · {pretty(formula_of(mt))} · `{smi}`")
                gl = groups_of(mt)
                if gl:
                    st.markdown("Nhóm chức/đặc điểm: " + ", ".join(gl))
            st.button("➡️ Câu tiếp theo", on_click=new_question, args=(kind,), key="pr_next")

    hist = st.session_state.ph
    if hist:
        st.divider()
        n_ok = sum(1 for h in hist if h["Kết quả"] == "dung")
        a, b, c = st.columns(3)
        a.metric("Đã làm", len(hist))
        b.metric("Đúng", n_ok)
        c.metric("Tỉ lệ đúng", f"{round(100 * n_ok / len(hist))}%")
        st.dataframe(hist, hide_index=True, width="stretch")
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(hist[0].keys()))
        w.writeheader()
        w.writerows(hist)
        st.download_button("⬇️ Tải phiếu kết quả (CSV, nộp cho giáo viên)", buf.getvalue().encode("utf-8-sig"),
                           file_name="phieu_ket_qua.csv", mime="text/csv", key="pr_dl")
    st.caption("Phần chấm công thức do RDKit so sánh cấu tạo nên chính xác. Phần đọc ảnh bài làm do AI thực hiện nên có thể đọc sai nét viết khó.")


# ---------------- Đánh giá độ chính xác (giáo viên / nhóm nghiên cứu) ----------------
def parse_truth(text):
    rows = {}
    for line in (text or "").splitlines():
        parts = [p.strip() for p in line.replace(";", ",").split(",")]
        if len(parts) >= 2 and parts[0] and parts[1]:
            rows[parts[0]] = (parts[1], parts[2] if len(parts) > 2 else "")
    return rows


def bench_row(pred, truth):
    mp, mt = Chem.MolFromSmiles(pred or ""), Chem.MolFromSmiles(truth)
    if mp is None or mt is None:
        return {"hop_le": False, "khop_cong_thuc": False, "khop_cau_tao": False, "khop_hoan_toan": False}
    return {"hop_le": True, "khop_cong_thuc": formula_of(mp) == formula_of(mt),
            "khop_cau_tao": noiso(mp) == noiso(mt),
            "khop_hoan_toan": Chem.MolToSmiles(mp) == Chem.MolToSmiles(mt)}


def bench_summary(rows):
    n = len(rows)
    pct = lambda k: round(100 * sum(1 for r in rows if r[k]) / n, 1) if n else 0
    return {"n": n, "hop_le": pct("hop_le"), "khop_cong_thuc": pct("khop_cong_thuc"),
            "khop_cau_tao": pct("khop_cau_tao"), "khop_hoan_toan": pct("khop_hoan_toan"),
            "tb_giay": round(sum(r["giay"] for r in rows) / n, 1) if n else 0}


def bench_ui():
    code = st.secrets.get("TEACHER_CODE", "")
    if code and st.text_input("Mã giáo viên", type="password", key="bm_code") != code:
        st.info("Tab này dành cho giáo viên hoặc nhóm nghiên cứu. Nhập mã giáo viên để dùng.")
        return
    st.markdown("Tải lên một bộ ảnh đã biết đáp án để đo **độ chính xác nhận dạng**. Kết quả có thể đưa vào báo cáo đề tài.")
    files = st.file_uploader("Ảnh kiểm thử", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key="bm_files")
    up = st.file_uploader("File đáp án CSV (tên_file, SMILES, nhóm)", type=["csv", "txt"], key="bm_csv")
    txt = st.text_area("Hoặc dán đáp án, mỗi dòng: tên_file, SMILES, nhóm", height=130, key="bm_truth",
                       placeholder="03_toluen.jpg, Cc1ccccc1, benzen\n06_etyl_axetat.jpg, CCOC(C)=O, công thức khung")
    method = st.radio("Phương pháp đọc ảnh", ["Đồ thị nguyên tử (trung thực với bài vẽ)", "SMILES trực tiếp"], horizontal=True, key="bm_method")
    src = up.getvalue().decode("utf-8-sig", errors="ignore") if up else txt
    truth = parse_truth(src)
    todo = [f for f in (files or []) if f.name in truth][:BENCH_MAX]
    if files and truth:
        st.caption(f"Có {len(todo)}/{len(files)} ảnh khớp tên trong file đáp án (tối đa {BENCH_MAX} ảnh mỗi lần chạy).")
    if st.button("Chạy đánh giá", type="primary", disabled=not todo, key="bm_go"):
        rows, bar = [], st.progress(0.0, text="Đang chạy…")
        for i, f in enumerate(todo):
            t0 = time.time()
            pred, err = "", ""
            try:
                pred = transcribe_student(f.getvalue())[0] if method.startswith("Đồ thị") else to_smiles(f.getvalue(), "")
            except Exception as e:
                err = type(e).__name__
            tr, grp = truth[f.name]
            r = {"tep": f.name, "nhom": grp, "dap_an": tr, "phuong_phap": "do_thi" if method.startswith("Đồ thị") else "smiles", "ket_qua_may": pred, "loi": err, "giay": round(time.time() - t0, 1)}
            r.update(bench_row(pred, tr))
            rows.append(r)
            bar.progress((i + 1) / len(todo), text=f"Đã xử lý {i + 1}/{len(todo)}")
        st.session_state.bm = rows
    rows = st.session_state.get("bm")
    if rows:
        sm = bench_summary(rows)
        a, b, c, d = st.columns(4)
        a.metric("Số ảnh", sm["n"])
        b.metric("Khớp cấu tạo", f"{sm['khop_cau_tao']}%")
        c.metric("Khớp công thức PT", f"{sm['khop_cong_thuc']}%")
        d.metric("SMILES hợp lệ", f"{sm['hop_le']}%")
        st.caption(f"Khớp hoàn toàn (kể cả lập thể): {sm['khop_hoan_toan']}% · thời gian trung bình: {sm['tb_giay']} giây/ảnh")
        grp = sorted({r["nhom"] for r in rows if r["nhom"]})
        if grp:
            st.markdown("**Theo nhóm nét vẽ**")
            st.dataframe([{"Nhóm": g, **{k: v for k, v in bench_summary([r for r in rows if r["nhom"] == g]).items()
                                         if k in ("n", "khop_cau_tao", "khop_cong_thuc", "hop_le")}} for g in grp],
                         hide_index=True, width="stretch")
        st.dataframe(rows, hide_index=True, width="stretch")
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
        st.download_button("⬇️ Tải bảng kết quả (CSV)", buf.getvalue().encode("utf-8-sig"), file_name="ket_qua_danh_gia.csv",
                           mime="text/csv", key="bm_dl")
    st.caption("Khớp cấu tạo: SMILES máy đọc được và đáp án giống nhau về cấu tạo (RDKit chuẩn hóa trước khi so sánh).")


# ---------------- Sơ đồ tư duy ----------------
MM_COLORS = ["#e11d48", "#f59e0b", "#16a34a", "#0ea5e9", "#8b5cf6", "#ec4899", "#14b8a6", "#f97316"]
MM_FONT = "Arial, Helvetica, sans-serif"

MM_PROMPT = """Bạn là giáo viên Hóa học THPT Việt Nam. Hãy lập SƠ ĐỒ TƯ DUY cho chủ đề được yêu cầu, đúng nội dung SGK Việt Nam.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"chu_de":"tên chủ đề ngắn gọn","chat_smiles":"SMILES nếu chủ đề là MỘT chất cụ thể (ví dụ etanol), nếu là nhóm chất hoặc khái niệm thì null",
"nhanh":[{"ten":"tên nhánh, tối đa 24 ký tự","y":["ý ngắn, tối đa 80 ký tự"]}]}
Quy tắc:
- Chọn các nhánh phù hợp chủ đề, ví dụ: Khái niệm và cấu tạo, Danh pháp, Tính chất vật lí, Tính chất hóa học, Điều chế, Ứng dụng, Nhận biết, Lưu ý. Với chủ đề phản ứng thì dùng: Định nghĩa, Điều kiện, Ví dụ, Quy tắc, Nhận biết.
- Mỗi ý ngắn gọn, dễ nhớ. Viết công thức bằng chữ thường (CH3-CH2-OH). Phương trình viết dạng: A + B → C (điều kiện).
- KHÔNG nêu số lượng đồng phân hay khối lượng mol (bộ máy tính toán sẽ thêm riêng). Không bịa số liệu."""


def _wrap(text, width):
    lines, cur = [], ""
    for w in str(text).split():
        if len(cur) + len(w) + (1 if cur else 0) <= width:
            cur = (cur + " " + w).strip()
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def _tint(hexc, a=0.12):
    r, g, b = int(hexc[1:3], 16), int(hexc[3:5], 16), int(hexc[5:7], 16)
    return f"rgba({r},{g},{b},{a})"


def mindmap_layout(title, branches):
    """Tính vị trí mọi khung và đường nối (đơn vị px). Dùng chung cho SVG, PowerPoint, draw.io và FreeMind."""
    LEAF_W, LH = 30, 19
    cw_leaf, cw_br, cw_root = 7.9, 9.8, 14.0
    root_lines = _wrap(title, 16)
    rw = max(len(x) for x in root_lines) * cw_root + 56
    rh = len(root_lines) * 28 + 36
    items = []
    for i, b in enumerate(branches):
        blines = _wrap(b["ten"], 18)
        bw = max(len(x) for x in blines) * cw_br + 36
        bh = len(blines) * 22 + 20
        leaves = []
        for t in b.get("y", []):
            ll = _wrap(t, LEAF_W)
            leaves.append((t, ll, max(len(x) for x in ll) * cw_leaf + 24, len(ll) * LH + 14))
        block = sum(l[3] for l in leaves) + 8 * max(0, len(leaves) - 1)
        items.append({"b": b, "blines": blines, "bw": bw, "bh": bh, "leaves": leaves, "block": block,
                      "h": max(block, bh), "color": b.get("color") or MM_COLORS[i % len(MM_COLORS)]})
    half = (len(items) + 1) // 2
    sides = {"R": items[:half], "L": items[half:]}
    GAP, MARG = 28, 50
    tot = {k: sum(x["h"] for x in v) + GAP * max(0, len(v) - 1) for k, v in sides.items()}
    H = max(max(tot.values()), rh) + 2 * MARG
    reach = 0
    for v in sides.values():
        for x in v:
            reach = max(reach, 70 + x["bw"] + 55 + max([l[2] for l in x["leaves"]] or [0]))
    W = 2 * (rw / 2 + reach + MARG)
    cx, cy = W / 2, H / 2
    nodes = [{"id": "n0", "kind": "root", "x": cx - rw / 2, "y": cy - rh / 2, "w": rw, "h": rh, "text": title,
              "lines": root_lines, "color": "#4f46e5", "side": None, "parent": None}]
    edges, k = [], 1
    for side, v in sides.items():
        sgn = 1 if side == "R" else -1
        y = cy - tot[side] / 2
        for x in v:
            col, yb = x["color"], y + x["h"] / 2
            bx0 = cx + sgn * (rw / 2 + 70)
            bx1 = bx0 + sgn * x["bw"]
            bid = f"n{k}"
            k += 1
            nodes.append({"id": bid, "kind": "branch", "x": min(bx0, bx1), "y": yb - x["bh"] / 2, "w": x["bw"], "h": x["bh"],
                          "text": x["b"]["ten"], "lines": x["blines"], "color": col, "side": side, "parent": "n0"})
            edges.append({"a": "n0", "b": bid, "color": col, "width": 5, "x1": cx + sgn * rw / 2, "y1": cy, "x2": bx0, "y2": yb, "side": side})
            ly = y + (x["h"] - x["block"]) / 2
            lx0 = bx1 + sgn * 55
            for text, ll, lw, lh in x["leaves"]:
                lid = f"n{k}"
                k += 1
                nodes.append({"id": lid, "kind": "leaf", "x": lx0 if side == "R" else lx0 - lw, "y": ly, "w": lw, "h": lh,
                              "text": text, "lines": ll, "color": col, "side": side, "parent": bid})
                edges.append({"a": bid, "b": lid, "color": col, "width": 2.5, "x1": bx1, "y1": yb, "x2": lx0, "y2": ly + lh / 2, "side": side})
                ly += lh + 8
            y += x["h"] + GAP
    return {"W": W, "H": H, "nodes": nodes, "edges": edges}


def svg_from_layout(L):
    esc = html.escape
    W, H = L["W"], L["H"]
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W:.0f}" height="{H:.0f}" viewBox="0 0 {W:.0f} {H:.0f}" '
           f'font-family="{MM_FONT}"><defs><linearGradient id="rg" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#4f46e5"/>'
           f'<stop offset="1" stop-color="#14b8a6"/></linearGradient></defs><rect width="100%" height="100%" fill="#ffffff"/>']
    for e in L["edges"]:
        xm = (e["x1"] + e["x2"]) / 2
        out.append(f'<path d="M{e["x1"]:.1f} {e["y1"]:.1f} C{xm:.1f} {e["y1"]:.1f} {xm:.1f} {e["y2"]:.1f} {e["x2"]:.1f} {e["y2"]:.1f}" '
                   f'fill="none" stroke="{e["color"]}" stroke-width="{e["width"]}" stroke-linecap="round"/>')
    texts = []
    for n in L["nodes"]:
        x, y, w, h, col = n["x"], n["y"], n["w"], n["h"], n["color"]
        cxn, cyn = x + w / 2, y + h / 2
        if n["kind"] == "root":
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{h / 2:.1f}" fill="url(#rg)"/>')
            ty = cyn - len(n["lines"]) * 28 / 2 + 21
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{cxn:.1f}" y="{ty + j * 28:.1f}" text-anchor="middle" font-size="24" font-weight="800" fill="#fff">{esc(ln)}</text>')
        elif n["kind"] == "branch":
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="14" fill="{col}"/>')
            ty = cyn - len(n["lines"]) * 22 / 2 + 16
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{cxn:.1f}" y="{ty + j * 22:.1f}" text-anchor="middle" font-size="17" font-weight="700" fill="#fff">{esc(ln)}</text>')
        else:
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="10" fill="{_tint(col)}" stroke="{col}" stroke-width="1.5"/>')
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{x + 12:.1f}" y="{y + 7 + 19 * (j + 1) - 5:.1f}" font-size="15" fill="#1f2937">{esc(ln)}</text>')
    return "".join(out + texts) + "</svg>"


def mindmap_svg(title, branches):
    L = mindmap_layout(title, branches)
    return svg_from_layout(L), L["W"], L["H"]


def _blend(hexc, a=0.12):
    r, g, b = int(hexc[1:3], 16), int(hexc[3:5], 16), int(hexc[5:7], 16)
    return "{:02x}{:02x}{:02x}".format(*(round(c * a + 255 * (1 - a)) for c in (r, g, b)))


def mm_pptx_bytes(L):
    """PowerPoint: mỗi khung là hình có chữ sửa được, đường nối dính vào khung nên kéo khung thì đường nối đi theo."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
    except ImportError:
        return None
    E = 9525
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(L["W"] * E)), Emu(int(L["H"] * E))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    shapes = {}
    for n in L["nodes"]:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(n["x"] * E)), Emu(int(n["y"] * E)), Emu(int(n["w"] * E)), Emu(int(n["h"] * E)))
        shp.adjustments[0] = min(0.5, (n["h"] / 2 if n["kind"] == "root" else (14 if n["kind"] == "branch" else 10)) / min(n["w"], n["h"]))
        shp.shadow.inherit = False
        shp.fill.solid()
        if n["kind"] == "leaf":
            shp.fill.fore_color.rgb = RGBColor.from_string(_blend(n["color"]).upper())
            shp.line.color.rgb = rgb(n["color"])
            shp.line.width = Pt(1.2)
        else:
            shp.fill.fore_color.rgb = rgb(n["color"])
            shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(int(10 * E))
        tf.margin_top = tf.margin_bottom = Emu(int(4 * E))
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if n["kind"] == "leaf" else PP_ALIGN.CENTER
        run = para.add_run()
        run.text = n["text"]
        run.font.name = "Arial"
        run.font.size = Pt({"root": 18, "branch": 13, "leaf": 11}[n["kind"]])
        run.font.bold = n["kind"] != "leaf"
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37) if n["kind"] == "leaf" else RGBColor(255, 255, 255)
        shapes[n["id"]] = shp
    tree = slide.shapes._spTree
    for e in L["edges"]:
        c = slide.shapes.add_connector(MSO_CONNECTOR.CURVE, Emu(int(e["x1"] * E)), Emu(int(e["y1"] * E)), Emu(int(e["x2"] * E)), Emu(int(e["y2"] * E)))
        right = e["side"] == "R"
        c.begin_connect(shapes[e["a"]], 3 if right else 1)
        c.end_connect(shapes[e["b"]], 1 if right else 3)
        c.line.color.rgb = rgb(e["color"])
        c.line.width = Pt(e["width"] * 0.75)
        tree.remove(c._element)
        tree.insert(2, c._element)  # đưa đường nối xuống dưới các khung
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def mm_drawio(L):
    """File .drawio (mở bằng app.diagrams.net hoặc draw.io Desktop)."""
    def attr(t):
        return html.escape(html.escape(str(t), quote=False), quote=True)
    cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
    for n in L["nodes"]:
        if n["kind"] == "root":
            st_ = "rounded=1;arcSize=50;whiteSpace=wrap;html=1;fillColor=#4f46e5;strokeColor=none;fontColor=#ffffff;fontStyle=1;fontSize=18;"
        elif n["kind"] == "branch":
            st_ = f"rounded=1;arcSize=30;whiteSpace=wrap;html=1;fillColor={n['color']};strokeColor=none;fontColor=#ffffff;fontStyle=1;fontSize=13;"
        else:
            st_ = (f"rounded=1;arcSize=20;whiteSpace=wrap;html=1;fillColor=#{_blend(n['color'])};strokeColor={n['color']};"
                   "fontColor=#1f2937;fontSize=11;align=left;spacingLeft=6;")
        cells.append(f'<mxCell id="{n["id"]}" value="{attr(n["text"])}" style="{st_}" vertex="1" parent="1">'
                     f'<mxGeometry x="{n["x"]:.0f}" y="{n["y"]:.0f}" width="{n["w"]:.0f}" height="{n["h"]:.0f}" as="geometry"/></mxCell>')
    for i, e in enumerate(L["edges"]):
        ex, nx = ("1", "0") if e["side"] == "R" else ("0", "1")
        cells.append(f'<mxCell id="e{i}" style="curved=1;endArrow=none;html=1;strokeColor={e["color"]};strokeWidth={e["width"]};'
                     f'exitX={ex};exitY=0.5;entryX={nx};entryY=0.5;" edge="1" parent="1" source="{e["a"]}" target="{e["b"]}">'
                     '<mxGeometry relative="1" as="geometry"/></mxCell>')
    return ('<mxfile host="app.diagrams.net"><diagram name="So do tu duy" id="mm1"><mxGraphModel grid="1" gridSize="10" page="0">'
            "<root>" + "".join(cells) + "</root></mxGraphModel></diagram></mxfile>")


def mm_freemind(L):
    """File .mm (FreeMind) mở được bằng XMind, Freeplane, MindManager, Mindomo."""
    q = lambda t: html.escape(str(t), quote=True)
    root = L["nodes"][0]
    out = [f'<map version="1.0.1"><node TEXT="{q(root["text"])}">']
    for b in [n for n in L["nodes"] if n["kind"] == "branch"]:
        pos = "right" if b["side"] == "R" else "left"
        out.append(f'<node TEXT="{q(b["text"])}" POSITION="{pos}" COLOR="{b["color"]}">')
        out += [f'<node TEXT="{q(l["text"])}"/>' for l in L["nodes"] if l["parent"] == b["id"]]
        out.append("</node>")
    out.append("</node></map>")
    return "".join(out)


def rows_to_branches(rows):
    """Bảng chỉnh sửa (Nhánh, Ý) -> danh sách nhánh. Dòng không ghi nhánh thì thuộc nhánh ngay trên."""
    order, data = [], {}
    for r_ in rows:
        n = str(r_.get("Nhánh") or "").strip()
        y = str(r_.get("Ý") or "").strip()
        if not n and not y:
            continue
        n = n or (order[-1] if order else "Nhánh mới")
        if n not in data:
            data[n] = []
            order.append(n)
        if y:
            data[n].append(y[:100])
    return [{"ten": n[:28], "y": data[n], **({"color": "#475569"} if n.startswith("Số liệu (máy tính)") else {})} for n in order[:10]]


def mm_outline(title, branches):
    lines = [f"# {title}"]
    for b in branches:
        lines.append(f"\n## {b['ten']}")
        lines += [f"- {t}" for t in b.get("y", [])]
    return "\n".join(lines)


def mm_clean(raw):
    """Giữ sơ đồ gọn: tối đa 8 nhánh, 5 ý mỗi nhánh, cắt câu quá dài."""
    out = []
    for b in (raw.get("nhanh") or [])[:8]:
        ten = str(b.get("ten", "")).strip()[:28]
        ys = [str(t).strip()[:100] for t in (b.get("y") or []) if str(t).strip()][:5]
        if ten and ys:
            out.append({"ten": ten, "y": ys})
    return out


def mm_engine_facts(smiles):
    """Số liệu do RDKit và bộ máy đồng phân tính, không phải AI."""
    m, _ = parse_student(smiles or "")
    if m is None:
        return None
    f = formula_of(m)
    facts = [f"Công thức phân tử: {pretty(f)}", f"Khối lượng mol: {Descriptors.MolWt(m):.2f} g/mol",
             "Nhóm chức / loại: " + ", ".join(iso_labels(m))]
    counts, nH = mol_counts(m)
    if sum(counts.values()) <= 10 and all(e in VALENCE for e in counts):
        try:
            res = build_iso_result(hill_key(counts, nH), None, dict(PRESETS["mono"]))
            if res["ok"] and res["kept"]:
                k = res["kept"]
                facts.append(f"Số đồng phân cấu tạo đơn chức mạch hở của {pretty(f)}: {len(k)} (kể cả lập thể: {sum(x['stereo'] for x in k)})")
        except Exception:
            pass
    return facts


def mm_component(svg, W, H):
    scale = min(1.0, 1000.0 / W)
    page = ('<div style="font-family:Arial,sans-serif"><button id="dl" style="background:#1f7a4d;color:#fff;border:0;border-radius:10px;'
            'padding:9px 16px;font-weight:700;cursor:pointer">⬇️ Tải ảnh PNG</button>'
            '<div style="overflow:auto;border:1px solid #d6ece5;border-radius:12px;margin-top:8px;background:#fff">'
            '<div style="width:max-content;max-width:100%">' + svg.replace("<svg ", '<svg style="max-width:100%;height:auto;display:block" ', 1) +
            '</div></div><script>document.getElementById("dl").onclick=function(){'
            'var s=document.querySelector("svg"),x=new XMLSerializer().serializeToString(s),i=new Image();'
            'i.onload=function(){var c=document.createElement("canvas");c.width=' + str(int(W * 2)) + ';c.height=' + str(int(H * 2)) + ';'
            'var g=c.getContext("2d");g.fillStyle="#fff";g.fillRect(0,0,c.width,c.height);g.drawImage(i,0,0,c.width,c.height);'
            'c.toBlob(function(b){var a=document.createElement("a");a.href=URL.createObjectURL(b);a.download="so_do_tu_duy.png";a.click();});};'
            'i.src="data:image/svg+xml;charset=utf-8,"+encodeURIComponent(x);};</script></div>')
    height = int(H * scale) + 90
    if hasattr(st, "iframe"):  # st.components.v1.html sẽ bị gỡ; nội dung trong page đã được escape, không chứa HTML từ AI
        st.iframe(page, height=height)
    else:
        components.html(page, height=height, scrolling=True)


MM_EXAMPLES = ["Ancol", "Anđehit và xeton", "Axit cacboxylic", "Este và lipit", "Amin và amino axit", "Benzen và đồng đẳng",
               "Phenol", "Etanol", "Phản ứng cộng"]


def _set_topic(t):
    st.session_state["mm_topic"] = t


MM_PHOTO_PROMPT = """Đây là ảnh SƠ ĐỒ TƯ DUY HÓA HỌC được vẽ tay trên giấy. Hãy đọc và số hóa lại ĐÚNG cấu trúc và nội dung đã vẽ.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"chu_de":"chủ đề ở giữa sơ đồ","nhanh":[{"ten":"tên nhánh chính nối trực tiếp với chủ đề, tối đa 28 ký tự","mau":"mã màu #rrggbb của nhánh nếu trên giấy có tô màu, nếu không thì null","y":["các ý thuộc nhánh này, theo thứ tự đọc"]}],"ghi_chu":"chữ khó đọc hoặc chỗ không chắc chắn, nếu không có thì null"}
Quy tắc:
- Giữ đúng chữ và cách diễn đạt trên giấy. KHÔNG thêm ý mới, KHÔNG sửa nội dung hóa học. Chỉ sửa lỗi chính tả hiển thị rõ ràng.
- Nhánh phụ cấp 2, cấp 3 gộp vào danh sách y của nhánh chính. Dùng dạng "Tên nhánh phụ: nội dung" hoặc "A → B" để giữ quan hệ.
- Viết công thức bằng chữ thường (CH3-CH2-OH). Phương trình viết dạng: A + B → C (điều kiện).
- Chữ không đọc được thì ghi [?] tại chỗ đó và nêu trong ghi_chu.
- Nếu ảnh không phải sơ đồ tư duy: trả {"chu_de":"","nhanh":[],"ghi_chu":"lý do"}."""

_HEX6 = re.compile(r"^#[0-9a-fA-F]{6}$")


def mm_clean_photo(raw):
    """Giữ gần nguyên bản vẽ tay: tối đa 10 nhánh, 8 ý mỗi nhánh, giữ màu nhánh nếu hợp lệ."""
    out = []
    for b in (raw.get("nhanh") or [])[:10]:
        ten = str(b.get("ten", "")).strip()[:28]
        ys = [str(t).strip()[:100] for t in (b.get("y") or []) if str(t).strip()][:8]
        if ten:
            item = {"ten": ten, "y": ys or [""]}
            c = str(b.get("mau") or "").strip()
            if _HEX6.match(c):
                item["color"] = c.lower()
            out.append(item)
    return out


def _mm_photo_inputs(top):
    src = st.radio("Nguồn ảnh", ["🖼️ Tải ảnh lên", "📷 Chụp ảnh"], horizontal=True, key="mm_src")
    if src.startswith("📷"):
        f = st.camera_input("Chụp sơ đồ vẽ tay (đủ sáng, chụp thẳng góc, chữ rõ nét)", key="mm_cam")
    else:
        f = st.file_uploader("Chọn ảnh sơ đồ tư duy vẽ tay (đủ sáng, chụp thẳng góc, chữ rõ nét)", type=["jpg", "jpeg", "png", "webp"], key="mm_up")
    img = f.getvalue() if f else None
    if img:
        st.image(img, caption="Sơ đồ bạn đã vẽ", width=380)
    if st.button("Số hóa sơ đồ này", type="primary", disabled=not img, key="mm_photo_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
            return
        with st.spinner("Đang đọc sơ đồ vẽ tay…"):
            try:
                raw = ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"), MM_PHOTO_PROMPT], temperature=0.1)
                branches = mm_clean_photo(raw)
                if not branches:
                    show_top_error(top, "Chưa đọc được sơ đồ trong ảnh. " + str(raw.get("ghi_chu") or "Hãy chụp rõ hơn, đủ sáng và thẳng góc."))
                else:
                    title = str(raw.get("chu_de") or "Sơ đồ tư duy").strip()[:60] or "Sơ đồ tư duy"
                    note = raw.get("ghi_chu")
                    st.session_state.mm = {"title": title, "branches": branches, "smiles": "", "rev": int(time.time() * 1000),
                                           "note": note if note and str(note).lower() != "null" else "", "photo": img}
                    add_history(title, "Sơ đồ")
                    st.session_state.n += 1
            except Exception as e:
                st.session_state.mm = None
                show_top_error(top, "Có lỗi khi đọc ảnh, hãy thử lại hoặc chụp rõ hơn.", e)


def _mm_topic_inputs(top):
    st.text_input("Chủ đề hoặc tên chất", key="mm_topic", placeholder="Ví dụ: Ancol, Etanol, Phản ứng thế, Este và lipit")
    cols = st.columns(len(MM_EXAMPLES))
    for c, t in zip(cols, MM_EXAMPLES):
        c.button(t, key=f"mm_ex_{t}", on_click=_set_topic, args=(t,), width="stretch")
    level = st.radio("Mức chi tiết", ["Gọn (4 nhánh, mỗi nhánh 2-3 ý)", "Đầy đủ (6 nhánh, mỗi nhánh 3-4 ý)"], horizontal=True, key="mm_level")
    topic = st.session_state.get("mm_topic", "").strip()
    if st.button("Tạo sơ đồ tư duy", type="primary", disabled=not topic, key="mm_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang vẽ sơ đồ…"):
                try:
                    want = "4 nhánh, mỗi nhánh 2-3 ý" if level.startswith("Gọn") else "6 nhánh, mỗi nhánh 3-4 ý"
                    raw = ask_ai([MM_PROMPT + f"\nSố nhánh và ý: {want}.\n\nChủ đề: {topic}"], temperature=0.2)
                    branches = mm_clean(raw)
                    if not branches:
                        show_top_error(top, "Chưa tạo được sơ đồ cho chủ đề này, hãy thử diễn đạt rõ hơn.")
                    else:
                        smi, facts = (raw.get("chat_smiles") or "").strip(), None
                        if smi and smi.lower() != "null":
                            facts = mm_engine_facts(smi)
                        if facts:
                            branches.append({"ten": "Số liệu (máy tính)", "y": facts, "color": "#475569"})
                        st.session_state.mm = {"title": str(raw.get("chu_de") or topic)[:60], "branches": branches,
                                               "smiles": smi if facts else "", "rev": int(time.time() * 1000)}
                        add_history(st.session_state.mm['title'], 'Sơ đồ')
                        st.session_state.n += 1
                except Exception as e:
                    st.session_state.mm = None
                    show_top_error(top, "Có lỗi khi tạo sơ đồ, hãy thử lại.", e)


def _mm_result():
    r = st.session_state.get("mm")
    if r:
        st.divider()
        rev = r.setdefault("rev", 0)
        if r.get("note"):
            st.info("Ghi chú khi đọc ảnh: " + str(r["note"]))
        if r.get("photo"):
            with st.expander("Ảnh gốc để đối chiếu"):
                st.image(r["photo"])
        title = st.text_input("Chủ đề ở giữa sơ đồ", value=r["title"], key=f"mm_title_{rev}").strip() or r["title"]
        with st.expander("✏️ Chỉnh sửa nội dung (sửa chữ, thêm hoặc xóa dòng, sơ đồ cập nhật ngay)", expanded=False):
            base = [{"Nhánh": b["ten"], "Ý": t} for b in r["branches"] for t in (b["y"] or [""])]
            edited = st.data_editor(base, num_rows="dynamic", hide_index=True, key=f"mm_edit_{rev}", width="stretch",
                                    column_config={"Nhánh": st.column_config.TextColumn("Nhánh", width="medium"),
                                                   "Ý": st.column_config.TextColumn("Ý", width="large")})
            st.caption("Muốn thêm nhánh mới: thêm dòng mới và gõ tên nhánh mới. Dòng để trống cột Nhánh sẽ thuộc nhánh ngay phía trên.")
        rows = edited.to_dict("records") if hasattr(edited, "to_dict") else list(edited)
        branches = rows_to_branches(rows)
        _cm = {b["ten"]: b.get("color") for b in r["branches"] if b.get("color")}
        for _b in branches:
            if not _b.get("color") and _b["ten"] in _cm:
                _b["color"] = _cm[_b["ten"]]
        if not branches:
            st.warning("Sơ đồ đang trống, hãy nhập ít nhất một nhánh.")
        else:
            L = mindmap_layout(title, branches)
            svg = svg_from_layout(L)
            mm_component(svg, L["W"], L["H"])
            st.markdown("**Tải về để chỉnh sửa tiếp trên máy**")
            c1, c2, c3 = st.columns(3)
            pp = mm_pptx_bytes(L)
            if pp:
                c1.download_button("📊 PowerPoint (.pptx)", pp, file_name="so_do_tu_duy.pptx", key="mm_pptx",
                                   mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", width="stretch")
            else:
                c1.caption("Thiếu thư viện python-pptx nên chưa xuất được PowerPoint.")
            c2.download_button("🧩 draw.io (.drawio)", mm_drawio(L).encode("utf-8"), file_name="so_do_tu_duy.drawio", key="mm_dio", width="stretch")
            c3.download_button("🗺️ XMind / FreeMind (.mm)", mm_freemind(L).encode("utf-8"), file_name="so_do_tu_duy.mm", key="mm_fm", width="stretch")
            d1, d2 = st.columns(2)
            d1.download_button("🖼️ SVG (vector, in sắc nét)", svg.encode("utf-8"), file_name="so_do_tu_duy.svg", mime="image/svg+xml", key="mm_svg", width="stretch")
            d2.download_button("📝 Dàn ý (Markdown)", mm_outline(title, branches).encode("utf-8"), file_name="so_do_tu_duy.md", key="mm_md", width="stretch")
            st.caption("PowerPoint: mỗi khung là hình có chữ sửa được và đường nối dính vào khung, kéo khung thì đường nối đi theo. "
                       "draw.io mở bằng app.diagrams.net. File .mm mở bằng XMind, Freeplane, MindManager.")
        if r.get("smiles"):
            with st.expander("Cấu trúc chất (kiểm tra xem đúng chất bạn muốn chưa)"):
                m = Chem.MolFromSmiles(r["smiles"])
                if m is not None:
                    st.image(Draw.MolToImage(m, size=(300, 200)))
                st.code(r["smiles"], language=None)


def mindmap_ui():
    top = st.container()
    st.session_state.setdefault("mm", None)
    mode = st.radio("Cách tạo sơ đồ", ["Từ chủ đề", "Từ ảnh sơ đồ tôi vẽ tay"], horizontal=True, key="mm_mode")
    if mode.startswith("Từ ảnh"):
        _mm_photo_inputs(top)
    else:
        _mm_topic_inputs(top)
    _mm_result()


# ---------------- Giao diện ----------------
# ====================================================================
#  GIAO DIỆN MolSnap (thay thế lớp giao diện cũ, giữ nguyên toàn bộ logic phía trên)
# ====================================================================
def _uri(svg):
    return "data:image/svg+xml;base64," + base64.b64encode(svg.encode("utf-8")).decode()


# ---------- Bộ biểu tượng nét (24x24) ----------
ICON_PATHS = {
    "home": "<path d='M3 11l9-8 9 8'/><path d='M5 10v10h5v-6h4v6h5V10'/>",
    "structure": "<circle cx='6' cy='6' r='2'/><circle cx='18' cy='6' r='2'/><circle cx='12' cy='12' r='2'/><circle cx='6' cy='18' r='2'/><circle cx='18' cy='18' r='2'/><path d='M7.5 7.5l3 3M16.5 7.5l-3 3M7.5 16.5l3-3M16.5 16.5l-3-3'/>",
    "flask": "<path d='M9 3h6M10 3v6l-5 9a2 2 0 0 0 2 3h10a2 2 0 0 0 2-3l-5-9V3'/><path d='M7.5 15h9'/>",
    "isomer": "<path d='M16 3h5v5M4 20L21 3M21 16v5h-5M15 15l6 6M4 4l5 5'/>",
    "equipment": "<path d='M7 3h10v3H7zM6 6h12v14a1 1 0 0 1-1 1H7a1 1 0 0 1-1-1z'/><path d='M9 12h6M9 16h6'/>",
    "mind": "<circle cx='18' cy='5' r='3'/><circle cx='6' cy='12' r='3'/><circle cx='18' cy='19' r='3'/><path d='M8.6 13.5l6.8 4M15.4 6.5l-6.8 4'/>",
    "practice": "<rect x='5' y='4' width='14' height='17' rx='2'/><path d='M9 4h6v3H9zM9 14l2 2 4-4'/>",
    "stats": "<path d='M5 21V11M12 21V4M19 21v-7'/>",
    "chat": "<path d='M21 12a8 8 0 0 1-11.5 7.2L4 21l1.8-5.3A8 8 0 1 1 21 12z'/>",
    "pen": "<path d='M12 20h9'/><path d='M16.5 3.5a2.1 2.1 0 0 1 3 3L7 19l-4 1 1-4z'/>",
    "search": "<circle cx='11' cy='11' r='7'/><path d='M21 21l-4.3-4.3'/>",
    "sun": "<circle cx='12' cy='12' r='4'/><path d='M12 2v2M12 20v2M4.9 4.9l1.4 1.4M17.7 17.7l1.4 1.4M2 12h2M20 12h2M4.9 19.1l1.4-1.4M17.7 6.3l1.4-1.4'/>",
    "bell": "<path d='M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9'/><path d='M10.3 21a1.9 1.9 0 0 0 3.4 0'/>",
    "image": "<rect x='3' y='3' width='18' height='18' rx='3'/><circle cx='9' cy='9' r='2'/><path d='M21 15l-5-5L5 21'/>",
    "arrow": "<path d='M5 12h14M13 6l6 6-6 6'/>",
    "camera": "<path d='M4 8h3l2-3h6l2 3h3v11H4z'/><circle cx='12' cy='13' r='3.5'/>",
}


def icon_svg(name, color="currentColor", size=22, sw=1.8):
    return (f"<svg xmlns='http://www.w3.org/2000/svg' width='{size}' height='{size}' viewBox='0 0 24 24' fill='none' stroke='{color}' "
            f"stroke-width='{sw}' stroke-linecap='round' stroke-linejoin='round'>{ICON_PATHS[name]}</svg>")


def _mask(name):
    return f'url("{_uri(icon_svg(name, "black"))}")'


# ---------- Hình minh họa ----------
LOGO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><defs>
<linearGradient id="lg" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#8b5cf6"/><stop offset=".55" stop-color="#4f6df5"/><stop offset="1" stop-color="#22d3ee"/></linearGradient></defs>
<path d="M32 5 L56 19 V45 L32 59 L8 45 V19 Z" fill="url(#lg)" stroke="url(#lg)" stroke-width="6" stroke-linejoin="round"/>
<path d="M42 23 C37 16 23 18 23 26 C23 34 41 31 41 40 C41 48 28 48 22 41" fill="none" stroke="#fff" stroke-width="5.5" stroke-linecap="round"/>
<circle cx="45" cy="21" r="3" fill="#fff" opacity=".9"/><circle cx="19" cy="43" r="3" fill="#fff" opacity=".9"/></svg>"""

HERO_ART = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 560 290"><defs>
<linearGradient id="liq" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#67e8f9"/><stop offset=".55" stop-color="#3b82f6"/><stop offset="1" stop-color="#4f46e5"/></linearGradient>
<linearGradient id="glass" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#fff" stop-opacity=".55"/><stop offset="1" stop-color="#bfdbfe" stop-opacity=".25"/></linearGradient>
<linearGradient id="leaf" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#86efac"/><stop offset="1" stop-color="#16a34a"/></linearGradient>
<style>.f1{animation:f 5s ease-in-out infinite}.f2{animation:f 6.5s ease-in-out 1s infinite}.f3{animation:f 7.5s ease-in-out .4s infinite}
@keyframes f{0%,100%{transform:translateY(0)}50%{transform:translateY(-8px)}}</style></defs>
<path d="M395 250 Q360 190 400 150 Q430 190 395 250Z" fill="url(#leaf)" opacity=".95"/>
<path d="M410 252 Q420 190 470 170 Q470 220 410 252Z" fill="#22c55e" opacity=".9"/>
<path d="M340 214 Q388 192 442 206 L442 266 Q388 252 340 272Z" fill="#ece9ff"/>
<path d="M442 206 Q500 190 552 212 L552 270 Q500 252 442 266Z" fill="#fff"/>
<path d="M442 206 V266" stroke="#c7c9f5" stroke-width="2"/>
<g stroke="#c4c8ef" stroke-width="2.5" stroke-linecap="round"><path d="M462 232 H520"/><path d="M462 244 H532"/><path d="M356 232 H420"/><path d="M356 244 H410"/></g>
<text x="466" y="226" font-family="Dancing Script,cursive" font-size="20" fill="#7a86c9" font-style="italic">H₂O</text>
<polygon points="378,226 388,220 398,226 398,238 388,244 378,238" fill="none" stroke="#7a86c9" stroke-width="2"/>
<g class="f1"><path d="M290 34 H322 V98 H290Z" fill="url(#glass)" stroke="#fff" stroke-opacity=".8" stroke-width="2"/>
<ellipse cx="306" cy="32" rx="24" ry="8" fill="none" stroke="#a5f3fc" stroke-width="6"/>
<path d="M290 98 H322 L372 196 Q384 230 352 232 H260 Q228 230 240 196Z" fill="url(#glass)" stroke="#fff" stroke-opacity=".85" stroke-width="2"/>
<path d="M254 170 H358 L372 196 Q384 230 352 232 H260 Q228 230 240 196Z" fill="url(#liq)"/>
<path d="M254 170 Q306 160 358 170" fill="none" stroke="#e0f2fe" stroke-width="3" opacity=".8"/>
<circle cx="276" cy="198" r="7" fill="#e0f2fe" opacity=".75"/><circle cx="312" cy="190" r="5" fill="#e0f2fe" opacity=".75"/><circle cx="332" cy="210" r="6" fill="#e0f2fe" opacity=".75"/><circle cx="300" cy="218" r="4" fill="#e0f2fe" opacity=".7"/>
<path d="M262 120 Q250 170 252 196" fill="none" stroke="#fff" stroke-width="5" stroke-linecap="round" opacity=".55" transform="translate(30 0)"/></g>
<g class="f2"><g stroke="#cfd8ff" stroke-width="5" stroke-linecap="round"><line x1="70" y1="132" x2="128" y2="88"/><line x1="128" y1="88" x2="190" y2="134"/><line x1="190" y1="134" x2="152" y2="186"/><line x1="128" y1="88" x2="214" y2="66"/><line x1="70" y1="132" x2="98" y2="186"/></g>
<circle cx="70" cy="132" r="15" fill="#4f6df5"/><circle cx="128" cy="88" r="19" fill="#8b5cf6"/><circle cx="190" cy="134" r="15" fill="#22d3ee"/><circle cx="152" cy="186" r="17" fill="#3b82f6"/><circle cx="214" cy="66" r="11" fill="#f59e0b"/><circle cx="98" cy="186" r="9" fill="#f97316"/>
<g fill="#fff" opacity=".6"><circle cx="65" cy="126" r="5"/><circle cx="122" cy="80" r="6"/><circle cx="185" cy="129" r="5"/><circle cx="146" cy="179" r="5"/><circle cx="211" cy="63" r="3.5"/></g></g>
<g class="f3"><ellipse cx="232" cy="38" rx="30" ry="12" fill="none" stroke="#7dd3fc" stroke-width="7" opacity=".85" transform="rotate(-14 232 38)"/>
<ellipse cx="452" cy="96" rx="20" ry="8" fill="none" stroke="#c4b5fd" stroke-width="6" opacity=".8" transform="rotate(12 452 96)"/></g>
<g fill="#fff" opacity=".9"><path d="M40 52 l3 9 9 3 -9 3 -3 9 -3-9 -9-3 9-3z"/><path d="M500 40 l2 6 6 2 -6 2 -2 6 -2-6 -6-2 6-2z"/><path d="M360 262 l2 5 5 2 -5 2 -2 5 -2-5 -5-2 5-2z" opacity=".6"/></g></svg>"""

MASCOT_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 250"><defs>
<linearGradient id="lf" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#4ade80"/><stop offset="1" stop-color="#15803d"/></linearGradient>
<linearGradient id="rb" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#fff"/><stop offset="1" stop-color="#dbe4ff"/></linearGradient>
<style>.bob{animation:b 4s ease-in-out infinite;transform-origin:150px 200px}@keyframes b{0%,100%{transform:translateY(0)}50%{transform:translateY(-5px)}}</style></defs>
<path d="M8 232 Q-6 180 40 160 Q52 205 8 232Z" fill="url(#lf)" opacity=".9"/><path d="M40 238 Q40 190 84 176 Q90 220 40 238Z" fill="#22c55e" opacity=".85"/>
<path d="M290 236 Q306 196 262 176 Q250 214 290 236Z" fill="url(#lf)" opacity=".9"/>
<rect x="50" y="206" width="200" height="26" rx="6" fill="#3b5bdb"/><rect x="50" y="206" width="200" height="8" rx="4" fill="#6c8cff"/>
<rect x="64" y="182" width="176" height="26" rx="6" fill="#7c3aed"/><rect x="64" y="182" width="176" height="8" rx="4" fill="#a78bfa"/>
<rect x="80" y="160" width="150" height="24" rx="6" fill="#0891b2"/><rect x="80" y="160" width="150" height="7" rx="3.5" fill="#67e8f9"/>
<g class="bob"><line x1="150" y1="30" x2="150" y2="48" stroke="#a5b4fc" stroke-width="4"/><circle cx="150" cy="26" r="7" fill="#22d3ee"/>
<rect x="104" y="46" width="92" height="70" rx="30" fill="url(#rb)" stroke="#c7d2fe" stroke-width="3"/>
<rect x="114" y="60" width="72" height="42" rx="20" fill="#16204f"/><ellipse cx="136" cy="80" rx="7" ry="9" fill="#38e1ff"/><ellipse cx="166" cy="80" rx="7" ry="9" fill="#38e1ff"/>
<path d="M143 92 Q151 98 159 92" fill="none" stroke="#38e1ff" stroke-width="3" stroke-linecap="round"/>
<rect x="98" y="66" width="10" height="28" rx="5" fill="#7c8cff"/><rect x="192" y="66" width="10" height="28" rx="5" fill="#7c8cff"/>
<rect x="116" y="116" width="68" height="46" rx="20" fill="url(#rb)" stroke="#c7d2fe" stroke-width="3"/>
<circle cx="150" cy="138" r="6" fill="#6d5bf0"/>
<path d="M118 126 Q96 140 104 158" fill="none" stroke="#dbe4ff" stroke-width="12" stroke-linecap="round"/>
<path d="M182 126 Q204 132 208 148" fill="none" stroke="#dbe4ff" stroke-width="12" stroke-linecap="round"/>
<g transform="translate(196 134)"><path d="M10 0 h10 v14 l12 22 q4 10 -6 10 h-22 q-10 0 -6 -10 l12 -22z" fill="#dff7ff" stroke="#fff" stroke-width="2"/><path d="M4 30 h32 l4 8 q3 8 -6 8 h-28 q-9 0 -6 -8z" fill="#4ade80"/></g></g></svg>"""

AI_FACE = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><rect x="8" y="12" width="48" height="40" rx="18" fill="#fff"/><rect x="14" y="20" width="36" height="24" rx="12" fill="#16204f"/><ellipse cx="25" cy="32" rx="4" ry="5" fill="#38e1ff"/><ellipse cx="39" cy="32" rx="4" ry="5" fill="#38e1ff"/><circle cx="32" cy="8" r="4" fill="#22d3ee"/><rect x="31" y="9" width="2" height="5" fill="#a5b4fc"/></svg>"""

PAPER_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 150 130"><defs><filter id="s" x="-10%" y="-10%" width="130%" height="130%"><feDropShadow dx="0" dy="3" stdDeviation="3" flood-color="#1b2150" flood-opacity=".25"/></filter></defs>
<g transform="rotate(4 75 65)"><rect x="18" y="14" width="112" height="100" rx="3" fill="#dcd5c8" filter="url(#s)"/></g>
<g transform="rotate(-3 75 65)"><rect x="14" y="10" width="112" height="100" rx="3" fill="#f6f3ec" filter="url(#s)"/>
<g fill="none" stroke="#5b6075" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round">
<polygon points="52,58 64,51 76,58 76,72 64,79 52,72"/><polyline points="55,60 55,70" stroke-width="1.6"/><polyline points="66,77 74,72" stroke-width="1.6" transform="translate(0 0)"/><polyline points="58,54 64,51 70,54" stroke-width="1.6"/>
<line x1="76" y1="58" x2="88" y2="51"/><line x1="88" y1="51" x2="88" y2="38"/><line x1="91" y1="51" x2="91" y2="38"/><line x1="88" y1="51" x2="100" y2="58"/></g>
<text x="86" y="33" font-family="Dancing Script,cursive" font-size="13" fill="#5b6075">O</text><text x="100" y="66" font-family="Dancing Script,cursive" font-size="13" fill="#5b6075">OH</text></g></svg>"""


# ---------- Tài nguyên giao diện ----------
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&family=Dancing+Script:wght@600;700&display=swap');
:root{--pri:#4f5cf0;--vio:#7c4dff;--cy:#19c3ff;--ink:#1b2150;--mut:#6b7299;--line:#e3e7fb;--soft:#f3f5ff;--navy:#0a0e30;
--grad:linear-gradient(135deg,#3b6cf6 0%,#6d4df0 100%)}
html{font-size:__ROOT__!important}
html,body,.stApp,.stMarkdown,p,label,input,textarea,button,li,h1,h2,h3,h4{font-family:'Be Vietnam Pro',system-ui,sans-serif}
.stApp{background:radial-gradient(64.286rem 30rem at 78% -8%,#e3e8ff 0%,rgba(227,232,255,0) 70%),#f5f7ff}
[data-testid="stHeader"]{background:transparent}
[data-testid="stAppDeployButton"]{display:none}
.block-container{max-width:95.714rem;padding:.6rem 1.6rem .8rem}
.block-container [data-testid="stVerticalBlock"]{gap:.7rem}
[data-testid="stSidebarUserContent"] [data-testid="stVerticalBlock"]{gap:.35rem}

/* ---------- Sidebar ---------- */
[data-testid="stSidebar"]{background:linear-gradient(185deg,#080c2c 0%,#0e1450 52%,#0a0f3a 100%);border-right:0}
[data-testid="stSidebarContent"]{background:transparent}
[data-testid="stSidebarUserContent"]{padding:0 .9rem .5rem}
[data-testid="stSidebar"] *{color:#e6eaff}
[data-testid="stSidebar"] [data-testid="stSidebarCollapseButton"] *{color:#c3cbff}
.sb-brand{display:flex;align-items:center;gap:0.857rem;margin:0 0 .6rem}
.sb-brand img{width:3rem;height:3rem;flex:0 0 auto;filter:drop-shadow(0 0.429rem 1rem rgba(110,90,255,.55))}
.sb-brand b{display:block;font-size:1.7rem;font-weight:800;letter-spacing:-.5px;line-height:1.05;color:#fff}
.sb-brand b i{font-style:normal;background:linear-gradient(90deg,#6cb8ff,#a78bfa);-webkit-background-clip:text;background-clip:text;color:transparent}
.sb-brand span{display:block;white-space:nowrap;font-style:italic;font-size:.72rem;color:#aab4ea;margin-top:0.214rem}
.st-key-sidenav [role="radiogroup"]{display:flex;flex-direction:column;gap:0.214rem;width:100%}
.st-key-sidenav [role="radiogroup"]>div{width:100%}
.st-key-sidenav [data-testid="stRadioOption"]>div{flex:1;min-width:0}
.st-key-sidenav [data-testid="stRadioOption"] p{white-space:nowrap}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]){display:flex;align-items:center;gap:0.857rem;margin:0;width:100%;padding:0.5rem 1rem;border-radius:0.857rem;cursor:pointer;border:1px solid transparent;transition:background .15s,transform .15s}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"])>div:first-child{display:none}
.st-key-sidenav [data-testid="stRadioOption"]>div{display:flex;align-items:center;gap:0;width:auto;padding:0}
.st-key-sidenav [data-testid="stRadioOption"]>div>div:first-child{display:none}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"])::before{content:"";flex:0 0 1.571rem;width:1.571rem;height:1.571rem;background-color:#dbe1ff;-webkit-mask:var(--ic) center/contain no-repeat;mask:var(--ic) center/contain no-repeat}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]) p{font-size:1rem;font-weight:500;color:#e6eaff}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]):hover{background:rgba(255,255,255,.08);transform:translateX(2px)}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]){background:linear-gradient(90deg,#3f63ee,#5b86ff);border-color:rgba(255,255,255,.28);box-shadow:0 0.571rem 1.571rem rgba(63,99,238,.45)}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"])::before{background-color:#fff}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]) p{color:#fff;font-weight:700}
__NAVICONS__
.aicard{display:flex;align-items:center;gap:0.857rem;padding:0.643rem 0.857rem;border-radius:1.143rem;background:linear-gradient(135deg,rgba(63,99,238,.42),rgba(109,77,240,.32));border:1px solid rgba(130,160,255,.5);box-shadow:0 0.714rem 1.857rem rgba(10,14,60,.45);margin-top:0.429rem}
.aicard img{width:2.714rem;height:2.714rem;flex:0 0 auto}
.aicard b{display:block;color:#fff;font-size:1rem}.aicard span{display:block;font-size:.78rem;color:#c9d2ff;line-height:1.35;margin-top:2px}
.aicard .go{margin-left:auto;flex:0 0 auto;width:2.143rem;height:2.143rem;border-radius:50%;background:rgba(255,255,255,.16);display:grid;place-items:center;color:#fff;font-weight:700}
.mascot{position:relative;margin:0 auto;width:72%}.mascot img{width:100%;display:block}
.mascot .hw{position:absolute;right:0.429rem;top:1rem;font-family:'Dancing Script',cursive;font-size:1.35rem;line-height:1.1;transform:rotate(-8deg);color:#e8ecff;text-align:center}

/* ---------- Ô bấm phủ lên thẻ (thẻ HTML + nút ẩn) ---------- */
div[class*="st-key-qk_"],.st-key-aibtn{position:relative;gap:0}
div[class*="st-key-qk_"] [data-testid="stElementContainer"]:has([data-testid="stButton"]),.st-key-aibtn [data-testid="stElementContainer"]:has([data-testid="stButton"]){position:absolute;inset:0;z-index:3;width:100%!important}
div[class*="st-key-qk_"] [data-testid="stButton"],.st-key-aibtn [data-testid="stButton"],div[class*="st-key-qk_"] [data-testid="stButton"] button,.st-key-aibtn [data-testid="stButton"] button{width:100%;height:100%;min-height:0}
div[class*="st-key-qk_"] button,.st-key-aibtn button{opacity:0;cursor:pointer}
div[class*="st-key-qk_"]:hover .tile{transform:translateY(-3px);box-shadow:0 0.857rem 1.857rem rgba(79,92,240,.2)}
.st-key-aibtn:hover .aicard{border-color:#9db4ff;transform:translateY(-1px)}.aicard{transition:all .15s}

/* ---------- Thanh trên ---------- */
.st-key-gsearch [data-testid="InputInstructions"]{display:none}
.st-key-gsearch div[data-baseweb="input"]{background:#fff;border:1px solid var(--line);border-radius:71.357rem;box-shadow:0 0.429rem 1.429rem rgba(79,92,240,.08);height:2.5rem}
.st-key-gsearch input{padding-left:3rem;background:url("__SEARCHIC__") no-repeat 1.1rem center/1.429rem;color:var(--ink)}
.tb-right{display:flex;justify-content:flex-end;align-items:center;gap:1.143rem;color:#5b6293}
.tb-right .av{width:2.571rem;height:2.571rem;border-radius:50%;background:var(--grad);color:#fff;display:grid;place-items:center;font-weight:700;box-shadow:0 0.286rem 0.857rem rgba(79,92,240,.4)}

/* ---------- Hero ---------- */
.hero3{position:relative;display:flex;justify-content:space-between;align-items:center;gap:0.714rem;color:#fff;border-radius:1.714rem;padding:1.143rem 2rem;margin:0 0 0.714rem;overflow:hidden;
background:radial-gradient(30rem 18.571rem at 88% 18%,rgba(25,195,255,.38),rgba(25,195,255,0) 70%),radial-gradient(27.143rem 18.571rem at 40% 110%,rgba(139,92,246,.5),rgba(139,92,246,0) 70%),linear-gradient(115deg,#152a9a 0%,#2f35c8 48%,#5a3bea 100%);
box-shadow:0 1.286rem 3.143rem rgba(47,53,200,.32)}
.hero3 .pill{display:inline-block;font-size:.85rem;background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.34);border-radius:71.357rem;padding:0.357rem 1.143rem;font-weight:600;font-size:.95rem;margin-bottom:0.571rem}
.hero3 h1{margin:0;padding:0;font-size:2.05rem;line-height:1.16;font-weight:800;color:#fff;letter-spacing:-.4px}
.hero3 h1 em{font-style:normal;font-size:2.6rem;background:linear-gradient(90deg,#22d3ee,#60a5fa 60%,#a5b4fc);-webkit-background-clip:text;background-clip:text;color:transparent}
.hero3 .sub{margin:0.429rem 0 0.714rem;font-size:1rem;font-weight:600;display:flex;gap:1rem;align-items:center;flex-wrap:wrap}.hero3 .sub i{opacity:.7;font-style:normal}
.hero3 .chips{display:flex;flex-wrap:wrap;gap:0.714rem}
.hero3 .chips span{background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.3);border-radius:71.357rem;padding:0.429rem 0.929rem;font-size:.85rem;font-weight:600;backdrop-filter:blur(0.429rem)}
.hero3 .hr{position:relative;flex:0 0 31%}.hero3 .art{width:100%;display:block}
.hero3 .slogan{position:absolute;right:0.714rem;top:0;font-family:'Dancing Script',cursive;font-weight:700;font-size:1.45rem;line-height:1.15;transform:rotate(-7deg);text-align:right;opacity:.96}
.hero3 .slogan u{display:block;height:2px;margin-top:0.429rem;background:linear-gradient(90deg,rgba(255,255,255,0),#fff);border-radius:2px}

/* ---------- Banner trang con ---------- */
.pbanner{display:flex;align-items:center;gap:1rem;color:#fff;border-radius:1.429rem;padding:1rem 1.429rem;margin:0.429rem 0 1.143rem;
background:radial-gradient(21.429rem 11.429rem at 92% 0%,rgba(25,195,255,.35),rgba(25,195,255,0) 70%),linear-gradient(115deg,#152a9a 0%,#3a37cf 55%,#5a3bea 100%);box-shadow:0 0.857rem 2.143rem rgba(47,53,200,.25);animation:fu .45s ease}
.pbanner .pb-ic{width:3.857rem;height:3.857rem;border-radius:1.143rem;background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.3);display:grid;place-items:center;flex:0 0 auto}
.pbanner b{display:block;font-size:1.35rem;font-weight:800;color:#fff}.pbanner span{display:block;color:#dbe3ff;font-size:.95rem;margin-top:2px}
@keyframes fu{from{opacity:0;transform:translateY(0.571rem)}to{opacity:1;transform:none}}
.st-key-panel{background:#fff;border:1px solid var(--line);border-radius:1.714rem;padding:1.571rem 1.857rem;box-shadow:0 0.714rem 2.143rem rgba(79,92,240,.08)}
.st-key-panel [role="radiogroup"]{gap:0.571rem;flex-wrap:wrap}
.st-key-panel :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]){background:#fff;border:0.107rem solid var(--line);border-radius:1rem;padding:0.571rem 1.143rem;margin:0}
.st-key-panel :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]){border-color:var(--pri);background:var(--soft)}

/* ---------- Thẻ tải ảnh (trang chủ) ---------- */
.st-key-upcard{background:linear-gradient(135deg,#e7f8f6 0%,#eef1ff 60%,#f1ecff 100%);border:1px solid #dfe5ff;border-radius:1.571rem;padding:0.857rem 1.286rem;box-shadow:0 0.857rem 2.286rem rgba(79,92,240,.1)}
.uphead{display:flex;align-items:center;gap:1.143rem;margin-bottom:0.571rem}
.uphead .hex{flex:0 0 3.714rem;height:3.714rem;display:grid;place-items:center;background:linear-gradient(135deg,#22c1a5,#3b82f6);clip-path:polygon(50% 0,93% 25%,93% 75%,50% 100%,7% 75%,7% 25%);filter:drop-shadow(0 0.571rem 1rem rgba(34,193,165,.4))}
.uphead h3{margin:0 0 0.286rem;padding:0;font-size:1.25rem;font-weight:800;color:var(--ink)}
.uphead p{margin:0;color:#59618f;font-size:.98rem;line-height:1.5}
.uphead .hw{margin-left:auto;font-family:'Dancing Script',cursive;font-size:1.25rem;color:#3a4690;transform:rotate(-5deg);text-align:right;line-height:1.1;flex:0 0 auto}
.demo{position:relative;display:flex;align-items:center;gap:0.714rem;justify-content:flex-end}
.demo .paper{width:clamp(5rem,8vw,8.571rem);flex:0 1 auto;min-width:0}
.demo .arrow{color:#6d4df0;flex:0 0 auto}
.demo .res{position:relative;flex:1;min-width:0;background:#fff;border-radius:1rem;padding:0.714rem 1rem;box-shadow:0 0.714rem 1.857rem rgba(79,92,240,.14);border:1px solid #e6e9ff;min-width:0}
.demo .res .badge{position:absolute;top:-13px;right:0.714rem;background:#12b886;color:#fff;font-size:.72rem;font-weight:700;border-radius:71.357rem;padding:0.286rem 0.857rem;box-shadow:0 0.429rem 1rem rgba(18,184,134,.35)}
.demo .res h4{margin:0;padding:0;font-size:1.05rem;color:var(--ink)}.demo .res h4 i{float:right;font-style:normal;color:#12b886}
.demo .res b{display:block;color:var(--ink);font-size:.9rem;margin:2px 0 0.571rem}
.demo .res div{font-size:.78rem;color:#4b5380;margin-top:0.286rem;line-height:1.35}
.demo .res small{display:block;margin-top:0.571rem;font-size:.68rem;color:#98a0c7}
[data-testid="stFileUploaderDropzone"]{display:flex;flex-direction:column;align-items:center;justify-content:center;gap:0.571rem;border:2px dashed #8f9dff;background:rgba(255,255,255,.72);border-radius:1.286rem;padding:0.714rem 0.714rem;min-height:11rem}
[data-testid="stFileUploaderDropzone"]::before{content:"";order:0;width:3.286rem;height:3.286rem;border-radius:1rem;background:#e7ebff url("__IMGIC__") center/1.714rem no-repeat}
[data-testid="stFileUploaderDropzone"]>span{order:1}
[data-testid="stFileUploaderDropzone"]>span>button *{font-size:0!important}
[data-testid="stFileUploaderDropzone"]>span>button{font-size:0;background:var(--grad);color:#fff;border:0;border-radius:71.357rem;padding:0.571rem 2.143rem;box-shadow:0 0.714rem 1.571rem rgba(79,92,240,.35)}
[data-testid="stFileUploaderDropzone"]>span>button *{color:#fff}
[data-testid="stFileUploaderDropzone"]>span>button p{display:none}
[data-testid="stFileUploaderDropzone"]>span>button::after{content:"Tải ảnh lên";font-size:1rem;font-weight:700;color:#fff}
[data-testid="stFileUploaderDropzone"]>span>button [data-testid="stIconMaterial"]{display:none}
[data-testid="stFileUploaderDropzoneInstructions"]{order:2;text-align:center}
[data-testid="stFileUploaderDropzoneInstructions"] *{font-size:0}
[data-testid="stFileUploaderDropzoneInstructions"] svg,[data-testid="stFileUploaderDropzoneInstructions"] [data-testid="stIconMaterial"]{display:none}
[data-testid="stFileUploaderDropzoneInstructions"]::before{content:"Kéo thả ảnh hoặc chọn file  |  JPG, PNG, WEBP (tối đa 200MB)";font-size:.78rem;color:#6b7299;line-height:1.4}

/* ---------- Khám phá nhanh ---------- */
.sec-title{display:flex;align-items:center;gap:0.714rem;font-weight:700;font-size:1.02rem;color:var(--ink);margin:0.429rem 0.286rem 0}
.sec-title svg{color:var(--pri)}
.tile{display:flex;align-items:center;justify-content:space-between;gap:0.429rem;border-radius:1.286rem;padding:1rem 0.857rem;height:7.4rem;box-sizing:border-box;border:1px solid rgba(255,255,255,.7);transition:all .18s ease}
.tile .tt{display:flex;flex-direction:column;gap:0.429rem}
.tile .ib{width:2rem;height:2rem;border-radius:50%;background:rgba(255,255,255,.85);display:grid;place-items:center}
.tile b{font-size:.88rem;font-weight:700;line-height:1.25;white-space:pre-line}
.tile .ar{align-self:flex-end;opacity:.85}
.t-pink{background:linear-gradient(135deg,#ffe4f5,#fbd3ee);color:#d6389b}
.t-blue{background:linear-gradient(135deg,#e3eeff,#d4e4ff);color:#2563eb}
.t-green{background:linear-gradient(135deg,#dcf7ea,#c8f1dd);color:#059669}
.t-orange{background:linear-gradient(135deg,#fff1dc,#ffe4bd);color:#d97706}
.t-violet{background:linear-gradient(135deg,#ece8ff,#ddd6ff);color:#6d4df0}

/* ---------- Hành trình + dải tính năng ---------- */
.st-key-journey{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.571rem 1.286rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.jr{display:flex;align-items:center;gap:1.143rem}.jr .tr{font-size:1.9rem;line-height:1}
.jr .tx{flex:1}.jr b{display:block;color:var(--ink);font-size:1.02rem}.jr small{color:var(--mut)}
.jr .bar{flex:1.2;height:0.643rem;border-radius:7.071rem;background:#e8ebff;overflow:hidden;min-width:8.571rem}.jr .bar i{display:block;height:100%;border-radius:7.071rem;background:var(--grad)}
.jr .cnt{font-weight:700;font-size:.88rem;color:var(--ink);white-space:nowrap}
.feat{display:flex;justify-content:space-around;align-items:center;margin:2px 0.286rem 0;color:#4b5380;font-weight:600;font-size:.9rem;flex-wrap:wrap;gap:0.571rem}
.feat span{display:flex;align-items:center;gap:0.571rem}.feat i{width:1px;height:1.286rem;background:#d9defa;display:inline-block}

/* ---------- Cột phải ---------- */
.st-key-quick{background:linear-gradient(135deg,#3b6cf6 0%,#7a52f4 100%);border-radius:1.429rem;padding:0.857rem 1rem 0.714rem;box-shadow:0 1rem 2.143rem rgba(79,92,240,.35)}
.qhead{display:flex;align-items:center;gap:0.857rem;color:#fff;margin-bottom:0.714rem}
.qhead .qi{width:2.714rem;height:2.714rem;border-radius:1rem;background:rgba(255,255,255,.2);display:grid;place-items:center;flex:0 0 auto}
.qhead b{display:block;font-size:1.1rem;color:#fff}.qhead span{font-size:.8rem;color:#e1e6ff}
.st-key-quick [data-testid="stForm"]{border:0;padding:0;background:transparent}
.st-key-quick [data-testid="InputInstructions"]{display:none}
.st-key-quick div[data-baseweb="input"]{background:#fff;border:0;border-radius:71.357rem;height:2.9rem}
.st-key-quick input{color:var(--ink);padding-left:1.1rem}
.st-key-quick [data-testid="stFormSubmitButton"] button{width:2.7rem;height:2.7rem;min-height:0;padding:0;border-radius:50%;background:var(--grad);border:2px solid #fff;color:#fff;font-size:1.15rem;line-height:1}
.st-key-quick [data-testid="stFormSubmitButton"] button p{color:#fff;font-size:1.2rem}
.rcard{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.714rem 1rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.rcard .rh{display:flex;align-items:center;gap:0.714rem;margin-bottom:0.286rem}.rcard .rh svg{color:var(--pri)}
.rcard .rh b{color:var(--ink);font-size:1rem}.rcard .rs{color:var(--mut);font-size:.78rem;margin:0 0 0.429rem 2.429rem}
.st-key-res{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.714rem 0.857rem 0.286rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.st-key-res{gap:0!important}
.st-key-res [data-testid="stElementContainer"],.st-key-res [data-testid="stButton"]{width:100%}
.st-key-res button>div{width:100%;justify-content:flex-start}
.st-key-res button::after{content:"›";margin-left:auto;font-size:1.2rem;color:#9aa3d6}
.st-key-res button{width:100%;justify-content:flex-start;gap:0.714rem;background:transparent;border:0;border-top:1px solid #eef0fb;border-radius:0;padding:.15rem .4rem;min-height:2.1rem;color:var(--ink);box-shadow:none}
.st-key-res button:hover{background:var(--soft);transform:none;box-shadow:none}
.st-key-res button p{font-weight:500;font-size:.9rem;color:var(--ink)}
.st-key-res button [data-testid="stIconMaterial"]{color:var(--pri)}
.hrow{display:flex;align-items:center;gap:0.857rem;padding:0.357rem 2px;border-top:1px solid #eef0fb}
.hrow:first-of-type{border-top:0}
.hrow .hc{width:2.286rem;height:2.286rem;border-radius:50%;display:grid;place-items:center;flex:0 0 auto}
.hrow b{display:block;color:var(--ink);font-size:.88rem}.hrow small{color:var(--mut);font-size:.72rem}
.hrow .hc.c0{background:#d9f7ee;color:#0f9f78}.hrow .hc.c1{background:#ffe0ef;color:#d6389b}.hrow .hc.c2{background:#ffeccd;color:#c2740a}.hrow .hc.c3{background:#e6e0ff;color:#6d4df0}
.rsh{display:flex;align-items:center;gap:0.857rem;color:var(--pri);padding:2px 0.286rem 0.857rem}
.rsh b{display:block;color:var(--ink);font-size:1rem;line-height:1.3}.rsh span{display:block;color:var(--mut);font-size:.78rem;line-height:1.3;margin-top:2px}
.hempty{color:var(--mut);font-size:.85rem;padding:0.571rem 2px}
.st-key-hclear button{border:0;background:transparent;color:var(--pri);font-size:.78rem;padding:0;min-height:0;box-shadow:none}
.quote{display:flex;align-items:center;gap:1rem;border-radius:1.286rem;padding:0.714rem 1rem;margin-top:0;background:linear-gradient(135deg,#f1ecff,#e6f6ff);border:1px solid #e3e6ff}
.quote .lf{font-size:2.2rem;line-height:1}.quote p{margin:0;font-style:italic;color:#3a4690;font-size:.85rem;line-height:1.5}.quote em{display:block;text-align:right;color:var(--pri);font-size:.8rem;font-weight:600;margin-top:0.286rem}

/* ---------- Hộp thoại tài liệu ---------- */
.ptwrap{overflow-x:auto;padding-bottom:0.429rem}
.ptgrid{display:grid;grid-template-columns:repeat(18,minmax(3.857rem,1fr));grid-template-rows:repeat(7,auto) 0.857rem auto auto;gap:0.214rem;min-width:71.429rem}
.ptc{border-radius:0.5rem;padding:0.214rem 0.286rem;text-align:center;line-height:1.15;display:flex;flex-direction:column;align-items:center;min-height:4.429rem;border:1px solid rgba(27,33,80,.08)}
.ptc i{font-style:normal;font-size:.62rem;align-self:flex-start;color:#4b5380}.ptc b{font-size:1.02rem;color:var(--ink)}
.ptc span{font-size:.54rem;color:#3a4172;white-space:nowrap}.ptc em{font-style:normal;font-size:.56rem;color:#4b5380;margin-top:auto}
.ptx{background:#eef0fb;justify-content:center}.ptx b{font-size:.7rem;color:#6b7299}
.ptleg{display:flex;flex-wrap:wrap;gap:0.429rem 1.143rem;margin:0.714rem 0 0.286rem;font-size:.78rem;color:#4b5380}.ptleg span{display:flex;align-items:center;gap:0.429rem}.ptleg i{width:1rem;height:1rem;border-radius:0.286rem;display:inline-block;border:1px solid rgba(27,33,80,.12)}
.kl-kiem{background:#ffd6d6}.kl-kt{background:#ffe3c2}.chuyen{background:#fff0a8}.sau-chuyen{background:#d8f0c8}.a-kim-loai{background:#c7eee3}.phi-kim{background:#c9e6ff}.halogen{background:#d9d4ff}.khi-hiem{background:#f1d2f5}.lantan{background:#ffd0e6}.actini{background:#f4c7d9}
.reftbl{overflow-x:auto}.reftbl table{width:100%;border-collapse:collapse;font-size:.88rem}
.reftbl th{background:#eef1ff;color:var(--ink);text-align:left;padding:0.571rem 0.714rem;position:sticky;top:0}
.reftbl td{padding:0.571rem 0.714rem;border-top:1px solid #eef0fb;color:#2c3466;vertical-align:top}
.rxr{display:flex;gap:0.857rem;align-items:center;padding:0.643rem 2px;border-top:1px solid #eef0fb}.rxr:first-child{border-top:0}
.rxt{flex:0 0 10.714rem;background:#eef1ff;color:#4f5cf0;border-radius:71.357rem;padding:0.214rem 0.714rem;font-size:.74rem;font-weight:700;text-align:center}
.rxe{display:flex;align-items:center;flex-wrap:wrap;gap:0.429rem;color:var(--ink);font-weight:600;font-size:.95rem}
.rxa{display:inline-flex;flex-direction:column;align-items:center;min-width:3.286rem;line-height:1}.rxa small{font-size:.62rem;color:#6b7299;font-weight:500;white-space:nowrap}.rxa i{font-style:normal;font-size:1.2rem;color:#4f5cf0}

/* ---------- Bảng tính tan ---------- */
.soltbl{overflow-x:auto}.soltbl table{border-collapse:separate;border-spacing:2px;font-size:.78rem;width:100%}
.soltbl th{background:#eef1ff;color:var(--ink);padding:.3rem .4rem;font-weight:700;white-space:nowrap}
.soltbl td{text-align:center;font-weight:700;padding:.32rem .2rem;border-radius:.3rem}
.sol-t{background:#d7f5e5;color:#087a4a}.sol-i{background:#fff0c2;color:#a96800}.sol-k{background:#ffd9d9;color:#c2272d}.sol-b{background:#d9e8ff;color:#1d5fd1}.sol-x{background:#eceef6;color:#6b7299}.sol-q{background:#f5f6fb;color:#9aa3c4}
.sol-sel{outline:3px solid #4f5cf0;outline-offset:-1px;box-shadow:0 0 0 4px rgba(79,92,240,.25)}
.solres{display:flex;align-items:center;flex-wrap:wrap;gap:.8rem;margin:.4rem 0 .7rem}.solres b{font-size:1.5rem;color:var(--ink)}
.solbadge{padding:.3rem .85rem;border-radius:999px;font-weight:700;font-size:.85rem}
.sollg{display:flex;flex-wrap:wrap;gap:.4rem .9rem;margin-top:.6rem;font-size:.78rem;color:#4b5380;align-items:center}.sollg span{padding:.1rem .5rem;border-radius:.3rem;font-weight:700}

/* ---------- Thành phần chung ---------- */
[data-testid="stWidgetLabel"] p{font-weight:700;color:var(--ink);border-left:0.286rem solid var(--pri);padding-left:0.643rem}
.st-key-gsearch [data-testid="stWidgetLabel"],.st-key-quick [data-testid="stWidgetLabel"]{display:none}
button[kind="primary"],[data-testid="stBaseButton-primary"],[data-testid="stBaseButton-primary"]{background:var(--grad);border:0;color:#fff;border-radius:0.857rem;font-weight:700;padding:.55rem 1.3rem}
[data-testid="stBaseButton-secondary"]{border-radius:0.857rem;border-color:var(--line)}
button{transition:transform .15s ease,box-shadow .15s ease}
button:hover{transform:translateY(-1px);box-shadow:0 0.429rem 1.143rem rgba(79,92,240,.22)}
.st-key-jbtn button{border-radius:71.357rem;background:var(--grad);color:#fff;border:0;padding:.6rem 1.4rem;font-weight:700;white-space:nowrap}
.st-key-jbtn button p{color:#fff}
div[data-baseweb="input"],div[data-baseweb="textarea"]{border-radius:0.857rem}
.errtop{color:#d00000;font-weight:800;font-size:1.08rem;line-height:1.5;background:rgba(220,0,0,.08);border:2px solid rgba(220,0,0,.45);border-radius:0.857rem;padding:0.857rem 1.143rem;margin:0.429rem 0 0.857rem}
button:disabled,[data-testid="stBaseButton-primary"]:disabled{background:#dfe3f7!important;color:#8a90b8!important;box-shadow:none!important;transform:none!important}
button:disabled p{color:#8a90b8!important}
.usage{display:flex;align-items:center;gap:0.714rem;background:#fff;border:1px solid var(--line);border-radius:1rem;padding:0.429rem 1rem;color:var(--mut);margin-top:0.429rem;font-size:.85rem}
.usage b{color:var(--pri)}
.disc{color:var(--mut);font-size:.75rem;margin:0.286rem 0.286rem 0}
@media(max-width:900px){.hero3{padding:1.429rem 1.429rem}.hero3 .hr{display:none}.hero3 h1{font-size:1.8rem}.hero3 h1 em{font-size:2.3rem}.st-key-panel{padding:1rem}.uphead .hw{display:none}.demo{flex-wrap:wrap}}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
@media(max-width:640px){
html{font-size:14.5px!important}
.block-container{padding:2.8rem .7rem 1rem!important}
.hero3{padding:1rem 1.1rem;border-radius:1.2rem}
.hero3 h1{font-size:1.5rem}.hero3 h1 em{font-size:1.9rem}.hero3 .sub{font-size:.9rem}
.hero3 .chips span{font-size:.78rem;padding:.35rem .7rem}
.st-key-topbar [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:nowrap!important;gap:.5rem!important}
.st-key-topbar [data-testid="stColumn"]{min-width:0!important;width:auto!important}
.st-key-topbar [data-testid="stColumn"]:first-child{flex:1 1 auto!important}
.st-key-topbar [data-testid="stColumn"]:last-child{flex:0 0 auto!important}
.st-key-topbar [data-testid="stColumn"]:last-child{display:none!important}
.st-key-topbar [data-testid="stColumn"]:first-child{flex:1 1 100%!important;width:100%!important}
.hero3 .sub{display:none}
[data-testid="stExpandSidebarButton"]{background:var(--grad)!important;border-radius:999px!important;width:auto!important;height:2.4rem!important;padding:0 .9rem!important;box-shadow:0 6px 16px rgba(79,92,240,.4)!important;display:inline-flex!important;align-items:center;gap:.35rem}
[data-testid="stExpandSidebarButton"] *{color:#fff!important;fill:#fff!important}
[data-testid="stExpandSidebarButton"]::after{content:"Menu";color:#fff;font-weight:700;font-size:.9rem}
.hero3 h1{margin-bottom:.7rem}
.st-key-tilesrow [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:wrap!important;gap:.6rem!important}
.st-key-tilesrow [data-testid="stHorizontalBlock"]>[data-testid="stColumn"]{flex:1 1 calc(50% - .6rem)!important;min-width:calc(50% - .6rem)!important}
.uphead .hw{display:none}.uphead{gap:.8rem}.uphead h3{font-size:1.05rem}.uphead p{font-size:.82rem}
.demo{flex-direction:column;align-items:stretch}.demo .paper,.demo .arrow{display:none}
.jr{flex-wrap:wrap}.jr .bar{flex:1 1 100%;order:3}
.feat{display:grid;grid-template-columns:1fr 1fr;gap:.4rem}.feat i{display:none}
.tile{height:5.2rem}
.st-key-panel{padding:.8rem .9rem}
.st-key-quick [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:nowrap!important;gap:.5rem!important}
.st-key-quick [data-testid="stColumn"]{min-width:0!important}
.st-key-quick [data-testid="stColumn"]:first-child{flex:1 1 auto!important}
.st-key-quick [data-testid="stColumn"]:last-child{flex:0 0 auto!important;width:auto!important}
.st-key-jbtn button{width:100%}
}
</style>
"""

# Biểu tượng cho từng mục điều hướng (thứ tự = thứ tự trong thanh bên)
NAV_ICON_ORDER = ["home", "structure", "flask", "isomer", "equipment", "mind"]
_nav_css = "".join(
    f'.st-key-sidenav label[data-baseweb="radio"]:nth-of-type({i + 1}){{--ic:{_mask(n)}}}.st-key-sidenav [role="radiogroup"]>div:nth-of-type({i + 1}) [data-testid="stRadioOption"]{{--ic:{_mask(n)}}}' for i, n in enumerate(NAV_ICON_ORDER)
)
_ROOT = "clamp(9px, 1.6vh, 16px)" if st.session_state.get("nav", "Trang chủ") == "Trang chủ" else "14px"
CSS = (CSS.replace("__ROOT__", _ROOT).replace("__NAVICONS__", _nav_css)
          .replace("__SEARCHIC__", _uri(icon_svg("search", "#7a82b8", 20)))
          .replace("__IMGIC__", _uri(icon_svg("image", "#4f5cf0", 30))))
st.markdown(CSS, unsafe_allow_html=True)

pw = st.secrets.get("APP_PASSWORD", "")
if pw and st.text_input("Mật khẩu lớp học", type="password") != pw:
    st.stop()

# Khi cập nhật mã, dữ liệu kết quả cũ trong phiên đang mở có thể thiếu trường mới. Xóa các kết quả cũ để tránh lỗi KeyError.
SCHEMA = 5
if st.session_state.get("_schema") != SCHEMA:
    for _k in ("mm", "iso", "iso_nm", "iso_und", "bm", "eq", "rx", "res", "pq"):
        st.session_state.pop(_k, None)
    st.session_state["_schema"] = SCHEMA
st.session_state.setdefault("n", 0)
st.session_state.setdefault("res", None)
st.session_state.setdefault("hist", [])
st.session_state.setdefault("chat", [])


def errbanner(*lines):
    body = "<br>".join(html.escape(str(x)) for x in lines if x)
    return f'<div class="errtop">❌ {body}</div>'


def show_top_error(top, msg, exc=None):
    """Hiện lỗi ở ĐẦU tab, chữ đỏ in đậm."""
    with top:
        st.markdown(errbanner(msg), unsafe_allow_html=True)
        if exc is not None:
            st.caption(f"Chi tiết kỹ thuật: {type(exc).__name__}: {str(exc)[:400]}")


def add_history(label, kind="Công thức"):
    """Ghi một lượt tra cứu vào lịch sử (tối đa 10 mục, bỏ trùng liên tiếp)."""
    h = st.session_state.hist
    label = (label or "").strip()
    if label and not (h and h[0]["q"] == label and h[0]["k"] == kind):
        h.insert(0, {"q": label, "k": kind, "t": datetime.now().strftime("%H:%M")})
        del h[10:]


# ---------------- Công thức cấu tạo viết tay ----------------
def _log_result():
    r = st.session_state.res
    smi = ((r[0].get("smiles") if r else "") or "").strip()
    if smi:
        mol, num = chem_numbers(smi)
        if num:
            add_history(num["formula"], "Công thức")


def _run_structure(top, img_bytes=None, smi=None):
    if st.session_state.n >= MAX_PER_SESSION:
        show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        return
    with st.spinner("Đang phân tích…"):
        try:
            if img_bytes:
                st.session_state.res = (analyze_image_strict(img_bytes), True)
                st.session_state.n += 1
            elif parse_student(smi)[0] is None:
                st.session_state.res = ({"smiles": smi}, False)  # SMILES sai: chấm lỗi ngay, không gọi AI
            else:
                st.session_state.res = (analyze_smiles(smi), False)
                st.session_state.n += 1
            _log_result()
        except Exception as e:
            st.session_state.res = None
            show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)


def structure_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    # Ảnh hoặc SMILES gửi từ trang chủ / ô Nhập nhanh: phân tích ngay
    pend_img = st.session_state.pop("pending_img", None)
    pend_smi = st.session_state.pop("pending_smi", None)
    if pend_img:
        _run_structure(top, img_bytes=pend_img)
    elif pend_smi:
        _run_structure(top, smi=pend_smi)

    mode = st.radio("Cách nhập", ["🖼️ Tải ảnh lên", "📷 Chụp ảnh", "⌨️ Nhập SMILES"], horizontal=True, key="st_mode")
    img_bytes, smi_in = None, None
    if mode.startswith("📷"):
        f = st.camera_input("Chụp công thức (viết rõ nét, nền sáng, mỗi ảnh một chất)", key="st_cam")
        img_bytes = f.getvalue() if f else None
    elif mode.startswith("🖼️"):
        f = st.file_uploader("Chọn ảnh công thức (viết rõ nét, nền sáng, mỗi ảnh một chất)", type=["jpg", "jpeg", "png", "webp"], key="st_up")
        img_bytes = f.getvalue() if f else None
    else:
        smi_in = st.text_input("Mã SMILES", placeholder="Ví dụ: CC(=O)O", key="st_smi").strip()

    if st.button("Phân tích", type="primary", disabled=not (img_bytes or smi_in), key="st_go"):
        _run_structure(top, img_bytes=img_bytes, smi=smi_in)

    if st.session_state.res:
        r, from_img = st.session_state.res
        smiles_r = (r.get("smiles") or "").strip()
        _, err_r = parse_student(smiles_r) if smiles_r else (None, None)
        if err_r:
            with top:
                st.markdown(errbanner("CÔNG THỨC VẼ SAI: " + err_r), unsafe_allow_html=True)
                im = draw_student(smiles_r)
                if im is not None:
                    st.markdown("**Hình máy đọc lại đúng như hình vẽ (nguyên tử sai khoanh đỏ):**")
                    st.image(im)
                if r.get("luu_y"):
                    st.caption("Ghi chú khi đọc ảnh: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))
            st.info("Vì công thức sai hóa trị nên chưa thể gọi tên hay tính khối lượng mol. Hãy sửa lại hình vẽ hoặc sửa SMILES ở ô bên dưới.")
        else:
            st.divider()
            render(r, from_img)
        st.divider()
        fix = st.text_input("Nhận dạng sai? Sửa SMILES rồi phân tích lại", value=r.get("smiles", ""), key="st_fix").strip()
        if st.button("Phân tích lại từ SMILES đã sửa", key="st_fix_go") and fix:
            if st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Đã hết số lượt của phiên này.")
            else:
                with st.spinner("Đang phân tích…"):
                    try:
                        st.session_state.res = (analyze_smiles(fix), False)
                        st.session_state.n += 1
                        _log_result()
                        st.rerun()
                    except Exception as e:
                        show_top_error(top, "Có lỗi khi phân tích lại.")


# ---------------- Trợ lý AI ----------------
ASSIST_RULES = ("Bạn là trợ lý học Hóa học thân thiện cho học sinh Việt Nam (THCS và THPT), bám sát SGK Việt Nam. "
                "Trả lời bằng tiếng Việt, ngắn gọn, rõ ràng; có phương trình, điều kiện, ví dụ khi cần. Dùng danh pháp theo SGK. "
                "Nếu không chắc chắn thì nói rõ và khuyên đối chiếu SGK. Nếu câu hỏi không liên quan đến hóa học, nhẹ nhàng đưa cuộc trò chuyện về hóa học.")


def chat_ai(history):
    contents = [types.Content(role="user" if m["role"] == "user" else "model", parts=[types.Part.from_text(text=m["text"])]) for m in history]
    models = [MODEL]
    fb = st.secrets.get("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite")
    if fb and fb != MODEL:
        models.append(fb)
    last = None
    for m in models:
        for i in range(2):
            try:
                resp = get_client().models.generate_content(
                    model=m, contents=contents, config=types.GenerateContentConfig(system_instruction=ASSIST_RULES))
                return (resp.text or "").strip() or "Mình chưa có câu trả lời, bạn hỏi lại giúp mình nhé."
            except errors.ServerError as e:
                last = e
                time.sleep(2 * (i + 1))
            except errors.ClientError as e:
                last = e
                if getattr(e, "code", None) == 429:
                    time.sleep(3 * (i + 1))
                else:
                    break
    raise last


def assistant_ui():
    chat = st.session_state.chat
    box = st.container()  # tin nhắn nằm trên, ô nhập nằm dưới
    q = st.session_state.pop("pending_chat", None) or st.chat_input("Hỏi bất cứ điều gì về hóa học…", key="chat_in")
    if q:
        chat.append({"role": "user", "text": q})
    with box:
        if not chat:
            st.markdown("Xin chào! Mình là trợ lý hóa học. Bạn cứ hỏi về công thức, phản ứng, danh pháp, đồng phân, cách làm bài tập... nhé.")
        for m in chat:
            with st.chat_message("user" if m["role"] == "user" else "assistant"):
                st.markdown(m["text"])
        if q:
            with st.chat_message("assistant"):
                if st.session_state.n >= MAX_PER_SESSION:
                    ans = "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục."
                    st.markdown(errbanner(ans), unsafe_allow_html=True)
                else:
                    with st.spinner("Đang suy nghĩ…"):
                        try:
                            ans = chat_ai(chat)
                            st.session_state.n += 1
                        except Exception as e:
                            ans = "Có lỗi khi hỏi trợ lý, bạn thử lại sau nhé."
                            st.caption(f"Chi tiết kỹ thuật: {type(e).__name__}: {str(e)[:300]}")
                    st.markdown(ans)
            chat.append({"role": "assistant", "text": ans})
        if chat and st.button("Xóa cuộc trò chuyện", key="chat_clear"):
            st.session_state.chat = []
            st.rerun()


# ---------------- Mô phỏng thí nghiệm (phim hoạt hình trong trình duyệt) ----------------
# Kịch bản thí nghiệm có sẵn: chạy ngay, không tốn lượt gọi AI. Mỗi kịch bản có phương trình đã kiểm tra cân bằng.
def _sp(he_so, cong_thuc, smiles):
    return {"he_so": he_so, "cong_thuc": cong_thuc, "smiles": smiles}


SIM_BUILTINS = {
    "zn_hcl": {
        "kw": [["kem", "hcl"], ["zn", "hcl"], ["kem", "axit clohidric"], ["zn", "axit clohidric"], ["kem", "axit"]],
        "ten": "Kẽm tác dụng với dung dịch axit clohiđric",
        "dung_cu": [{"id": "A", "loai": "binh_tam_giac", "nhan": "Bình tam giác"}],
        "thu_khi": {"tu": "A", "khi": "H2", "mau": "#e6e9f2"},
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 6, "mo_ta": "Dùng ống nhỏ giọt cho khoảng 10 mL dung dịch HCl loãng vào bình tam giác. Lắp nút cao su có ống dẫn khí, đầu ống dẫn khí đặt dưới miệng ống thu khí úp ngược trong chậu nước.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd HCl", "mau": "#eef3ff", "muc": 0.4}]},
            {"tieu_de": "Cho kẽm vào", "thoi_gian": 7, "mo_ta": "Thả vài viên kẽm vào bình. Trên bề mặt viên kẽm xuất hiện nhiều bọt khí, viên kẽm nhỏ dần.",
             "hanh_dong": [{"loai": "them_ran", "dung_cu": "A", "chat": "Zn", "mau": "#9aa4ad", "kieu": "vien"}, {"loai": "sui_bot", "dung_cu": "A", "cuong_do": 0.75},
                           {"loai": "tan", "dung_cu": "A", "con_lai": 0.35, "thoi_gian": 8, "dong_thoi": True}]},
            {"tieu_de": "Thu khí hiđro", "thoi_gian": 7, "mo_ta": "Khí không màu, không mùi thoát ra theo ống dẫn khí và đẩy nước trong ống thu khí xuống. H2 ít tan trong nước nên thu được bằng cách đẩy nước.",
             "hanh_dong": [{"loai": "thu_khi", "luong": 0.8, "thoi_gian": 6}]},
            {"tieu_de": "Thử khí", "thoi_gian": 5, "mo_ta": "Đưa que đóm đang cháy lại gần miệng ống thu khí: khí cháy với ngọn lửa xanh nhạt và nghe tiếng nổ nhỏ “póp”, chứng tỏ khí là hiđro.",
             "hanh_dong": [{"loai": "que_dom", "ket_qua": "no_nho"}]},
        ],
        "hien_tuong": ["Có bọt khí không màu thoát ra trên bề mặt viên kẽm", "Viên kẽm tan dần", "Dung dịch trong suốt, không màu", "Khí thu được cháy có tiếng nổ nhỏ"],
        "phuong_trinh": "Zn + 2HCl -> ZnCl2 + H2↑",
        "chat_tham_gia": [_sp(1, "Zn", "[Zn]"), _sp(2, "HCl", "Cl")], "san_pham": [_sp(1, "ZnCl2", "Cl[Zn]Cl"), _sp(1, "H2", "[H][H]")],
        "giai_thich": "Kẽm đứng trước hiđro trong dãy hoạt động hóa học nên đẩy được hiđro ra khỏi dung dịch axit loãng.",
        "an_toan": ["Đeo kính bảo hộ và găng tay", "Không để lửa gần bình khi đang sinh nhiều khí H2", "Axit HCl ăn mòn, tránh để dính vào da"],
    },
    "cuso4_naoh": {
        "kw": [["cuso4", "naoh"], ["dong sunfat", "naoh"], ["cuso4", "natri hidroxit"], ["dong ii sunfat", "natri hidroxit"], ["dong (ii) sunfat", "naoh"]],
        "ten": "Dung dịch đồng(II) sunfat tác dụng với dung dịch natri hiđroxit",
        "dung_cu": [{"id": "A", "loai": "ong_nghiem", "nhan": "Ống nghiệm"}],
        "thu_khi": None,
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 5, "mo_ta": "Cho khoảng 2 mL dung dịch CuSO4 vào ống nghiệm. Dung dịch có màu xanh lam.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd CuSO4", "mau": "#2f7be0", "muc": 0.45}]},
            {"tieu_de": "Nhỏ dung dịch NaOH", "thoi_gian": 7, "mo_ta": "Nhỏ từ từ dung dịch NaOH vào ống nghiệm. Xuất hiện kết tủa màu xanh lam lơ lửng rồi lắng xuống đáy ống.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd NaOH", "mau": "#f3f7ff", "muc": 0.2},
                           {"loai": "ket_tua", "dung_cu": "A", "chat": "Cu(OH)2", "mau": "#4f93f2", "luong": 0.85, "thoi_gian": 4},
                           {"loai": "doi_mau", "dung_cu": "A", "mau": "#cfe3fb", "thoi_gian": 4, "dong_thoi": True}]},
            {"tieu_de": "Quan sát", "thoi_gian": 4, "mo_ta": "Kết tủa Cu(OH)2 màu xanh lam không tan, lắng ở đáy ống nghiệm; phần dung dịch phía trên nhạt màu đi.", "hanh_dong": []},
        ],
        "hien_tuong": ["Dung dịch CuSO4 màu xanh lam", "Khi nhỏ NaOH xuất hiện kết tủa màu xanh lam", "Kết tủa lắng xuống đáy ống nghiệm"],
        "phuong_trinh": "CuSO4 + 2NaOH -> Cu(OH)2↓ + Na2SO4",
        "chat_tham_gia": [_sp(1, "CuSO4", "[Cu+2].[O-]S(=O)(=O)[O-]"), _sp(2, "NaOH", "[Na+].[OH-]")],
        "san_pham": [_sp(1, "Cu(OH)2", "[Cu+2].[OH-].[OH-]"), _sp(1, "Na2SO4", "[Na+].[Na+].[O-]S(=O)(=O)[O-]")],
        "giai_thich": "Đây là phản ứng trao đổi: ion Cu2+ kết hợp với ion OH- tạo hiđroxit không tan Cu(OH)2.",
        "an_toan": ["Đeo kính bảo hộ và găng tay", "NaOH có tính ăn mòn, tránh dính vào da và mắt"],
    },
    "agno3_nacl": {
        "kw": [["agno3", "nacl"], ["bac nitrat", "natri clorua"], ["agno3", "muoi an"], ["agno3", "natri clorua"], ["bac nitrat", "nacl"]],
        "ten": "Dung dịch bạc nitrat tác dụng với dung dịch natri clorua",
        "dung_cu": [{"id": "A", "loai": "ong_nghiem", "nhan": "Ống nghiệm"}],
        "thu_khi": None,
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 5, "mo_ta": "Cho khoảng 2 mL dung dịch AgNO3 vào ống nghiệm. Dung dịch trong suốt, không màu.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd AgNO3", "mau": "#f2f6ff", "muc": 0.45}]},
            {"tieu_de": "Nhỏ dung dịch NaCl", "thoi_gian": 7, "mo_ta": "Nhỏ dung dịch NaCl vào ống nghiệm. Xuất hiện ngay kết tủa màu trắng dạng sữa, vón cục, làm đục dung dịch.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd NaCl", "mau": "#f3f7ff", "muc": 0.2},
                           {"loai": "ket_tua", "dung_cu": "A", "chat": "AgCl", "mau": "#fbfbf7", "luong": 0.9, "thoi_gian": 3.5},
                           {"loai": "doi_mau", "dung_cu": "A", "mau": "#e9ecf3", "thoi_gian": 3, "dong_thoi": True}]},
            {"tieu_de": "Quan sát", "thoi_gian": 4, "mo_ta": "Kết tủa trắng AgCl không tan trong nước và trong axit, lắng dần xuống đáy ống nghiệm.", "hanh_dong": []},
        ],
        "hien_tuong": ["Xuất hiện kết tủa màu trắng", "Kết tủa vón cục, lắng xuống đáy ống nghiệm"],
        "phuong_trinh": "AgNO3 + NaCl -> AgCl↓ + NaNO3",
        "chat_tham_gia": [_sp(1, "AgNO3", "[Ag+].[O-][N+](=O)[O-]"), _sp(1, "NaCl", "[Na+].[Cl-]")],
        "san_pham": [_sp(1, "AgCl", "[Ag+].[Cl-]"), _sp(1, "NaNO3", "[Na+].[O-][N+](=O)[O-]")],
        "giai_thich": "Ion Ag+ kết hợp với ion Cl- tạo AgCl không tan. Phản ứng này dùng để nhận biết ion clorua.",
        "an_toan": ["Đeo kính bảo hộ và găng tay", "AgNO3 làm đen da khi dính vào, rửa ngay bằng nhiều nước"],
    },
    "kmno4": {
        "kw": [["kmno4", "nhiet"], ["kali pemanganat", "nhiet"], ["kmno4", "oxi"], ["kmno4", "o2"], ["thuoc tim", "nhiet"], ["dieu che oxi"], ["dieu che khi oxi"]],
        "ten": "Nhiệt phân kali pemanganat điều chế khí oxi",
        "dung_cu": [{"id": "A", "loai": "ong_nghiem", "nhan": "Ống nghiệm"}],
        "thu_khi": {"tu": "A", "khi": "O2", "mau": "#dfeaf7"},
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 6, "mo_ta": "Cho một ít tinh thể KMnO4 màu tím đen vào ống nghiệm, nhét một nhúm bông ở miệng ống rồi đậy bằng nút cao su có ống dẫn khí. Kẹp ống nghiệm trên giá, đặt ống thu khí đầy nước úp trong chậu nước.",
             "hanh_dong": [{"loai": "them_ran", "dung_cu": "A", "chat": "KMnO4", "mau": "#5a1b6f", "kieu": "bot"}]},
            {"tieu_de": "Đun nóng", "thoi_gian": 8, "mo_ta": "Đốt đèn cồn và đun nóng ống nghiệm. Chất rắn đổi màu, khí oxi thoát ra theo ống dẫn khí.",
             "hanh_dong": [{"loai": "dun_nong", "bat": True, "dung_cu": "A"}, {"loai": "doi_mau_ran", "dung_cu": "A", "mau": "#2b2623", "thoi_gian": 5}]},
            {"tieu_de": "Thu khí oxi", "thoi_gian": 7, "mo_ta": "Khí O2 ít tan trong nước nên thu được bằng cách đẩy nước. Khí đẩy nước trong ống thu khí xuống.",
             "hanh_dong": [{"loai": "thu_khi", "luong": 0.8, "thoi_gian": 6}]},
            {"tieu_de": "Thử khí", "thoi_gian": 5, "mo_ta": "Đưa que đóm còn tàn đỏ vào miệng ống thu khí: tàn đỏ bùng cháy thành ngọn lửa, chứng tỏ khí là oxi.",
             "hanh_dong": [{"loai": "que_dom", "ket_qua": "bung_chay"}]},
            {"tieu_de": "Kết thúc", "thoi_gian": 5, "mo_ta": "Tháo ống dẫn khí ra khỏi chậu nước trước, sau đó mới tắt đèn cồn để nước không bị hút ngược vào ống nghiệm nóng.",
             "hanh_dong": [{"loai": "dun_nong", "bat": False, "dung_cu": "A"}]},
        ],
        "hien_tuong": ["Chất rắn màu tím đen chuyển sang màu sẫm hơn khi đun nóng", "Có khí không màu thoát ra, thu được khí trong ống thu khí", "Que đóm còn tàn đỏ bùng cháy trong khí thu được"],
        "phuong_trinh": "2KMnO4 -> K2MnO4 + MnO2 + O2↑",
        "chat_tham_gia": [_sp(2, "KMnO4", "[K+].[O-][Mn](=O)(=O)=O")], "san_pham": [_sp(1, "K2MnO4", "[K+].[K+].[O-][Mn](=O)(=O)[O-]"), _sp(1, "MnO2", "O=[Mn]=O"), _sp(1, "O2", "O=O")],
        "giai_thich": "KMnO4 bị nhiệt phân, giải phóng khí oxi. Nhúm bông ở miệng ống giữ bụi chất rắn không bị cuốn vào ống dẫn khí.",
        "an_toan": ["Đeo kính bảo hộ", "Tháo ống dẫn khí khỏi chậu nước trước khi tắt đèn cồn", "Ống nghiệm nóng không chạm tay, để nguội rồi mới rửa"],
    },
    "fe_cuso4": {
        "kw": [["fe", "cuso4"], ["sat", "cuso4"], ["dinh sat", "dong sunfat"], ["sat", "dong sunfat"], ["fe", "dong sunfat"], ["sat", "dong ii sunfat"]],
        "ten": "Sắt tác dụng với dung dịch đồng(II) sunfat",
        "dung_cu": [{"id": "A", "loai": "coc", "nhan": "Cốc thủy tinh"}],
        "thu_khi": None,
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 5, "mo_ta": "Cho khoảng 50 mL dung dịch CuSO4 màu xanh lam vào cốc thủy tinh.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd CuSO4", "mau": "#2f7be0", "muc": 0.55}]},
            {"tieu_de": "Nhúng đinh sắt", "thoi_gian": 8, "mo_ta": "Nhúng một chiếc đinh sắt đã làm sạch vào dung dịch. Sau một lúc, đinh sắt được phủ một lớp kim loại màu đỏ (đồng), màu xanh của dung dịch nhạt dần.",
             "hanh_dong": [{"loai": "them_ran", "dung_cu": "A", "chat": "Fe (đinh sắt)", "mau": "#8b9199", "kieu": "dinh"},
                           {"loai": "doi_mau_ran", "dung_cu": "A", "mau": "#c0642a", "thoi_gian": 6},
                           {"loai": "doi_mau", "dung_cu": "A", "mau": "#bfe6cf", "thoi_gian": 7, "dong_thoi": True}]},
            {"tieu_de": "Quan sát", "thoi_gian": 4, "mo_ta": "Dung dịch chuyển từ xanh lam sang xanh lục nhạt (FeSO4). Trên đinh sắt có lớp đồng màu đỏ bám vào.", "hanh_dong": []},
        ],
        "hien_tuong": ["Có lớp kim loại màu đỏ bám trên đinh sắt", "Dung dịch từ màu xanh lam chuyển sang xanh lục nhạt"],
        "phuong_trinh": "Fe + CuSO4 -> FeSO4 + Cu↓",
        "chat_tham_gia": [_sp(1, "Fe", "[Fe]"), _sp(1, "CuSO4", "[Cu+2].[O-]S(=O)(=O)[O-]")], "san_pham": [_sp(1, "FeSO4", "[Fe+2].[O-]S(=O)(=O)[O-]"), _sp(1, "Cu", "[Cu]")],
        "giai_thich": "Sắt hoạt động hóa học mạnh hơn đồng nên đẩy đồng ra khỏi dung dịch muối của nó.",
        "an_toan": ["Đeo găng tay và kính bảo hộ", "Muối đồng độc, không để dính vào miệng và mắt"],
    },
    "naoh_hcl": {
        "kw": [["naoh", "hcl"], ["phenolphtalein"], ["trung hoa"], ["natri hidroxit", "axit clohidric"]],
        "ten": "Phản ứng trung hòa NaOH bằng HCl với chất chỉ thị phenolphtalein",
        "dung_cu": [{"id": "A", "loai": "coc", "nhan": "Cốc thủy tinh"}],
        "thu_khi": None,
        "buoc": [
            {"tieu_de": "Chuẩn bị", "thoi_gian": 5, "mo_ta": "Cho khoảng 30 mL dung dịch NaOH loãng vào cốc, nhỏ vài giọt phenolphtalein. Dung dịch chuyển sang màu hồng.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd NaOH", "mau": "#f4f8ff", "muc": 0.45}, {"loai": "doi_mau", "dung_cu": "A", "mau": "#f07db8", "thoi_gian": 2.5}]},
            {"tieu_de": "Nhỏ từ từ dung dịch HCl", "thoi_gian": 8, "mo_ta": "Nhỏ từng giọt dung dịch HCl vào cốc và khuấy đều. Màu hồng nhạt dần rồi mất hẳn khi axit vừa đủ trung hòa hết NaOH.",
             "hanh_dong": [{"loai": "them_long", "dung_cu": "A", "chat": "dd HCl", "mau": "#f4f8ff", "muc": 0.2, "thoi_gian": 4},
                           {"loai": "doi_mau", "dung_cu": "A", "mau": "#f3f7ff", "thoi_gian": 5, "dong_thoi": True}]},
            {"tieu_de": "Quan sát", "thoi_gian": 4, "mo_ta": "Dung dịch trở về không màu: bazơ đã bị trung hòa hết, tạo muối NaCl và nước.", "hanh_dong": []},
        ],
        "hien_tuong": ["Dung dịch NaOH có phenolphtalein màu hồng", "Nhỏ HCl vào, màu hồng nhạt dần rồi mất hẳn"],
        "phuong_trinh": "NaOH + HCl -> NaCl + H2O",
        "chat_tham_gia": [_sp(1, "NaOH", "[Na+].[OH-]"), _sp(1, "HCl", "Cl")], "san_pham": [_sp(1, "NaCl", "[Na+].[Cl-]"), _sp(1, "H2O", "O")],
        "giai_thich": "Phenolphtalein có màu hồng trong môi trường bazơ và không màu trong môi trường axit hoặc trung tính, nên đổi màu báo hiệu điểm trung hòa.",
        "an_toan": ["Đeo kính bảo hộ và găng tay", "NaOH và HCl đều ăn mòn, nhỏ từ từ và tránh bắn"],
    },
}


def _fold_sim(s):
    s = unicodedata.normalize("NFD", (s or "").lower().replace("đ", "d"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def sim_match_builtin(text):
    """Tìm kịch bản có sẵn khớp với câu người dùng gõ (khớp khi chứa đủ các từ khóa của một nhóm)."""
    f = _fold_sim(text)
    for key, sc in SIM_BUILTINS.items():
        for grp in sc["kw"]:
            if all(tok in f for tok in grp):
                return key
    return None


SIM_TEMPLATE = r'''<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Mô phỏng thí nghiệm</title>
<style>
*{box-sizing:border-box}body{margin:0;font-family:Arial,"DejaVu Sans",sans-serif;background:#f5f7ff;color:#1b2150}
#wrap{max-width:960px;margin:0 auto}
#cv{width:100%;height:auto;display:block;border-radius:14px;background:#0f1640;box-shadow:0 8px 26px rgba(40,50,160,.25)}
#bar{display:flex;align-items:center;gap:8px;flex-wrap:wrap;margin-top:8px}
button,select{font:inherit;border:1px solid #d9defa;background:#fff;color:#1b2150;border-radius:10px;padding:7px 12px;cursor:pointer;font-weight:600}
button.pri{background:linear-gradient(135deg,#3b6cf6,#6d4df0);color:#fff;border:0}
button:hover{filter:brightness(.97)}button.off{background:#eef0fb;color:#8a90b8}
#sl{flex:1 1 200px;accent-color:#4f5cf0}
#tm{font-size:13px;color:#6b7299;min-width:84px;text-align:right}
#st{font-size:12.5px;color:#6b7299;margin-top:6px;min-height:16px}
</style></head><body><div id="wrap">
<canvas id="cv" width="960" height="540"></canvas>
<div id="bar">
<button id="bR" title="Về đầu">⏮</button><button id="bP" title="Bước trước">⏪</button>
<button id="bPlay" class="pri" style="min-width:92px">▶ Phát</button>
<button id="bN" title="Bước sau">⏩</button>
<input id="sl" type="range" min="0" max="1000" value="0">
<span id="tm">0:00</span>
<span style="display:flex;align-items:center;gap:6px;font-size:13px;color:#6b7299" title="Tốc độ phát tùy ý">⏱ <input id="sp" type="range" min="0.25" max="3" step="0.05" value="1" style="width:110px;accent-color:#4f5cf0"><b id="spl" style="min-width:42px">1,00×</b></span><button id="bVoice" title="Bật hoặc tắt giọng đọc lời dẫn">🗣️ Giọng đọc: Bật</button><button id="bSfx" title="Bật hoặc tắt hiệu ứng âm thanh">🔊 Âm thanh: Bật</button><button id="bTeach" title="Hiện hoặc ẩn cô giáo làm thí nghiệm">👩‍🏫 Cô giáo: Bật</button>
<select id="vsel" style="display:none" title="Chọn giọng đọc của trình duyệt"></select>
<button id="bRec" title="Ghi lại thành video .mp4">🎬 Tải video (.mp4)</button>
<button id="bHtml" title="Lưu bản mô phỏng tự chạy được">💾 Tải .html</button>
</div><div id="st"></div></div>
<script>
const D = __DATA__;
const cv = document.getElementById("cv"), g = cv.getContext("2d"), W = 960, H = 540;
const FONT = 'Arial,"DejaVu Sans",sans-serif';
const clamp = (x, a = 0, b = 1) => Math.max(a, Math.min(b, x));
const lerp = (a, b, t) => a + (b - a) * t;
const hex2 = (h) => { h = String(h || "#ffffff").replace("#", ""); if (h.length === 3) h = h.split("").map(c => c + c).join(""); return [parseInt(h.substr(0, 2), 16) || 0, parseInt(h.substr(2, 2), 16) || 0, parseInt(h.substr(4, 2), 16) || 0]; };
const mixc = (a, b, t) => [lerp(a[0], b[0], t), lerp(a[1], b[1], t), lerp(a[2], b[2], t)];
const rgba = (c, a = 1) => `rgba(${c[0] | 0},${c[1] | 0},${c[2] | 0},${a})`;
const rnd = (i) => { const x = Math.sin(i * 127.1 + 311.7) * 43758.5453; return x - Math.floor(x); };
const SUBD = "₀₁₂₃₄₅₆₇₈₉";
const sub = (s) => String(s || "").replace(/(?<=[A-Za-z\)])(\d+)/g, m => m.split("").map(d => SUBD[+d]).join("")).replace(/->/g, "→");
const ease = (p) => p < .5 ? 2 * p * p : 1 - Math.pow(-2 * p + 2, 2) / 2;

// ---------- dụng cụ ----------
const GEO = { ong_nghiem: { w: 66, h: 230, cap: 190 }, coc: { w: 150, h: 170, cap: 140 }, binh_tam_giac: { w: 170, h: 190, cap: 105 } };
const VS = (D.dung_cu || []).slice(0, 3);
const GAS = D.thu_khi && D.thu_khi.tu ? D.thu_khi : null;
const GAS_V = GAS ? Math.max(0, VS.findIndex(v => v.id === GAS.tu)) : -1;
const lampV = (() => { let k = -1; (D.buoc || []).forEach(s => (s.hanh_dong || []).forEach(a => { if (a.loai === "dun_nong" && k < 0) k = Math.max(0, VS.findIndex(v => v.id === a.dung_cu)); })); return k; })();
VS.forEach((v, i) => {
  const n = VS.length, left = GAS ? 40 : 60, right = GAS ? 560 : 900;
  v.x = left + (right - left) * (i + .5) / n;
  v.base = (i === lampV) ? 336 : 398;
  v.geo = GEO[v.loai] || GEO.ong_nghiem;
});
let SC = 1;
if (!GAS && lampV < 0 && VS.length) { const mt = Math.min(...VS.map(v => v.base - v.geo.h)); SC = Math.max(1, Math.min(VS.length === 1 ? 1.4 : VS.length === 2 ? 1.25 : 1.1, (398 - 66) / (398 - (mt - 72)))); }
function vIndex(id) { const k = VS.findIndex(v => v.id === id); return k >= 0 ? k : 0; }

// ---------- kịch bản thời gian ----------
const narrDur = (txt) => String(txt || "").length / 15 + 1.0;   // ước lượng khi chưa có tệp giọng đọc
const BLOCK = { them_long: 2.6, them_ran: 1.8, dun_nong: 2.6, doi_mau: 3, doi_mau_ran: 3, ket_tua: 2.8, tan: 3, thu_khi: 4, que_dom: 3.2 };
const steps = (D.buoc || []).map(s => Object.assign({}, s));
const LEAD = 1.0;   // 1 giây đầu: cô bước vào bàn thí nghiệm
const events = []; let T = LEAD;
steps.forEach((s, si) => {
  s.t0 = T; let tt = .35, lastOff = .35; const evs = [];
  (s.hanh_dong || []).forEach(a => {
    const e = Object.assign({}, a), dur = Math.max(.4, +a.thoi_gian || BLOCK[a.loai] || .5);
    e.off0 = a.dong_thoi ? lastOff : tt; e.dur = dur; lastOff = e.off0;
    if (BLOCK[a.loai] && !a.dong_thoi) tt = e.off0 + dur + .15; else if (!BLOCK[a.loai]) tt = Math.max(tt, e.off0 + .4);
    evs.push(e);
  });
  const speech = s.am_dur ? s.am_dur + .6 : narrDur(s.mo_ta);
  s.auto = !s.thoi_gian_cung;
  const dur = s.thoi_gian_cung ? Math.max(2, +s.thoi_gian_cung) : Math.max(+s.thoi_gian || 4, tt + .9, speech);
  // kéo giãn (hoặc nén) chuyển động để phim chạy liền mạch suốt lời đọc, không đứng hình
  const body = tt - .35, room = dur - .35 - .9, k = body > 0 ? clamp(room / body, .5, 3) : 1;
  evs.forEach(e => { e.t0 = T + .35 + (e.off0 - .35) * k; e.t1 = e.t0 + e.dur * k; e.si = si; events.push(e); });
  s.k = (s.am_dur && s.thoi_gian_cung) ? clamp(s.am_dur / Math.max(1, dur - .8), .7, 1.6) : 1;   // tốc độ đọc để vừa với thời lượng tự đặt
  s.t1 = T + dur; T += dur;
});
const endText = "Kết luận. " + (D.hien_tuong || []).slice(0, 5).join(". ") + ". " + (D.giai_thich || "");
const endDur = D.thoi_gian_ket ? Math.max(3, +D.thoi_gian_ket) : Math.max(8, D.am_ket_dur ? D.am_ket_dur + 1 : narrDur(endText));
const endCard = { tieu_de: "Kết luận", t0: T, t1: T + endDur, end: true, am: D.am_ket, am_dur: D.am_ket_dur, mo_ta: endText, k: (D.am_ket_dur && D.thoi_gian_ket) ? clamp(D.am_ket_dur / Math.max(1, endDur - 1), .7, 1.6) : 1 };
steps.push(endCard); const TOTAL = endCard.t1;

// ---------- trạng thái tại thời điểm t ----------
function stateAt(t) {
  const vs = VS.map(() => ({ vol: 0, col: [235, 240, 255], solid: null, scol: null, ssize: 1, pr: 0, prcol: [255, 255, 255], rate: 0, items: [], prNew: 0, fx: [] }));
  let lamp = false, gasVol = 0, gasCol = hex2(GAS && GAS.mau || "#d8dde8"), splint = null, gasOn = false, capOn = false, lift = 0;
  for (let ei = 0; ei < events.length; ei++) {
    const e = events[ei]; if (t < e.t0) continue;
    const p = clamp((t - e.t0) / (e.t1 - e.t0)), k = vIndex(e.dung_cu), v = vs[k] || vs[0]; if (!v) continue;
    switch (e.loai) {
      case "them_long": {
        const dv = clamp(+e.muc || .3, .03, .85) * (TEACH_ON ? dripP(e, t) : p), v0 = v.vol, c1 = hex2(e.mau || "#dfe8ff");
        v.col = v0 + dv > 0 ? mixc(v.col, c1, v0 < .001 ? 1 : dv / (v0 + dv)) : c1; v.vol = clamp(v0 + dv, 0, .95);
        if (p < 1) v.fx.push({ t: "long", c: c1, p, id: ei }); if (p > .5 && !v.items.includes(e.chat)) v.items.push(e.chat); break; }
      case "them_ran": {
        let sp = p; if (TEACH_ON) { const tr = relT(e); if (t - e.t0 < tr) { v.ranPend = true; break; } sp = clamp((t - e.t0 - tr) / 1.2); }   // chỉ rơi khi cô buông tay
        v.solid = e.kieu || "vien"; v.scol = hex2(e.mau || "#999999"); v.ssize = 1; v.sp = sp;
        if (sp < 1) v.fx.push({ t: "ran", c: v.scol, p: sp, kieu: v.solid, id: ei }); if (sp > .5 && !v.items.includes(e.chat)) v.items.push(e.chat); break; }
      case "doi_mau": v.col = mixc(v.col, hex2(e.mau || "#ffffff"), ease(p)); break;
      case "doi_mau_ran": if (v.scol) v.scol = mixc(v.scol, hex2(e.mau || "#999999"), ease(p)); break;
      case "ket_tua": v.prcol = hex2(e.mau || "#ffffff"); v.pr = clamp(v.pr + clamp(+e.luong || .5) * p * 1, 0, 1); if (p < 1) { v.prNew = p; v.fx.push({ t: "tua", p, id: ei }); } if (p > .5 && e.chat && !v.items.includes(e.chat)) v.items.push(e.chat); break;
      case "sui_bot": if (TEACH_ON && v.ranPend) break; v.rate = clamp(+e.cuong_do || 0); break;
      case "tan": v.ssize = lerp(1, clamp(e.con_lai === undefined ? .15 : +e.con_lai), p); break;
      case "dun_nong": { if (TEACH_ON) { const el = t - e.t0, A = phs(e).A; if (e.bat) { lamp = el >= A + .4; capOn = false; } else { lamp = el < A + .2; capOn = el >= A + .2; } } else { lamp = !!e.bat; capOn = false; } break; }
      case "thu_khi": gasOn = p < 1; gasVol = clamp(gasVol + clamp(+e.luong || .7) * p); break;
      case "que_dom": { if (TEACH_ON) { const si = splintInfo(e, t); lift = si.lift; splint = si.vis ? { p, kq: e.ket_qua || "bung_chay", id: ei, eh: si.eh } : null; } else splint = { p, kq: e.ket_qua || "bung_chay", id: ei }; break; }
    }
  }
  return { vs, lamp, gasVol, gasCol, splint, gasOn, capOn, lift };
}

// ---------- vẽ ----------
function outline(v) {
  const p = new Path2D(), x = v.x, b = v.base, G = v.geo, w = G.w, h = G.h, top = b - h;
  if (v.loai === "coc") { p.moveTo(x - w / 2, top); p.lineTo(x - w / 2 + 4, b - 8); p.quadraticCurveTo(x - w / 2 + 4, b, x - w / 2 + 14, b); p.lineTo(x + w / 2 - 14, b); p.quadraticCurveTo(x + w / 2 - 4, b, x + w / 2 - 4, b - 8); p.lineTo(x + w / 2, top); }
  else if (v.loai === "binh_tam_giac") { const nw = 46; p.moveTo(x - nw / 2, top); p.lineTo(x - nw / 2, top + 66); p.lineTo(x - w / 2 + 8, b - 10); p.quadraticCurveTo(x - w / 2, b, x - w / 2 + 14, b); p.lineTo(x + w / 2 - 14, b); p.quadraticCurveTo(x + w / 2, b, x + w / 2 - 8, b - 10); p.lineTo(x + nw / 2, top + 66); p.lineTo(x + nw / 2, top); }
  else { const r = w / 2; p.moveTo(x - r, top); p.lineTo(x - r, b - r); p.arc(x, b - r, r, Math.PI, 0, true); p.lineTo(x + r, top); }
  return p;
}
function surfaceY(v, vol) { const G = v.geo; return v.base - 5 - vol * G.cap; }
function halfW(v, y) { // nửa bề rộng bên trong tại độ cao y
  const G = v.geo, top = v.base - G.h;
  if (v.loai === "binh_tam_giac") { const nb = top + 66; if (y <= nb) return 20; const r = clamp((y - nb) / (v.base - nb)); return lerp(20, G.w / 2 - 8, r); }
  if (v.loai === "coc") return G.w / 2 - 6; return G.w / 2 - 6;
}
function drawBench() {
  const gr = g.createLinearGradient(0, 0, 0, H); gr.addColorStop(0, "#1c2670"); gr.addColorStop(.75, "#141b52"); gr.addColorStop(1, "#0d1240");
  g.fillStyle = gr; g.fillRect(0, 0, W, H);
  g.fillStyle = "rgba(255,255,255,.04)"; for (let i = 0; i < 16; i++) { g.beginPath(); g.arc(80 + i * 58, 70 + rnd(i) * 160, 20 + rnd(i + 9) * 40, 0, 7); g.fill(); }
  const wd = g.createLinearGradient(0, 398, 0, 424); wd.addColorStop(0, "#caa270"); wd.addColorStop(1, "#8d6a3f");
  g.fillStyle = wd; g.fillRect(0, 398, W, 26); g.fillStyle = "rgba(0,0,0,.18)"; g.fillRect(0, 424, W, 6);
}
function drawLamp(x, on, t, cap) {
  g.fillStyle = "#cfd8e8"; g.strokeStyle = "#8aa0c0"; g.lineWidth = 2;
  g.beginPath(); g.moveTo(x - 16, 382); g.lineTo(x - 26, 394); g.quadraticCurveTo(x - 28, 398, x - 20, 398); g.lineTo(x + 20, 398); g.quadraticCurveTo(x + 28, 398, x + 26, 394); g.lineTo(x + 16, 382); g.closePath(); g.fill(); g.stroke();
  g.fillStyle = "rgba(255,210,120,.55)"; g.fillRect(x - 21, 388, 42, 9);
  g.fillStyle = "#6b7280"; g.fillRect(x - 8, 372, 16, 11); g.fillStyle = "#333"; g.fillRect(x - 1.5, 362, 3, 11);
  if (cap) { capShape(x, 357); g.fillStyle = "rgba(205,228,250,.62)"; g.fill(); g.strokeStyle = "rgba(140,172,215,.98)"; g.lineWidth = 2; g.stroke(); g.fillStyle = "#cfd8e8"; g.beginPath(); g.arc(x, 344, 3.4, 0, 7); g.fill(); }
  if (on) {
    const sw = Math.sin(t * 11) * 2, hh = 36 + Math.sin(t * 17) * 3;
    const fl = (col, s) => { g.fillStyle = col; g.beginPath(); g.moveTo(x, 366); g.bezierCurveTo(x - 13 * s, 354, x - 8 * s, 366 - hh * s * .6, x + sw, 366 - hh * s); g.bezierCurveTo(x + 8 * s, 366 - hh * s * .6, x + 13 * s, 354, x, 366); g.fill(); };
    const gl = g.createRadialGradient(x, 340, 4, x, 340, 70); gl.addColorStop(0, "rgba(255,170,60,.35)"); gl.addColorStop(1, "rgba(255,170,60,0)"); g.fillStyle = gl; g.fillRect(x - 80, 270, 160, 130);
    fl("rgba(255,120,20,.75)", 1); fl("rgba(255,210,70,.9)", .68); fl("rgba(120,170,255,.9)", .32);
  }
}
function drawStand(v) {
  const x = v.x - v.geo.w / 2 - 34; g.fillStyle = "#6b7280"; g.fillRect(x - 3, 168, 6, 232); g.fillStyle = "#4b5563"; g.fillRect(x - 24, 396, 48, 6);
  const y = v.base - v.geo.h + 40; g.fillStyle = "#6b7280"; g.fillRect(x, y - 3, v.x - x, 6); g.fillStyle = "#374151"; g.fillRect(v.x - v.geo.w / 2 - 5, y - 9, v.geo.w + 10, 8);
}
function drawGlass(v, st, t) {
  const path = outline(v), G = v.geo, b = v.base, x = v.x;
  g.save(); g.clip(path);
  if (st.vol > 0.002) {
    const sy = surfaceY(v, st.vol), grd = g.createLinearGradient(0, sy, 0, b); grd.addColorStop(0, rgba(st.col, .72)); grd.addColorStop(1, rgba(st.col, .92));
    g.fillStyle = grd; g.fillRect(x - G.w, sy, G.w * 2, b - sy + 4);
    g.fillStyle = rgba(mixc(st.col, [255, 255, 255], .5), .8); g.beginPath(); g.ellipse(x, sy, halfW(v, sy), 3, 0, 0, 7); g.fill();
  }
  // chất rắn
  if (st.solid && st.ssize > .02) {
    const sc = st.scol || [150, 150, 150], s = st.ssize, sp = st.sp === undefined ? 1 : st.sp;
    const fall = sp >= 1 ? 0 : (1 - ease(clamp(sp / .5))) * (TEACH_ON ? G.h * .6 : 120);
    g.fillStyle = rgba(sc, 1); g.strokeStyle = rgba(mixc(sc, [0, 0, 0], .35), 1); g.lineWidth = 1.5;
    if (st.solid === "bot") { g.beginPath(); g.ellipse(x, b - 4 - fall, 28 * s, 15 * s + 2, 0, Math.PI, 0); g.fill(); g.stroke(); }
    else if (st.solid === "dinh") { g.save(); g.translate(x + 2, b - 6 - fall); g.rotate(-.12); g.fillRect(-3, -G.h * .5, 6, G.h * .5 * clamp(s + .15, 0, 1)); g.restore(); }
    else if (st.solid === "mieng") { const fy = surfaceY(v, st.vol) - 5 + Math.sin(t * 5) * 1.5; g.fillRect(x - 10 * s + Math.sin(t * 3) * 6 * (st.rate), (st.vol > .05 ? fy : b - 10) - fall, 20 * s, 10 * s); }
    else { g.beginPath(); g.ellipse(x - 8, b - 9 - fall, 15 * s + 1, 8 * s + 1, -.2, 0, 7); g.fill(); g.stroke(); }
  }
  // kết tủa
  if (st.pr > .01) {
    const hgt = 6 + st.pr * G.cap * .24; g.fillStyle = rgba(st.prcol, .96);
    g.beginPath(); g.moveTo(x - G.w, b + 2); for (let i = 0; i <= 24; i++) { const xx = x - G.w / 2 + 6 + (G.w - 12) * i / 24; g.lineTo(xx, b - hgt * (.8 + .25 * rnd(i + 3))); } g.lineTo(x + G.w, b + 2); g.closePath(); g.fill();
    if (st.prNew > 0 && st.prNew < 1 && st.vol > .05) { const top = surfaceY(v, st.vol); for (let i = 0; i < 46; i++) { const px = x + (rnd(i) - .5) * halfW(v, b - 20) * 1.7, ph = (t * (22 + rnd(i + 5) * 26) + rnd(i + 9) * 100) % (b - top - hgt); g.fillStyle = rgba(st.prcol, .85); g.beginPath(); g.arc(px, top + ph, 2 + rnd(i + 2) * 2, 0, 7); g.fill(); } }
  }
  // bọt khí
  if (st.rate > .02 && (st.vol > .04 || st.solid)) {
    const sy = st.vol > .04 ? surfaceY(v, st.vol) : b - 30, n = Math.round(8 + st.rate * 34), hh = b - 8 - sy;
    for (let i = 0; i < n; i++) { const sp = 26 + rnd(i + 90) * 46, ph = (t * sp * (.6 + st.rate) + rnd(i) * hh) % hh, hw = Math.max(6, halfW(v, b - ph) - 6), bx = x + (rnd(i + 50) - .5) * hw * 1.5 + Math.sin(t * 4 + i) * 2, r = 1.8 + rnd(i + 70) * 3.2;
      g.strokeStyle = "rgba(255,255,255,.85)"; g.fillStyle = "rgba(255,255,255,.28)"; g.lineWidth = 1.2; g.beginPath(); g.arc(bx, b - 8 - ph, r, 0, 7); g.fill(); g.stroke(); }
  }
  g.restore();
  g.fillStyle = "rgba(210,230,255,.10)"; g.fill(path);
  g.strokeStyle = "rgba(190,215,250,.95)"; g.lineWidth = 3; g.stroke(path);
  g.strokeStyle = "rgba(255,255,255,.45)"; g.lineWidth = 3; g.beginPath(); g.moveTo(x - G.w / 2 + 9, b - G.h + 22); g.lineTo(x - G.w / 2 + 9, b - 38); g.stroke();
}
function drawFx(v, st, t) { // dòng chảy, giọt, vật rơi
  const x = v.x, top = v.base - v.geo.h;
  if (TEACH_ON) { const dr = POSE && POSE.drip; if (dr && dr.vi === VS.indexOf(v)) { const sy = surfaceY(v, Math.max(st.vol, .02)); g.fillStyle = rgba(dr.c, .95); g.strokeStyle = "rgba(120,150,200,.6)"; g.lineWidth = 1; for (let k = 0; k < 3; k++) { const ph = (t * 1.6 + k / 3) % 1; g.beginPath(); g.ellipse(dr.x, dr.y + 4 + ph * Math.max(2, sy - dr.y - 4), 2.7, 4.3, 0, 0, 7); g.fill(); g.stroke(); } }
    st.fx.forEach(f => { if (f.t === "ran" && f.kieu === "bot" && f.p < .6) { g.fillStyle = rgba(f.c, .95); for (let i = 0; i < 16; i++) { const pr = clamp(f.p / .6 * 1.5 - i * .045, 0, 1); if (pr <= 0 || pr >= 1) continue; g.beginPath(); g.arc(x + Math.sin(i * 7.3) * 4 * pr, lerp(top - 26, v.base - 14, pr * pr), 1.7, 0, 7); g.fill(); } } });   // bột chảy từ thìa xuống đáy
    return; }
  st.fx.forEach(f => {
    if (f.t === "long") {
      const sy = surfaceY(v, Math.max(st.vol, .02)), tx = x + 40;
      g.fillStyle = "#8b3a3a"; g.beginPath(); g.arc(tx, top - 78, 11, 0, 7); g.fill();
      g.fillStyle = "rgba(235,240,255,.85)"; g.strokeStyle = "rgba(150,170,210,.9)"; g.lineWidth = 2; g.fillRect(tx - 6, top - 70, 12, 50); g.strokeRect(tx - 6, top - 70, 12, 50);
      g.beginPath(); g.moveTo(tx - 6, top - 20); g.lineTo(tx - 2, top - 8); g.lineTo(tx + 2, top - 8); g.lineTo(tx + 6, top - 20); g.closePath(); g.fill(); g.stroke();
      g.fillStyle = rgba(f.c, .95); g.fillRect(tx - 4, top - 62, 8, 38);
      for (let k = 0; k < 4; k++) { const ph = (t * 1.6 + k / 4) % 1, yy = top - 8 + ph * (sy - top + 8); g.beginPath(); g.ellipse(tx + (x - tx) * ph * .98, yy, 3, 5, 0, 0, 7); g.fill(); }
    } else if (f.t === "ran") {
      const y = lerp(top - 60, v.base - 20, ease(clamp(f.p / .5))); if (f.p < .55) { g.fillStyle = rgba(f.c, 1); g.beginPath(); g.ellipse(x - 6, y, 13, 7, -.3, 0, 7); g.fill(); }
      g.fillStyle = "rgba(255,255,255,.6)"; g.fillRect(x - 38, top - 70, 70, 5); g.fillRect(x - 38, top - 70, 5, 14);
    }
  });
}
function drawGasKit(st, t) {
  if (!GAS) return; const v = VS[GAS_V], x = v.x, top = v.base - v.geo.h, tx = 690;
  // nút cao su + ống dẫn khí
  g.fillStyle = "#8b5a2b"; g.beginPath(); g.moveTo(x - 17, top - 6); g.lineTo(x + 17, top - 6); g.lineTo(x + 13, top + 18); g.lineTo(x - 13, top + 18); g.closePath(); g.fill();
  g.strokeStyle = "rgba(210,230,250,.95)"; g.lineWidth = 6; g.lineJoin = "round"; g.beginPath(); g.moveTo(x, top + 6); g.lineTo(x, top - 46); g.lineTo(tx - 40, top - 46); g.lineTo(tx - 40, 372); g.stroke();
  g.strokeStyle = "rgba(60,80,120,.5)"; g.lineWidth = 1.5; g.stroke();
  // chậu nước
  const bx = tx - 130, by = 340, bw = 260, bh = 58;
  g.fillStyle = "rgba(90,160,230,.55)"; g.fillRect(bx + 3, by + 14, bw - 6, bh - 17);
  g.strokeStyle = "rgba(190,215,250,.95)"; g.lineWidth = 4; g.strokeRect(bx, by, bw, bh);
  g.strokeStyle = "rgba(255,255,255,.5)"; g.lineWidth = 2; g.beginPath(); g.moveTo(bx + 4, by + 14); for (let i = 0; i <= 26; i++) g.lineTo(bx + 4 + i * 9.6, by + 14 + Math.sin(t * 3 + i) * 1.3); g.stroke();
  // ống thu khí úp ngược
  const LF = st.lift || 0, jx = tx - 40, jw = 70, jt = 232 - LIFT_PX * LF, jb = 372 - LIFT_PX * LF, inner = jb - jt;
  const waterTop = lerp(jt + st.gasVol * inner * .95, jb + 6, ease(clamp(LF * 2.4)));   // nhấc lên khỏi chậu thì nước chảy hết
  g.save(); g.beginPath(); g.rect(jx - jw / 2 + 2, jt + 2, jw - 4, inner); g.clip();
  g.fillStyle = "rgba(90,160,230,.55)"; g.fillRect(jx - jw / 2, waterTop, jw, Math.max(0, jb - waterTop));
  if (st.gasVol > .01) { g.fillStyle = rgba(st.gasCol, .35); g.fillRect(jx - jw / 2, jt, jw, waterTop - jt); }
  g.restore();
  g.strokeStyle = "rgba(190,215,250,.95)"; g.lineWidth = 3; g.beginPath(); g.moveTo(jx - jw / 2, jb + 8); g.lineTo(jx - jw / 2, jt + 8); g.quadraticCurveTo(jx - jw / 2, jt, jx - jw / 2 + 8, jt); g.lineTo(jx + jw / 2 - 8, jt); g.quadraticCurveTo(jx + jw / 2, jt, jx + jw / 2, jt + 8); g.lineTo(jx + jw / 2, jb + 8); g.stroke();
  g.fillStyle = "rgba(255,255,255,.9)"; g.font = `bold 14px ${FONT}`; g.textAlign = "center"; g.fillText(sub(GAS.khi || "Khí"), jx, jt - 8);
  if (st.gasOn) for (let i = 0; i < 9; i++) { const ph = (t * 1.3 + i / 9) % 1, by2 = 372 - ph * 90; g.fillStyle = "rgba(255,255,255,.8)"; g.beginPath(); g.arc(jx + Math.sin(i * 3 + t * 4) * 6, by2, 2.5 + rnd(i) * 2, 0, 7); g.fill(); }
}
function drawSplint(st, t) { if (TEACH_ON) { if (POSE && POSE.splint) drawSplintFx(POSE.splint, t); return; } drawSplintOld(st, t); }
function drawSplintOld(st, t) {
  if (!st.splint) return; const p = st.splint.p, kq = st.splint.kq, jx = 650, top = 232;
  const mv = ease(clamp(p / .55)), x = lerp(860, jx + 6, mv), y = lerp(120, top - 20, mv);
  g.save(); g.translate(x, y); g.rotate(.55); g.fillStyle = "#c9a36a"; g.fillRect(-4, -110, 8, 110); g.restore();
  const tipx = x + Math.sin(.55) * 0, tipy = y;
  const lit = kq !== "tat" ? 1 : clamp(1 - (p - .6) * 3);
  if (p > .5 && kq === "bung_chay") { const r = 8 + (p - .5) * 60; const gr = g.createRadialGradient(tipx, tipy, 2, tipx, tipy, r); gr.addColorStop(0, "rgba(255,255,200,1)"); gr.addColorStop(.4, "rgba(255,190,60,.9)"); gr.addColorStop(1, "rgba(255,120,0,0)"); g.fillStyle = gr; g.beginPath(); g.arc(tipx, tipy, r, 0, 7); g.fill(); }
  else if (p > .75 && kq === "no_nho") { const r = (p - .75) * 160; g.strokeStyle = `rgba(255,230,120,${1 - (p - .75) * 3})`; g.lineWidth = 4; g.beginPath(); g.arc(tipx, tipy, r, 0, 7); g.stroke(); g.fillStyle = "rgba(255,240,180,.9)"; g.font = `bold 22px ${FONT}`; g.textAlign = "center"; g.fillText("pop!", tipx, tipy - r - 8); }
  else { g.fillStyle = `rgba(255,${90 + lit * 50},30,${.5 + .5 * lit})`; g.beginPath(); g.arc(tipx, tipy, 5, 0, 7); g.fill(); }
}
function wrap(text, maxW, font) { g.font = font; const words = String(text || "").split(/\s+/), lines = []; let cur = ""; words.forEach(w => { const tt = cur ? cur + " " + w : w; if (g.measureText(tt).width > maxW && cur) { lines.push(cur); cur = w; } else cur = tt; }); if (cur) lines.push(cur); return lines; }
function rrect(x, y, w, h, r) { g.beginPath(); g.moveTo(x + r, y); g.arcTo(x + w, y, x + w, y + h, r); g.arcTo(x + w, y + h, x, y + h, r); g.arcTo(x, y + h, x, y, r); g.arcTo(x, y, x + w, y, r); g.closePath(); }
function curStep(t) { for (let i = 0; i < steps.length; i++) if (t < steps[i].t1) return i; return steps.length - 1; }
function drawHud(t) {
  const si = curStep(t), s = steps[si], n = steps.length - 1;
  g.fillStyle = "rgba(8,12,40,.7)"; g.fillRect(0, 0, W, 44);
  g.fillStyle = "#fff"; g.textAlign = "left"; g.font = `bold 21px ${FONT}`; g.fillText(sub(D.ten || "Mô phỏng thí nghiệm"), 18, 29);
  g.fillStyle = "#9fb3ff"; g.font = `13px ${FONT}`; g.textAlign = "right"; g.fillText("MolSnap · Mô phỏng thí nghiệm", W - 16, 28);
  if (s.end) return;
  const x = 20, y = 440, w = W - 40, h = 90; g.fillStyle = "rgba(8,12,40,.86)"; rrect(x, y, w, h, 14); g.fill();
  g.fillStyle = "#6ee7ff"; g.textAlign = "left"; g.font = `bold 16px ${FONT}`; g.fillText(`Bước ${si + 1}/${n}  ·  ${s.tieu_de || ""}`, x + 16, y + 24);
  const lines = wrap(s.mo_ta, w - 32, `16px ${FONT}`).slice(0, 3); g.fillStyle = "#fff"; g.font = `16px ${FONT}`; lines.forEach((l, i) => g.fillText(l, x + 16, y + 48 + i * 21));
  g.fillStyle = "rgba(255,255,255,.14)"; g.fillRect(x, y + h + 2, w, 3); g.fillStyle = "#6ee7ff"; g.fillRect(x, y + h + 2, w * clamp((t - s.t0) / (s.t1 - s.t0)), 3);
}
function drawEnd(t) {
  const p = clamp((t - endCard.t0) / 1.2); g.fillStyle = `rgba(8,12,40,${.98 * p})`; g.fillRect(0, 44, W, H - 44);
  g.globalAlpha = p; g.textAlign = "center"; g.fillStyle = "#6ee7ff"; g.font = `bold 20px ${FONT}`; g.fillText("KẾT LUẬN", W / 2, 90);
  g.fillStyle = "#fff"; g.font = `bold 28px ${FONT}`; const eq = wrap(sub(D.phuong_trinh || ""), W - 100, `bold 28px ${FONT}`).slice(0, 2); eq.forEach((l, i) => g.fillText(l, W / 2, 134 + i * 36));
  let y = 134 + eq.length * 36 + 22; g.textAlign = "left"; g.fillStyle = "#9fb3ff"; g.font = `bold 17px ${FONT}`; g.fillText("Hiện tượng quan sát được:", 70, y); y += 28;
  g.fillStyle = "#fff"; g.font = `17px ${FONT}`;
  (D.hien_tuong || []).slice(0, 5).forEach(h => { wrap("• " + h, W - 150, `17px ${FONT}`).slice(0, 2).forEach((l, i) => { g.fillText((i ? "   " : "") + l, 70, y); y += 24; }); y += 4; });
  if (D.giai_thich && y < H - 70) { g.fillStyle = "#9fb3ff"; g.font = `bold 17px ${FONT}`; y += 6; g.fillText("Giải thích:", 70, y); y += 26; g.fillStyle = "#fff"; g.font = `16px ${FONT}`; wrap(D.giai_thich, W - 150, `16px ${FONT}`).slice(0, 3).forEach(l => { g.fillText(l, 70, y); y += 22; }); }
  g.globalAlpha = 1;
}
// ---------- cô giáo làm thí nghiệm ----------
// Cô đứng sau bàn thí nghiệm, tự đi đến từng dụng cụ, tay làm đúng thao tác, mắt nhìn theo thí nghiệm, miệng cử động theo giọng đọc.
let TEACH_ON = true, POSE = null, mouthOpen = 0;
const C_SK = "#f2c8a2", C_SKD = "#d4a07a", C_HAIR = "#2a1c1e", C_COAT = "#f8faff", C_COATD = "#b9c5e4", C_GL = "#7cc3f0", C_GLD = "#3b8fc8";
const ARM1 = 106, ARM2 = 106, SH_Y = 172, BENCH = 398, SH_W = 50;
const MANUAL = { them_long: 1, them_ran: 1, dun_nong: 1, que_dom: 1, thu_khi: 1 };
const JAR = GAS ? { x: 650, top: 232, bot: 372 } : null;   // ống thu khí (khớp với drawGasKit)
const JAR_ST = 775, LIFT_PX = 118;
const lampX = () => lampV >= 0 ? VS[lampV].x : (VS[0] ? VS[0].x : 480);
const rotOf = (u) => Math.atan2(u.x, -u.y);                 // góc xoay để trục "lên" của dụng cụ trùng hướng u
const along = (p, u, d) => ({ x: p.x + u.x * d, y: p.y + u.y * d });
const tiltFor = (len, tipY) => Math.acos(clamp((tipY - 56) / len, .3, .96));   // nghiêng dụng cụ để tay/bầu cao su không chạm thanh tiêu đề
const unit = (x, y) => { const d = Math.hypot(x, y) || 1; return { x: x / d, y: y / d }; };

// các pha của một thao tác (giây): A = đưa tay tới, R = rút tay về; Le = thời lượng thao tác thực sự
function phs(e) {
  const L = e.t1 - e.t0, Le = e.loai === "dun_nong" ? Math.min(L, 3.4) : e.loai === "them_ran" ? Math.min(L, 4.4) : e.loai === "que_dom" ? Math.min(L, 7) : L;
  return { L: Le, A: Math.min(1.1, Le * .32), R: Math.min(1, Le * .3) };
}
function relT(e) { const q = phs(e); return q.A + Math.min(.55, Math.max(.15, (q.L - q.A - q.R) * .4)); }   // lúc buông chất rắn
function dripP(e, t) { const q = phs(e), s = q.A + .15, f = Math.max(s + .3, q.L - q.R - .1); return clamp((t - e.t0 - s) / (f - s)); }
function splintInfo(e, t) {
  const q = phs(e), el = Math.max(0, t - e.t0), L = q.L, A = q.A, R = q.R, jx = JAR ? JAR.x : 650;
  const hitT = A + .45, endT = L - R;
  const lift = ease(clamp(el / A)) * (1 - ease(clamp((el - endT) / R)));
  const mouthY = (JAR ? JAR.bot : 372) + 8 - LIFT_PX * lift;
  const t0p = { x: JAR_ST - 20, y: 318 }, t1p = { x: jx + 4, y: mouthY + 14 };
  const s1 = ease(clamp((el - .25) / Math.max(.3, hitT - .25))), s2 = ease(clamp((el - endT) / (R * .75))), qq = s1 * (1 - s2);
  const tip = { x: lerp(t0p.x, t1p.x, qq), y: lerp(t0p.y, t1p.y, qq) - Math.sin(Math.PI * qq) * 26 };
  return { tip, lift, eh: el - hitT, hit: el >= hitT, vis: el <= L + .05, qq };
}

// ---------- đường đi của cô ----------
function stationX(e) {
  if (e.loai === "que_dom" || e.loai === "thu_khi") return JAR_ST;
  const vi = e.loai === "dun_nong" && lampV >= 0 ? lampV : vIndex(e.dung_cu), v = VS[vi] || { x: 480 };
  let x = v.x - (e.loai === "dun_nong" ? 108 : 125); if (x < 100) x = v.x + (e.loai === "dun_nong" ? 108 : 125);
  return clamp(x, 100, 860);
}
function targetX(e) { return (e.loai === "que_dom" || e.loai === "thu_khi") ? (JAR ? JAR.x : 650) : e.loai === "dun_nong" ? lampX() : (VS[vIndex(e.dung_cu)] || { x: 480 }).x; }
const TASKS = events.filter(e => MANUAL[e.loai]).map(e => ({ e, x: stationX(e), t0: e.t0, t1: e.t1 })).sort((a, b) => a.t0 - b.t0);
const HOME = TASKS.length ? Math.max(-70, TASKS[0].x - 330) : clamp((VS[0] ? VS[0].x : 480) - 125, 100, 860);
const WALK = [];
(function () {
  let px = HOME, pend = 0;
  TASKS.forEach(k => {
    const need = clamp(Math.abs(k.x - px) / 280, .5, 2.2), b = k.t0 - .1; let a = Math.max(pend, b - need);
    if (b - a < .3) a = Math.max(0, b - .3);
    WALK.push({ a, b, x0: px, x1: k.x, task: k }); px = k.x; pend = k.t1;
  });
})();
function teacherPos(t) {
  let x = HOME, walking = false, dist = 0, task = null;
  for (let i = 0; i < WALK.length; i++) {
    const w = WALK[i];
    if (t >= w.b) { x = w.x1; continue; }
    if (t > w.a && Math.abs(w.x1 - w.x0) > 2) { const q = ease(clamp((t - w.a) / (w.b - w.a))); x = lerp(w.x0, w.x1, q); walking = q > .001 && q < .999; dist = Math.abs(x - w.x0); task = w.task; }
    else x = w.x0;
    break;
  }
  return { x, walking, dist, task };
}
function activeTask(t) { let best = null; for (const k of TASKS) if (t >= k.t0 - .05 && t <= k.t1 + .1) best = k; return best; }
function lastTask(t) { let best = null; for (const k of TASKS) if (k.t0 - .6 <= t) best = k; return best || TASKS[0] || null; }
function focusPt(t) {
  let e = null; for (const x of events) if (x.t0 <= t + .05 && (!e || x.t0 >= e.t0)) e = x;
  if (!e) e = events[0]; if (!e) return { x: 480, y: 250 };
  if (e.loai === "que_dom" || e.loai === "thu_khi") return { x: JAR ? JAR.x : 650, y: 300 };
  if (e.loai === "dun_nong") return { x: lampX(), y: 352 };
  const v = VS[vIndex(e.dung_cu)] || VS[0]; return { x: v.x, y: v.base - v.geo.h * .45 };
}

// ---------- tay: tính khớp khuỷu ----------
function ik(sx, sy, hx, hy, outS, cx) {
  const dx = hx - sx, dy = hy - sy, d = Math.hypot(dx, dy) || 1, dd = Math.min(d, ARM1 + ARM2 - 1.5), ux = dx / d, uy = dy / d;
  const a = (ARM1 * ARM1 - ARM2 * ARM2 + dd * dd) / (2 * dd), h = Math.sqrt(Math.max(0, ARM1 * ARM1 - a * a));
  const bx = sx + ux * a, by = sy + uy * a, c1 = { x: bx - uy * h, y: by + ux * h }, c2 = { x: bx + uy * h, y: by - ux * h };
  const score = c => (c.y - (sy + hy) / 2) + .7 * outS * (c.x - (sx + hx) / 2);   // khuỷu hướng xuống và ra ngoài
  const el = score(c1) >= score(c2) ? c1 : c2;
  return { sx, sy, ex: el.x, ey: el.y, hx: sx + ux * dd, hy: sy + uy * dd };
}

// ---------- tư thế tại thời điểm t ----------
function teacherPose(t, ta) {
  const P = teacherPos(t), k = activeTask(t), lk = k || P.task || lastTask(t);
  const cx0 = P.x, dir = lk ? (targetX(lk.e) >= lk.x ? 1 : -1) : 1;
  const ph = P.dist / 24, sw = P.walking ? Math.sin(ph) : 0, bob = P.walking ? Math.abs(Math.sin(ph)) * 3.2 : 0;
  const cx = cx0 + Math.sin(ta * 1.3) * (P.walking ? 0 : 1.2);
  const rest = { n: { x: cx + dir * 64 + sw * 12 * dir, y: 378 - sw * 8 }, f: { x: cx - dir * 64 - sw * 12 * dir, y: 378 + sw * 8 } };
  const out = { cx, dir, walking: P.walking, tools: [], drip: null, busyV: -1, splint: null, hn: rest.n, hf: rest.f, look: null, bob };
  if (k) {
    const e = k.e, q = phs(e), L = q.L, A = q.A, R = q.R, el = clamp(t - e.t0, 0, L + .3);
    const bl = ease(clamp(el / .45)) * (1 - ease(clamp((el - (L - .45)) / .45)));
    const ap = ease(clamp(el / A)), rt = ease(clamp((el - (L - R)) / R)), qq = ap * (1 - rt);
    const arc = (a, b, h) => ({ x: lerp(a.x, b.x, qq), y: lerp(a.y, b.y, qq) - Math.sin(Math.PI * qq) * h });
    if (e.loai === "them_long") {
      const vi = vIndex(e.dung_cu), v = VS[vi], top = v.base - v.geo.h, bx = cx - dir * 30, by = 270;
      const tip0 = { x: bx, y: by - 40 }, tip1 = { x: v.x, y: top - 18 }, tip = arc(tip0, tip1, 30);
      if (qq > .999) tip.y += Math.sin(el * 2.3) * 1.2;
      const al = Math.max(.3, tiltFor(98, tip1.y)), u = unit(-dir * Math.sin(al), -Math.cos(al)), grip = along(tip, u, 42), dp = dripP(e, t);
      const dripping = qq > .98 && el > A + .15 && el < L - R - .05;
      out.tools.push({ k: "bottle", x: bx, y: by, col: e.mau, label: e.chat || "", a: bl });
      out.tools.push({ k: "dropper", tip, u, fill: .95 - .8 * dp, col: e.mau, sq: dripping ? .5 + .5 * Math.cos(t * 1.6 * 6.2832) : 0, a: bl });
      out.hn = { x: lerp(rest.n.x, grip.x, bl), y: lerp(rest.n.y, grip.y, bl) }; out.hf = { x: lerp(rest.f.x, bx, bl), y: lerp(rest.f.y, by + 15, bl) };
      if (dripping) out.drip = { vi, x: tip.x, y: tip.y, c: hex2(e.mau || "#dfe8ff") };
      out.busyV = vi; out.look = { x: tip.x, y: tip.y + 20 };
    } else if (e.loai === "them_ran") {
      const vi = vIndex(e.dung_cu), v = VS[vi], top = v.base - v.geo.h, kind = e.kieu || "vien", tr = relT(e);
      const tip0 = { x: cx + dir * 44, y: 282 }, tip1 = { x: v.x, y: top - 34 }, tip = arc(tip0, tip1, 28);
      const pour = kind === "bot" ? ease(clamp((el - A) / Math.max(.25, tr - A))) * (1 - rt) : 0, al = Math.max(.4, tiltFor(82, tip1.y));
      const u = unit(-dir * (Math.sin(al) + pour * .75), -Math.cos(al) * (1 - pour * .6)), grip = along(tip, u, 44);
      out.tools.push({ k: kind === "bot" ? "spoon" : "tweezers", tip, u, kind, col: e.mau || "#9aa4ad", carry: el < tr, gap: ease(clamp((el - tr) / .3)), a: bl });
      out.hn = { x: lerp(rest.n.x, grip.x, bl), y: lerp(rest.n.y, grip.y, bl) };
      out.busyV = vi; out.look = { x: tip.x, y: tip.y + 14 };
    } else if (e.loai === "dun_nong") {
      const lx = lampX(), wick = { x: lx - 3, y: 366 };
      if (e.bat) {
        const tip0 = { x: cx + dir * 46, y: 330 }, tip = arc(tip0, wick, 22), u = unit(-dir * .82, .57), grip = along(tip, u, 30);
        out.tools.push({ k: "match", tip, u, lit: el > A * .55 && el < A + .9, a: bl });
        out.hn = { x: lerp(rest.n.x, grip.x, bl), y: lerp(rest.n.y, grip.y, bl) }; out.look = { x: tip.x, y: tip.y - 8 };
      } else {
        const c0 = { x: cx + dir * 46, y: 318 }, c1 = { x: lx, y: 357 }, cp = arc(c0, c1, 24), done = el >= A + .2;
        const capPos = done ? c1 : cp, grip = { x: capPos.x - dir * 21, y: capPos.y - 2 };
        if (!done) out.tools.push({ k: "cap", x: capPos.x, y: capPos.y, a: bl });
        out.hn = { x: lerp(rest.n.x, grip.x, bl), y: lerp(rest.n.y, grip.y, bl) }; out.look = { x: capPos.x, y: capPos.y - 6 };
      }
    } else if (e.loai === "que_dom") {
      const si = splintInfo(e, t), jx = JAR ? JAR.x : 650, jm = ((JAR ? JAR.top : 232) + (JAR ? JAR.bot : 372)) / 2 - LIFT_PX * si.lift;
      const u = unit(-dir * .82, .57), grip = along(si.tip, u, 58), gj = { x: jx + 37, y: jm + 12 };
      out.tools.push({ k: "splint", tip: si.tip, u, a: bl });
      out.hn = { x: lerp(rest.n.x, gj.x, bl), y: lerp(rest.n.y, gj.y, bl) }; out.hf = { x: lerp(rest.f.x, grip.x, bl), y: lerp(rest.f.y, grip.y, bl) };
      out.splint = { tip: si.tip, kq: e.ket_qua || "bung_chay", eh: si.eh, a: bl }; out.look = { x: si.tip.x, y: si.tip.y };
    } else if (e.loai === "thu_khi") {
      const gj = { x: (JAR ? JAR.x : 650) + 37, y: 296 };
      out.hn = { x: lerp(rest.n.x, gj.x, bl), y: lerp(rest.n.y, gj.y, bl) }; out.look = { x: gj.x - 30, y: 270 };
    }
  }
  // ánh nhìn: vào đầu dụng cụ đang cầm; nếu không làm gì thì nhìn vào chỗ đang xảy ra hiện tượng
  if (!out.look) { const tg = P.walking && P.task ? { x: targetX(P.task.e), y: 280 } : focusPt(t); out.look = tg; }
  const hy = Math.max(out.hn.y, out.hf.y), lean = 22 * clamp((hy - 380) / 22, 0, 1);
  out.sy = SH_Y + lean + Math.sin(ta * 1.8) * 1 + bob;
  const aN = ik(cx + dir * SH_W, out.sy, out.hn.x, out.hn.y, dir, cx), aF = ik(cx - dir * SH_W, out.sy, out.hf.x, out.hf.y, -dir, cx);
  out.aN = aN; out.aF = aF;
  return out;
}

// ---------- vẽ cô giáo (phần thân, sau dụng cụ thủy tinh) ----------
function drawSleeve(a) {
  g.lineCap = "round"; g.lineJoin = "round";
  [[C_COATD, 26], [C_COAT, 22]].forEach(([c, w]) => { g.strokeStyle = c; g.lineWidth = w; g.beginPath(); g.moveTo(a.sx, a.sy); g.lineTo(a.ex, a.ey); g.lineTo(a.hx, a.hy); g.stroke(); });
  g.strokeStyle = "#dbe3f6"; g.lineWidth = 3; g.beginPath(); const ux = a.hx - a.ex, uy = a.hy - a.ey, d = Math.hypot(ux, uy) || 1;   // viền tay áo
  g.moveTo(a.hx - ux / d * 9 - uy / d * 10, a.hy - uy / d * 9 + ux / d * 10); g.lineTo(a.hx - ux / d * 9 + uy / d * 10, a.hy - uy / d * 9 - ux / d * 10); g.stroke();
}
function drawTeacherBack(P, ta) {
  const cx = P.cx, sy = P.sy, blink = (ta * 1000) % 3400 < 130;
  g.save(); g.beginPath(); g.rect(0, 0, W, BENCH); g.clip();
  // thân áo blouse
  const gr = g.createLinearGradient(cx - 64, 0, cx + 64, 0); gr.addColorStop(0, "#e5ebfa"); gr.addColorStop(.5, "#ffffff"); gr.addColorStop(1, "#e0e7f7");
  g.fillStyle = gr; g.strokeStyle = C_COATD; g.lineWidth = 2.5; g.beginPath();
  g.moveTo(cx - 22, sy - 16); g.bezierCurveTo(cx - 44, sy - 12, cx - 66, sy - 6, cx - 66, sy + 26); g.lineTo(cx - 70, BENCH + 6); g.lineTo(cx + 70, BENCH + 6); g.lineTo(cx + 66, sy + 26);
  g.bezierCurveTo(cx + 66, sy - 6, cx + 44, sy - 12, cx + 22, sy - 16); g.closePath(); g.fill(); g.stroke();
  g.fillStyle = "#77a9e6"; g.beginPath(); g.moveTo(cx - 21, sy - 15); g.lineTo(cx, sy + 44); g.lineTo(cx + 21, sy - 15); g.closePath(); g.fill();      // áo trong
  g.fillStyle = C_SK; g.fillRect(cx - 10, sy - 38, 20, 26); g.fillStyle = "rgba(0,0,0,.10)"; g.fillRect(cx - 10, sy - 18, 20, 5);                    // cổ
  g.fillStyle = C_SK; g.beginPath(); g.moveTo(cx - 12, sy - 14); g.lineTo(cx, sy + 14); g.lineTo(cx + 12, sy - 14); g.closePath(); g.fill();
  g.strokeStyle = C_COATD; g.lineWidth = 2.2; g.beginPath(); g.moveTo(cx - 24, sy - 15); g.lineTo(cx - 4, sy + 52); g.moveTo(cx + 24, sy - 15); g.lineTo(cx + 4, sy + 52); g.moveTo(cx, sy + 52); g.lineTo(cx, BENCH + 6); g.stroke();
  g.fillStyle = "#c9d3ec"; for (let i = 0; i < 3; i++) { g.beginPath(); g.arc(cx + 3, sy + 90 + i * 52, 3, 0, 7); g.fill(); }
  g.strokeStyle = C_COATD; g.lineWidth = 2; rrect(cx + 24, sy + 100, 32, 34, 5); g.stroke(); g.fillStyle = "#3b6cf6"; g.fillRect(cx + 33, sy + 94, 3.5, 16); g.fillStyle = "#e0457b"; g.fillRect(cx + 40, sy + 96, 3.5, 14);   // túi và bút
  // tay áo
  drawSleeve(P.aF); drawSleeve(P.aN);
  // đầu
  const lk = P.look, hx0 = cx, hy0 = sy - 60 + Math.sin(ta * 7) * mouthOpen * 1.3, ddx = lk.x - hx0, ddy = lk.y - hy0;
  const turn = clamp(ddx / 90, -1, 1), gx = clamp(ddx / 110, -1, 1), gy = clamp(ddy / 130, -1, 1), tilt = clamp(ddx / 700, -.08, .08) - clamp(ddy / 2500, -.03, .03) * 0;
  const hx = hx0 + turn * 3, hy = hy0;
  g.save(); g.translate(hx, hy + 34); g.rotate(tilt); g.translate(-hx, -(hy + 34));
  g.fillStyle = C_HAIR; g.beginPath(); g.ellipse(hx, hy - 4, 37, 39, 0, 0, 7); g.fill();                                                               // tóc sau
  g.beginPath(); g.arc(hx + turn * 2, hy - 42, 15, 0, 7); g.fill(); g.strokeStyle = "#e0457b"; g.lineWidth = 3; g.beginPath(); g.arc(hx + turn * 2, hy - 42, 15, 3.5, 5.9); g.stroke();   // búi tóc và dây buộc
  g.fillStyle = C_SK; g.beginPath(); g.arc(hx - 32, hy + 4, 7, 0, 7); g.arc(hx + 32, hy + 4, 7, 0, 7); g.fill();                                         // tai
  g.beginPath(); g.ellipse(hx, hy + 3, 31, 34, 0, 0, 7); g.fill();                                                                                       // mặt
  g.fillStyle = C_HAIR; g.beginPath(); g.moveTo(hx - 33, hy - 2); g.bezierCurveTo(hx - 37, hy - 42, hx + 37, hy - 42, hx + 33, hy - 2);
  g.bezierCurveTo(hx + 24, hy - 24, hx - 4, hy - 27, hx - 33, hy - 2); g.fill();                                                                         // tóc mái
  const ey = hy - 1, fx = turn * 4;
  g.strokeStyle = "#1e3a8a"; g.lineWidth = 6; g.lineCap = "butt"; g.beginPath(); g.moveTo(hx - 34, ey); g.lineTo(hx + 34, ey); g.stroke();             // dây kính
  for (const s of [-1, 1]) {                                                                                                                              // mắt
    const ex = hx + s * 14 + fx; g.fillStyle = "#fff"; g.beginPath(); g.ellipse(ex, ey, 7.4, blink ? .9 : 6.4, 0, 0, 7); g.fill();
    if (!blink) { g.fillStyle = "#2b2233"; g.beginPath(); g.arc(ex + gx * 3.2, ey + gy * 2.4, 3.5, 0, 7); g.fill(); g.fillStyle = "#fff"; g.beginPath(); g.arc(ex + gx * 3.2 + 1.2, ey + gy * 2.4 - 1.2, 1.1, 0, 7); g.fill(); }
    g.strokeStyle = C_HAIR; g.lineWidth = 2.4; g.lineCap = "round"; g.beginPath(); g.moveTo(ex - 7, ey - 15 - mouthOpen * 1.5); g.quadraticCurveTo(ex, ey - 19 - mouthOpen * 1.5, ex + 7, ey - 15.5 - mouthOpen); g.stroke();   // lông mày
  }
  g.fillStyle = "rgba(200,235,255,.28)"; g.strokeStyle = "#2554c7"; g.lineWidth = 3; rrect(hx + fx - 30, ey - 12.5, 29, 25, 9); g.fill(); g.stroke(); rrect(hx + fx + 1, ey - 12.5, 29, 25, 9); g.fill(); g.stroke();   // kính bảo hộ
  g.strokeStyle = "rgba(255,255,255,.8)"; g.lineWidth = 2; g.beginPath(); g.moveTo(hx + fx - 25, ey - 6); g.lineTo(hx + fx - 17, ey - 9); g.moveTo(hx + fx + 6, ey - 6); g.lineTo(hx + fx + 14, ey - 9); g.stroke();
  g.fillStyle = "rgba(240,120,130,.28)"; g.beginPath(); g.ellipse(hx + fx - 21, ey + 17, 6, 3.4, 0, 0, 7); g.ellipse(hx + fx + 21, ey + 17, 6, 3.4, 0, 0, 7); g.fill();   // má
  g.strokeStyle = C_SKD; g.lineWidth = 1.8; g.beginPath(); g.moveTo(hx + fx * .8 - 2, ey + 12); g.quadraticCurveTo(hx + fx * .8 + 2, ey + 17, hx + fx * .8 - 1, ey + 18); g.stroke();   // mũi
  const mx = hx + fx * .9, my = ey + 27, o = mouthOpen;                                                                                                  // miệng
  if (o > .06) { g.fillStyle = "#7d2433"; g.beginPath(); g.ellipse(mx, my + o * 2.5, 6.5 + o * 2.2, 1.6 + o * 7.5, 0, 0, 7); g.fill(); g.strokeStyle = "#c4545e"; g.lineWidth = 1.6; g.stroke(); if (o > .35) { g.fillStyle = "#f6f0f0"; g.fillRect(mx - 4.5, my - o * 2.6 + 1.6, 9, 2.4); } }
  else { g.strokeStyle = "#c4545e"; g.lineWidth = 2.2; g.lineCap = "round"; g.beginPath(); g.moveTo(mx - 7, my - 1); g.quadraticCurveTo(mx, my + 4, mx + 7, my - 1); g.stroke(); }
  g.restore(); g.restore();
}

// ---------- dụng cụ cầm trên tay (vẽ phía trước kính) ----------
function drawBottle(tl) {
  const { x, y } = tl; g.save(); g.globalAlpha = tl.a; const c = mixc(hex2(tl.col || "#dfe8ff"), [110, 140, 200], .22);
  g.fillStyle = "rgba(225,235,250,.6)"; g.strokeStyle = "rgba(110,140,195,.95)"; g.lineWidth = 2; rrect(x - 18, y - 22, 36, 48, 8); g.fill(); g.stroke();
  g.save(); rrect(x - 18, y - 22, 36, 48, 8); g.clip(); g.fillStyle = rgba(c, .92); g.fillRect(x - 20, y - 8, 40, 36); g.restore();
  g.fillStyle = "rgba(225,235,250,.75)"; g.fillRect(x - 8, y - 34, 16, 13); g.strokeRect(x - 8, y - 34, 16, 13);
  g.fillStyle = "#fff"; rrect(x - 15, y - 18, 30, 17, 3); g.fill(); g.fillStyle = "#1b2150"; g.font = `bold 8px ${FONT}`; g.textAlign = "center"; g.fillText(sub(tl.label).slice(0, 9), x, y - 6);
  g.restore();
}
function drawDropper(tl) {
  g.save(); g.globalAlpha = tl.a; g.translate(tl.tip.x, tl.tip.y); g.rotate(rotOf(tl.u));
  const tube = () => { g.beginPath(); g.moveTo(-1.6, 0); g.lineTo(-4.6, -15); g.lineTo(-4.6, -64); g.lineTo(4.6, -64); g.lineTo(4.6, -15); g.lineTo(1.6, 0); g.closePath(); };
  tube(); g.fillStyle = "rgba(238,245,255,.6)"; g.fill();
  if (tl.fill > .02) { g.save(); tube(); g.clip(); g.fillStyle = rgba(mixc(hex2(tl.col || "#dfe8ff"), [110, 140, 200], .25), .95); g.fillRect(-6, -2 - tl.fill * 60, 12, tl.fill * 60 + 2); g.restore(); }
  tube(); g.strokeStyle = "rgba(100,130,185,.95)"; g.lineWidth = 1.6; g.stroke();
  g.fillStyle = "#222c4f"; g.fillRect(-6.6, -69, 13.2, 6);
  g.fillStyle = "#c8403a"; g.strokeStyle = "#8f2a27"; g.lineWidth = 1.5; g.beginPath(); g.ellipse(0, -82, 8.4 - tl.sq * 2, 13.5 - tl.sq * 3.6, 0, 0, 7); g.fill(); g.stroke();
  g.fillStyle = "rgba(255,255,255,.38)"; g.beginPath(); g.ellipse(-2.8, -86, 1.8, 5, 0, 0, 7); g.fill();
  g.restore();
}
function drawTweezers(tl) {
  g.save(); g.globalAlpha = tl.a; g.translate(tl.tip.x, tl.tip.y); g.rotate(rotOf(tl.u)); const op = tl.gap * 5;
  g.lineCap = "round"; g.strokeStyle = "#aeb8c8"; g.lineWidth = 3.2; g.beginPath(); g.moveTo(-1.5 - op, 0); g.lineTo(-6.5, -74); g.moveTo(1.5 + op, 0); g.lineTo(6.5, -74); g.stroke();
  g.strokeStyle = "#7d889b"; g.lineWidth = 1; g.beginPath(); g.moveTo(-1.5 - op, 0); g.lineTo(-6.5, -74); g.moveTo(1.5 + op, 0); g.lineTo(6.5, -74); g.stroke();
  if (tl.carry) { const c = hex2(tl.col); g.fillStyle = rgba(c, 1); g.strokeStyle = rgba(mixc(c, [0, 0, 0], .35), 1); g.lineWidth = 1.3;
    if (tl.kind === "dinh") { rrect(-2.6, 1, 5.2, 46, 2); g.fill(); g.stroke(); g.beginPath(); g.ellipse(0, 1, 5.6, 2.2, 0, 0, 7); g.fill(); g.stroke(); }
    else if (tl.kind === "mieng") { rrect(-10, 1, 20, 5, 1.5); g.fill(); g.stroke(); }
    else { g.beginPath(); g.ellipse(0, 3, 8, 5.2, 0, 0, 7); g.fill(); g.stroke(); } }
  g.restore();
}
function drawSpoon(tl) {
  g.save(); g.globalAlpha = tl.a; g.translate(tl.tip.x, tl.tip.y); g.rotate(rotOf(tl.u));
  g.lineCap = "round"; g.strokeStyle = "#a9b3c4"; g.lineWidth = 3.4; g.beginPath(); g.moveTo(0, -4); g.lineTo(0, -76); g.stroke();
  g.fillStyle = "#c6cedb"; g.strokeStyle = "#7d889b"; g.lineWidth = 1.2; g.beginPath(); g.ellipse(0, -7, 5.6, 10, 0, 0, 7); g.fill(); g.stroke();
  if (tl.carry) { const c = hex2(tl.col); g.fillStyle = rgba(c, 1); g.beginPath(); g.ellipse(0, -7, 4.4, 8, 0, 0, 7); g.fill(); }
  g.restore();
}
function flameAt(x, y, s, t, blue) {
  const sw = Math.sin(t * 11) * 2 * s, hh = (36 + Math.sin(t * 17) * 3) * s;
  const gl = g.createRadialGradient(x, y - 14 * s, 3, x, y - 14 * s, 46 * s); gl.addColorStop(0, blue ? "rgba(120,170,255,.4)" : "rgba(255,170,60,.38)"); gl.addColorStop(1, "rgba(255,170,60,0)"); g.fillStyle = gl; g.fillRect(x - 50 * s, y - 60 * s, 100 * s, 100 * s);
  const fl = (col, k) => { g.fillStyle = col; g.beginPath(); g.moveTo(x, y); g.bezierCurveTo(x - 13 * k * s, y - 12 * s, x - 8 * k * s, y - hh * k * .6, x + sw, y - hh * k); g.bezierCurveTo(x + 8 * k * s, y - hh * k * .6, x + 13 * k * s, y - 12 * s, x, y); g.fill(); };
  if (blue) { fl("rgba(110,160,255,.7)", 1); fl("rgba(190,215,255,.9)", .6); } else { fl("rgba(255,120,20,.78)", 1); fl("rgba(255,210,70,.92)", .68); fl("rgba(255,250,200,.9)", .32); }
}
function drawMatch(tl, ta) {
  g.save(); g.globalAlpha = tl.a; g.translate(tl.tip.x, tl.tip.y); g.rotate(rotOf(tl.u));
  g.fillStyle = "#d9b27a"; g.fillRect(-1.8, -48, 3.6, 48); g.fillStyle = tl.lit ? "#2f2420" : "#b13a22"; g.beginPath(); g.arc(0, 0, 3.8, 0, 7); g.fill(); g.restore();
  if (tl.lit) { g.save(); g.globalAlpha = tl.a; flameAt(tl.tip.x, tl.tip.y - 2, .5, ta, false); g.restore(); }
}
function capShape(x, y) { g.beginPath(); g.moveTo(x - 12, y + 16); g.lineTo(x - 12, y + 2); g.quadraticCurveTo(x - 12, y - 12, x, y - 12); g.quadraticCurveTo(x + 12, y - 12, x + 12, y + 2); g.lineTo(x + 12, y + 16); g.closePath(); }
function drawCap(tl) {
  g.save(); g.globalAlpha = tl.a; capShape(tl.x, tl.y); g.fillStyle = "rgba(205,228,250,.62)"; g.fill(); g.strokeStyle = "rgba(140,172,215,.98)"; g.lineWidth = 2; g.stroke();
  g.fillStyle = "#cfd8e8"; g.beginPath(); g.arc(tl.x, tl.y - 13, 3.4, 0, 7); g.fill(); g.restore();
}
function drawSplintStick(tl) {
  g.save(); g.globalAlpha = tl.a; g.translate(tl.tip.x, tl.tip.y); g.rotate(rotOf(tl.u));
  g.fillStyle = "#d2a86c"; g.fillRect(-2.6, -98, 5.2, 98); g.fillStyle = "#3a2b20"; g.fillRect(-2.6, -12, 5.2, 12); g.restore();
}
function glove(x, y, dir, grip) {
  g.fillStyle = C_GL; g.strokeStyle = C_GLD; g.lineWidth = 2;
  g.beginPath(); g.ellipse(x, y, 11.5, 10.5, 0, 0, 7); g.fill(); g.stroke();
  g.beginPath(); g.ellipse(x + dir * 9, y - 6, 4.6, 6.2, dir * .5, 0, 7); g.fill(); g.stroke();                       // ngón cái
  if (grip) { g.strokeStyle = "rgba(60,140,200,.7)"; g.lineWidth = 1.5; g.beginPath(); g.moveTo(x - 7, y - 3); g.lineTo(x + 7, y - 3); g.moveTo(x - 7, y + 2); g.lineTo(x + 7, y + 2); g.stroke(); }
}
function drawTeacherFront(P, ta) {
  g.save(); g.beginPath(); g.rect(0, 0, W, BENCH); g.clip();
  P.tools.forEach(tl => { if (tl.a < .02) return; ({ bottle: drawBottle, dropper: drawDropper, tweezers: drawTweezers, spoon: drawSpoon, match: m => drawMatch(m, ta), cap: drawCap, splint: drawSplintStick })[tl.k](tl); });
  const holding = P.tools.some(t => t.a > .3);
  glove(P.aF.hx, P.aF.hy, -P.dir, holding); glove(P.aN.hx, P.aN.hy, P.dir, holding);
  g.restore();
}
// hiệu ứng ở đầu que đóm (tàn đỏ, bùng cháy, nổ nhỏ, tắt)
function drawSplintFx(S, t) {
  const x = S.tip.x, y = S.tip.y, kq = S.kq, eh = S.eh; g.save(); g.globalAlpha = S.a;
  const ember = () => { const r = 9 + Math.sin(t * 9) * 1.2, gr = g.createRadialGradient(x, y, 1, x, y, r); gr.addColorStop(0, "rgba(255,110,40,.95)"); gr.addColorStop(1, "rgba(255,70,20,0)"); g.fillStyle = gr; g.beginPath(); g.arc(x, y, r, 0, 7); g.fill(); g.fillStyle = "#ff7a3a"; g.beginPath(); g.arc(x, y, 2.6, 0, 7); g.fill(); };
  if (kq === "bung_chay") {
    if (eh < .05) ember(); else { const s = .8 + clamp((eh - .05) / .8) * .9, r = 14 + clamp((eh - .05) / .8) * 34, gr = g.createRadialGradient(x, y, 2, x, y, r); gr.addColorStop(0, "rgba(255,255,210,1)"); gr.addColorStop(.4, "rgba(255,190,60,.9)"); gr.addColorStop(1, "rgba(255,120,0,0)"); g.fillStyle = gr; g.beginPath(); g.arc(x, y, r, 0, 7); g.fill(); flameAt(x, y, s, t, false); }
  } else if (kq === "no_nho") {
    if (eh < .12) flameAt(x, y, .7, t, false);
    else { const pp = clamp((eh - .12) / .45);
      if (pp < 1) { g.strokeStyle = `rgba(255,230,120,${1 - pp})`; g.lineWidth = 4; g.beginPath(); g.arc(x, y - 10, pp * 56, 0, 7); g.stroke(); g.fillStyle = `rgba(255,243,190,${1 - pp * .6})`; g.font = `bold 22px ${FONT}`; g.textAlign = "center"; g.fillText("póp!", x, y - 14 - pp * 56); }
      flameAt(x, y, .7, t, true); }
  } else { if (eh < .1) flameAt(x, y, .7, t, false); else { const k2 = clamp((eh - .1) / .5); g.globalAlpha = S.a * (1 - k2); flameAt(x, y, .7 * (1 - k2 * .7), t, false); g.globalAlpha = S.a; if (k2 > .1) { g.fillStyle = "rgba(200,205,220,.5)"; for (let i = 0; i < 4; i++) { g.beginPath(); g.arc(x + Math.sin(t * 3 + i) * 5, y - 12 - ((t * 24 + i * 11) % 36), 2.5 + i * .6, 0, 7); g.fill(); } } ember(); } }
  g.restore();
}
let an = null, anBuf = null;
function updMouth(now) {   // miệng cử động theo biên độ giọng đọc (giọng AI), hoặc nhịp giả khi trình duyệt tự đọc
  if (!playing || !narrOn) { mouthOpen = 0; return; }
  let mt = 0;
  if (narrSrc && an) { an.getByteTimeDomainData(anBuf); let s = 0; for (let i = 0; i < anBuf.length; i++) { const v = (anBuf[i] - 128) / 128; s += v * v; } mt = clamp(Math.sqrt(s / anBuf.length) * 6.5); }
  else if (pending > 0) mt = .25 + .55 * Math.abs(Math.sin(now / 105 + Math.sin(now / 410) * 2));
  mouthOpen += (mt - mouthOpen) * (mt > mouthOpen ? .6 : .35); if (mouthOpen < .03) mouthOpen = 0;
}
let curT = 0;
function render(t, ta) {
  ta = ta === undefined ? t : ta; curT = t; const st = stateAt(t); drawBench();
  POSE = TEACH_ON ? teacherPose(t, ta) : null; if (POSE) drawTeacherBack(POSE, ta);
  if (lampV >= 0) { const v = VS[lampV]; drawStand(v); drawLamp(v.x, st.lamp, ta, st.capOn); }
  VS.forEach((v, i) => { drawGlass(v, st.vs[i], ta); });
  drawGasKit(st, ta); drawSplint(st, ta);
  VS.forEach((v, i) => drawFx(v, st.vs[i], ta));
  VS.forEach((v, i) => { g.textAlign = "center"; g.fillStyle = "#3b2a14"; g.font = `bold 13px ${FONT}`; g.fillText(v.nhan || v.id, v.x, 416);
    const its = st.vs[i].items.slice(0, 2).join(" + ");
    if (its && !(POSE && POSE.busyV === i)) { const s = sub(its).slice(0, 38), yy = v.base - v.geo.h - (i === lampV ? 6 : 10) - (!POSE && st.vs[i].fx.length ? 76 : 0); g.font = `12.5px ${FONT}`;
      if (POSE) { const w = g.measureText(s).width + 16; g.fillStyle = "rgba(8,12,40,.72)"; rrect(v.x - w / 2, yy - 14, w, 21, 8); g.fill(); }
      g.fillStyle = "rgba(255,255,255,.95)"; g.fillText(s, v.x, yy); } });
  if (POSE) drawTeacherFront(POSE, ta);
  if (st.lamp) { g.fillStyle = "rgba(255,200,120,.95)"; g.font = `bold 13px ${FONT}`; g.textAlign = "left"; g.fillText("🔥 Đang đun nóng", 18, 66); }
  drawHud(t); if (steps[curStep(t)].end) drawEnd(t);
}

// ---------- điều khiển, âm thanh, giọng đọc ----------
let playing = false, last = 0, speed = 1, recording = null, AMB = 0;
let AC = null, master = null, sfxBus = null, narBus = null, rdest = null, noiseBuf = null, flame = null, sfxOn = true, narrOn = true;
const bufs = {};
function initAudio() {
  if (AC) { if (AC.state === "suspended") AC.resume(); return; }
  try {
    AC = new (window.AudioContext || window.webkitAudioContext)();
    master = AC.createGain(); master.gain.value = .95; master.connect(AC.destination);
    try { rdest = AC.createMediaStreamDestination(); master.connect(rdest); } catch (e) { rdest = null; }
    sfxBus = AC.createGain(); sfxBus.gain.value = steps.some(s => s.am) ? .45 : 1; sfxBus.connect(master);   // hiệu ứng nhỏ lại khi có giọng đọc
    narBus = AC.createGain(); narBus.gain.value = 1; narBus.connect(master);
    try { an = AC.createAnalyser(); an.fftSize = 512; anBuf = new Uint8Array(an.fftSize); narBus.connect(an); } catch (e) { an = null; }
    noiseBuf = AC.createBuffer(1, AC.sampleRate * 2, AC.sampleRate); const d = noiseBuf.getChannelData(0); for (let i = 0; i < d.length; i++) d[i] = Math.random() * 2 - 1;
    const src = AC.createBufferSource(); src.buffer = noiseBuf; src.loop = true;
    const lp = AC.createBiquadFilter(); lp.type = "lowpass"; lp.frequency.value = 520; flame = AC.createGain(); flame.gain.value = 0;
    src.connect(lp); lp.connect(flame); flame.connect(sfxBus); src.start();
    steps.forEach((s, i) => { if (s.am) fetch(s.am).then(r => r.arrayBuffer()).then(ab => new Promise((ok, no) => AC.decodeAudioData(ab, ok, no))).then(b => { bufs[i] = b; if (playing && curStep(curT) === i) playNarr(i); }).catch(() => { }); });
  } catch (e) { AC = null; }
}
function blip(f0, f1, dur, vol) {
  if (!AC || !sfxOn) return; const o = AC.createOscillator(), gn = AC.createGain(), t = AC.currentTime;
  o.type = "sine"; o.frequency.setValueAtTime(f0, t); o.frequency.exponentialRampToValueAtTime(Math.max(30, f1), t + dur);
  gn.gain.setValueAtTime(Math.min(1, vol * 4), t); gn.gain.exponentialRampToValueAtTime(.0001, t + dur); o.connect(gn); gn.connect(sfxBus); o.start(t); o.stop(t + dur + .02);
}
function burst(dur, vol, type, freq) {
  if (!AC || !sfxOn) return; const s = AC.createBufferSource(), f = AC.createBiquadFilter(), gn = AC.createGain(), t = AC.currentTime;
  s.buffer = noiseBuf; f.type = type; f.frequency.value = freq; gn.gain.setValueAtTime(Math.min(1, vol * 4), t); gn.gain.exponentialRampToValueAtTime(.0001, t + dur);
  s.connect(f); f.connect(gn); gn.connect(sfxBus); s.start(t, Math.random()); s.stop(t + dur + .02);
}
const fired = new Set(); const acc = { b: 0, d: 0, g: 0 }; let prevStep = -1;
function audioTick(t, st, rdt) {
  if (!AC) return;
  if (flame) flame.gain.setTargetAtTime(sfxOn && st.lamp ? .16 : 0, AC.currentTime, .15);
  if (!sfxOn) return;
  let rate = 0; st.vs.forEach(v => { rate = Math.max(rate, v.rate); });
  acc.b += rdt * rate * 12; while (acc.b >= 1) { acc.b -= 1; const f = 280 + Math.random() * 520; blip(f, f * 1.9, .06 + Math.random() * .04, .035 + Math.random() * .03); }
  if (st.gasOn) { acc.g += rdt * 7; while (acc.g >= 1) { acc.g -= 1; const f = 420 + Math.random() * 360; blip(f, f * 1.7, .05, .03); } }
  let drip = false; st.vs.forEach(v => v.fx.forEach(f => { if (f.t === "long") drip = true; }));
  if (drip) { acc.d += rdt * 4.5; while (acc.d >= 1) { acc.d -= 1; blip(1000 + Math.random() * 500, 520, .08, .05); } }
  st.vs.forEach(v => v.fx.forEach(f => {
    if (f.t === "ran" && f.p > .5 && !fired.has("r" + f.id)) { fired.add("r" + f.id); burst(.14, .16, "lowpass", 260); blip(190, 90, .12, .12); }
    if (f.t === "tua" && f.p > .02 && !fired.has("k" + f.id)) { fired.add("k" + f.id); burst(.7, .035, "bandpass", 2200); }
  }));
  if (st.splint) { const s = st.splint;
    if (s.kq === "bung_chay" && (TEACH_ON ? s.eh > .06 : s.p > .5) && !fired.has("q" + s.id)) { fired.add("q" + s.id); burst(.6, .12, "bandpass", 1400); }
    if (s.kq === "no_nho" && (TEACH_ON ? s.eh > .14 : s.p > .78) && !fired.has("q" + s.id)) { fired.add("q" + s.id); burst(.09, .4, "lowpass", 1800); blip(220, 80, .1, .25); } }
  const si = curStep(t); if (si !== prevStep) { if (prevStep >= 0 && si > prevStep) { blip(660, 990, .22, .05); if (steps[si].end) setTimeout(() => blip(990, 1320, .3, .05), 180); } prevStep = si; }
}
// --- giọng đọc: ưu tiên tệp giọng AI đã thu sẵn; nếu không có thì dùng giọng của trình duyệt ---
let narrSrc = null, spokenStep = -1, pending = 0, holdFor = 0;
function stopNarr() { if (narrSrc) { try { narrSrc.onended = null; narrSrc.stop(); } catch (e) { } narrSrc = null; } stopSpeech(); }
function playNarr(si, offMovie = 0) {
  stopNarr(); if (!narrOn) return; const s = steps[si];
  if (AC && bufs[si]) {
    const src = AC.createBufferSource(); src.buffer = bufs[si]; src.playbackRate.value = clamp(s.k * speed, .5, 3); src.connect(narBus);
    const start = Math.max(0, offMovie - .25) * s.k;
    if (start < bufs[si].duration - .05) { src.start(0, start); narrSrc = src; }
    return;
  }
  if (s.am && !bufs[si]) return;   // đang giải mã, sẽ tự phát khi xong
  speak(s.end ? endText : s.mo_ta);
}
const HAS_TTS = "speechSynthesis" in window; let viVoice = null; const vsel = document.getElementById("vsel");
function pickVoice() {
  if (!HAS_TTS) return; const vi = speechSynthesis.getVoices().filter(v => /^vi([-_]|$)/i.test(v.lang));
  const score = v => (/natural|neural|online/i.test(v.name) ? 3 : 0) + (/google/i.test(v.name) ? 2 : 0) + (v.localService ? 0 : 1);
  vi.sort((a, b) => score(b) - score(a)); viVoice = vi[0] || null;
  if (vsel) { vsel.innerHTML = vi.map((v, i) => `<option value="${i}">${v.name.replace(/Microsoft |Google /, "")}</option>`).join(""); vsel.style.display = vi.length > 1 ? "" : "none"; vsel.onchange = () => { viVoice = vi[+vsel.value]; }; }
}
if (HAS_TTS) { pickVoice(); speechSynthesis.onvoiceschanged = pickVoice; }
const SAY = [["H2SO4", "a xít sun phu ríc"], ["HCl", "a xít clo hi đríc"], ["Cu(OH)2", "đồng hai hi đrô xít"], ["NaOH", "na tri hi đrô xít"], ["CuSO4", "đồng hai sun phát"], ["FeSO4", "sắt hai sun phát"],
  ["Na2SO4", "na tri sun phát"], ["Na2CO3", "na tri các bô nát"], ["AgNO3", "bạc ni trát"], ["AgCl", "bạc clo rua"], ["NaCl", "na tri clo rua"], ["NaNO3", "na tri ni trát"], ["KMnO4", "ka li pê man ga nát"],
  ["K2MnO4", "ka li măng ga nát"], ["MnO2", "man gan đi ô xít"], ["ZnCl2", "kẽm clo rua"], ["CO2", "khí các bô níc"], ["H2O", "nước"], ["H2", "khí hi đrô"], ["O2", "khí ô xi"], ["Zn", "kẽm"], ["Fe", "sắt"], ["Cu", "đồng"]];
function spoken(text) { let s = String(text || ""); SAY.forEach(([k, v]) => { s = s.split(k).join(v); }); return s.replace(/->|→/g, " tạo thành ").replace(/[↑↓]/g, "").replace(/°C/g, " độ xê").replace(/\bmL\b/g, "mi li lít").replace(/\s+/g, " ").trim(); }
function chunks(s) { const parts = s.split(/(?<=[.!?])\s+/), out = []; let cur = ""; parts.forEach(p => { if ((cur + " " + p).length > 230 && cur) { out.push(cur); cur = p; } else cur = cur ? cur + " " + p : p; }); if (cur) out.push(cur); return out; }
function stopSpeech() { if (HAS_TTS) { try { speechSynthesis.cancel(); } catch (e) { } } pending = 0; }
function speak(text) {
  if (!HAS_TTS || !narrOn) return; stopSpeech();
  chunks(spoken(text)).forEach(p => { const u = new SpeechSynthesisUtterance(p); u.lang = "vi-VN"; if (viVoice) u.voice = viVoice; u.rate = clamp(speed * .95, .6, 2); u.pitch = 1;
    pending++; u.onend = u.onerror = () => { pending = Math.max(0, pending - 1); }; try { speechSynthesis.speak(u); } catch (e) { pending = Math.max(0, pending - 1); } });
}
// --- bộ điều khiển phát ---
const sl = document.getElementById("sl"), tm = document.getElementById("tm"), bPlay = document.getElementById("bPlay"), stl = document.getElementById("st"), spR = document.getElementById("sp"), spL = document.getElementById("spl");
const fmt = (s) => Math.floor(s / 60) + ":" + String(Math.floor(s % 60)).padStart(2, "0");
function setT(t) { curT = clamp(t, 0, TOTAL); render(curT, AMB); sl.value = Math.round(curT / TOTAL * 1000); tm.textContent = fmt(curT) + " / " + fmt(TOTAL); }
function seekTo(t) { fired.clear(); stopNarr(); spokenStep = -1; prevStep = -1; holdFor = 0; AMB = clamp(t, 0, TOTAL); setT(t); if (playing) { spokenStep = curStep(curT); playNarr(spokenStep, curT - steps[spokenStep].t0); } }
function tick(now) {
  const rdt = Math.min(.1, (now - last) / 1000); last = now;
  if (playing) {
    AMB += rdt * speed; let t = curT + rdt * speed; const si0 = curStep(curT);
    // chưa có tệp giọng AI mà trình duyệt vẫn đang đọc: giữ phim ở cuối bước tối đa 6 giây để không cắt giữa câu
    if (t >= steps[si0].t1 && !steps[si0].am && narrOn && pending > 0 && holdFor < 6) { t = steps[si0].t1 - .001; holdFor += rdt; }
    if (t >= TOTAL) { t = TOTAL; playing = false; bPlay.textContent = "↻ Xem lại"; if (recording) finishRec(); }
    updMouth(now); setT(t); const si = curStep(t);
    if (si !== spokenStep && playing) { spokenStep = si; holdFor = 0; playNarr(si, 0); }
    audioTick(t, stateAt(t), rdt);
  }
  requestAnimationFrame(tick);
}
function pauseUI() { playing = false; bPlay.textContent = "▶ Phát"; stopNarr(); spokenStep = -1; mouthOpen = 0; setT(curT); }
bPlay.onclick = () => { initAudio(); if (!playing) { if (curT >= TOTAL - .05) seekTo(0); playing = true; spokenStep = -1; last = performance.now(); bPlay.textContent = "⏸ Dừng"; } else pauseUI(); };
document.getElementById("bR").onclick = () => { pauseUI(); seekTo(0); };
document.getElementById("bN").onclick = () => { const i = curStep(curT); seekTo(steps[Math.min(i + 1, steps.length - 1)].t0 + .01); };
document.getElementById("bP").onclick = () => { const i = curStep(curT); seekTo((curT - steps[i].t0 > 1.5 ? steps[i].t0 : steps[Math.max(i - 1, 0)].t0) + .01); };
sl.oninput = () => { seekTo(sl.value / 1000 * TOTAL); };
spR.oninput = () => { speed = +spR.value; spL.textContent = speed.toFixed(2).replace(".", ",") + "×"; if (playing) { const si = curStep(curT); playNarr(si, curT - steps[si].t0); } };
const bVoice = document.getElementById("bVoice"), bSfx = document.getElementById("bSfx");
bVoice.onclick = () => { narrOn = !narrOn; bVoice.textContent = "🗣️ Giọng đọc: " + (narrOn ? "Bật" : "Tắt"); bVoice.classList.toggle("off", !narrOn); if (!narrOn) stopNarr(); else if (playing) { const si = curStep(curT); playNarr(si, curT - steps[si].t0); } };
const bTeach = document.getElementById("bTeach");
bTeach.onclick = () => { TEACH_ON = !TEACH_ON; bTeach.textContent = "👩‍🏫 Cô giáo: " + (TEACH_ON ? "Bật" : "Tắt"); bTeach.classList.toggle("off", !TEACH_ON); setT(curT); };
bSfx.onclick = () => { sfxOn = !sfxOn; bSfx.textContent = "🔊 Âm thanh: " + (sfxOn ? "Bật" : "Tắt"); bSfx.classList.toggle("off", !sfxOn); initAudio(); };
setTimeout(() => { if (!steps.some(s => s.am) && HAS_TTS && !viVoice && !stl.textContent) stl.textContent = "Máy này chưa có giọng đọc tiếng Việt nên chưa nghe được lời đọc (hiệu ứng âm thanh vẫn có). Hãy cài giọng Tiếng Việt trong phần Giọng nói của hệ điều hành."; }, 2500);
document.getElementById("bHtml").onclick = () => { const b = new Blob(["<!DOCTYPE html>" + document.documentElement.outerHTML], { type: "text/html" }); const a = document.createElement("a"); a.href = URL.createObjectURL(b); a.download = "mo_phong_thi_nghiem.html"; a.click(); };
function finishRec() { try { recording.rec.stop(); } catch (e) { } }
document.getElementById("bRec").onclick = () => {
  if (recording) return;
  if (!window.MediaRecorder || !cv.captureStream) { stl.textContent = "Trình duyệt này chưa hỗ trợ ghi video. Hãy dùng Chrome hoặc Edge, hoặc bấm “Tải .html” để lưu bản mô phỏng."; return; }
  try {
    initAudio();
    const vt = cv.captureStream(30).getVideoTracks(), at = (AC && rdest) ? rdest.stream.getAudioTracks() : [];
    const MP4 = ["video/mp4;codecs=avc1.42E01E,mp4a.40.2", "video/mp4;codecs=avc1,mp4a.40.2", "video/mp4;codecs=avc1.42E01E,opus", "video/mp4;codecs=avc1", "video/mp4;codecs=vp9,opus", "video/mp4"];
    const WEBM = at.length ? ["video/webm;codecs=vp9,opus", "video/webm;codecs=vp8,opus", "video/webm"] : ["video/webm;codecs=vp9", "video/webm;codecs=vp8", "video/webm"];
    let mt = MP4.find(m => MediaRecorder.isTypeSupported(m)) || "", warn = "";
    if (!mt) { mt = WEBM.find(m => MediaRecorder.isTypeSupported(m)) || ""; warn = "Trình duyệt này chưa ghi được .mp4 nên sẽ tải ra .webm. Hãy dùng Chrome hoặc Edge bản mới (từ 126) để tải .mp4. "; }
    const rec = new MediaRecorder(new MediaStream(vt.concat(at)), mt ? { mimeType: mt, videoBitsPerSecond: 4000000 } : {}), chunks = [];
    rec.ondataavailable = e => { if (e.data && e.data.size) chunks.push(e.data); };
    rec.onstop = () => { const blob = new Blob(chunks, { type: rec.mimeType || "video/webm" }); const a = document.createElement("a"); a.href = URL.createObjectURL(blob); a.download = "mo_phong_thi_nghiem." + ((rec.mimeType || "").includes("mp4") ? "mp4" : "webm"); a.click(); recording = null; stl.textContent = steps.some(s => s.am) ? "Đã lưu video (có giọng đọc và hiệu ứng âm thanh)." : "Đã lưu video (có hiệu ứng âm thanh; giọng đọc của trình duyệt không thu được vào video)."; document.getElementById("bRec").textContent = "🎬 Tải video (.mp4)"; };
    recording = { rec }; stl.textContent = warn + "Đang ghi video, vui lòng đợi đến hết phim…"; document.getElementById("bRec").textContent = "⏺ Đang ghi…";
    speed = 1; spR.value = 1; spL.textContent = "1,00×"; seekTo(0); rec.start(); playing = true; spokenStep = -1; last = performance.now(); bPlay.textContent = "⏸ Dừng";
  } catch (e) { stl.textContent = "Không ghi được video: " + e.message; recording = null; }
};
window.__seek = seekTo; window.__total = TOTAL; window.__steps = steps;
setT(0); requestAnimationFrame(tick);
</script></body></html>
'''

SIM_EXAMPLES = ["Kẽm + dung dịch HCl", "CuSO4 + NaOH", "AgNO3 + NaCl", "Nhiệt phân KMnO4 thu khí oxi", "Sắt + dung dịch CuSO4", "Trung hòa NaOH bằng HCl (phenolphtalein)"]
SIM_PROMPT = """Bạn là giáo viên Hóa học trung học phổ thông ở Việt Nam. Hãy lập KỊCH BẢN MÔ PHỎNG THÍ NGHIỆM cho yêu cầu: "__Q__".
Kịch bản này sẽ được máy dựng thành phim hoạt hình nên phải rõ ràng, cụ thể, đúng hiện tượng thật.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"ten":"tên thí nghiệm bằng tiếng Việt",
 "dung_cu":[{"id":"A","loai":"ong_nghiem | coc | binh_tam_giac","nhan":"tên dụng cụ ngắn"}],
 "thu_khi":null hoặc {"tu":"id dụng cụ sinh khí","khi":"công thức khí, ví dụ H2","mau":"#rrggbb"},
 "buoc":[{"tieu_de":"ngắn","mo_ta":"2 câu: thao tác cụ thể (lượng hóa chất, cách làm) và hiện tượng quan sát được","thoi_gian":6,
   "hanh_dong":[ ... ]}],
 "hien_tuong":["các hiện tượng quan sát được, mỗi ý một câu ngắn"],
 "phuong_trinh":"phương trình phản ứng, viết chỉ số dạng số thường, ví dụ Zn + 2HCl -> ZnCl2 + H2↑",
 "chat_tham_gia":[{"he_so":1,"cong_thuc":"Zn","smiles":"[Zn]"}],
 "san_pham":[{"he_so":1,"cong_thuc":"ZnCl2","smiles":"Cl[Zn]Cl"}],
 "giai_thich":"1-2 câu giải thích bản chất",
 "an_toan":["lưu ý an toàn cụ thể"]}
Các hành động hợp lệ trong "hanh_dong" (mỗi hành động là một đối tượng có "loai"):
- them_long: {"loai":"them_long","dung_cu":"A","chat":"tên chất","mau":"#rrggbb màu dung dịch thêm vào","muc":0.05-0.8 (lượng tương đối)}
- them_ran: {"loai":"them_ran","dung_cu":"A","chat":"tên chất","mau":"#rrggbb","kieu":"vien | bot | dinh | mieng"}
- doi_mau: dung dịch đổi sang màu mới {"loai":"doi_mau","dung_cu":"A","mau":"#rrggbb","thoi_gian":3}
- doi_mau_ran: chất rắn đổi màu {"loai":"doi_mau_ran","dung_cu":"A","mau":"#rrggbb","thoi_gian":4}
- ket_tua: xuất hiện kết tủa {"loai":"ket_tua","dung_cu":"A","chat":"công thức","mau":"#rrggbb","luong":0.3-1}
- sui_bot: bọt khí trong dung dịch {"loai":"sui_bot","dung_cu":"A","cuong_do":0-1} (0 là dừng)
- tan: chất rắn nhỏ dần {"loai":"tan","dung_cu":"A","con_lai":0-1,"thoi_gian":6}
- dun_nong: bật hoặc tắt đèn cồn {"loai":"dun_nong","bat":true,"dung_cu":"A"}
- thu_khi: thu khí vào ống thu khí úp trong chậu nước (chỉ khi có "thu_khi") {"loai":"thu_khi","luong":0.3-0.9,"thoi_gian":5}
- que_dom: thử khí bằng que đóm {"loai":"que_dom","ket_qua":"bung_chay | no_nho | tat"}
Có thể thêm "dong_thoi":true để một hành động chạy cùng lúc với hành động trước.
Quy tắc:
- Dùng tối đa 3 dụng cụ, 3-7 bước, mỗi bước 4-8 giây. Mô tả đúng thứ tự thao tác thật trong phòng thí nghiệm.
- Màu sắc theo hiện tượng thật: Cu2+ xanh lam #2f7be0; Fe2+ xanh lục nhạt #bfe6cf; Fe3+ vàng nâu #c98a2b; dung dịch KMnO4 tím #7b1fa2; I2 nâu #8b4513; Cu(OH)2 xanh lam #4f93f2; Fe(OH)3 nâu đỏ #a0522d; kết tủa trắng (AgCl, BaSO4, CaCO3) #fbfbf7; dung dịch không màu #f3f7ff; phenolphtalein trong bazơ hồng #f07db8; quỳ tím: axit đỏ #e0463a, bazơ xanh #3b6cf6.
- Phương trình phải cân bằng; "smiles" của chất vô cơ ion viết dạng như [Na+].[Cl-], khí như [H][H], O=O. Không chắc thì bỏ trống smiles.
- Chỉ mô phỏng thí nghiệm có trong chương trình hoặc an toàn, quen thuộc. Nếu yêu cầu nguy hiểm (gây nổ, chất độc, ma túy, vũ khí) hoặc không phải thí nghiệm hóa học, trả {"tu_choi":"lý do ngắn bằng tiếng Việt"}.
- Không bịa hiện tượng. Nếu thí nghiệm cần điều kiện đặc biệt, nêu rõ trong mo_ta."""

SIM_BLOCK = {"them_long": 2.6, "them_ran": 1.8, "dun_nong": 2.6, "doi_mau": 3, "doi_mau_ran": 3, "ket_tua": 2.8, "tan": 3, "thu_khi": 4, "que_dom": 3.2}
SIM_SAY = [("H2SO4", "a xít sun phu ríc"), ("HCl", "a xít clo hi đríc"), ("Cu(OH)2", "đồng hai hi đrô xít"), ("NaOH", "na tri hi đrô xít"), ("CuSO4", "đồng hai sun phát"), ("FeSO4", "sắt hai sun phát"),
           ("Na2SO4", "na tri sun phát"), ("Na2CO3", "na tri các bô nát"), ("AgNO3", "bạc ni trát"), ("AgCl", "bạc clo rua"), ("NaCl", "na tri clo rua"), ("NaNO3", "na tri ni trát"), ("KMnO4", "ka li pê man ga nát"),
           ("K2MnO4", "ka li măng ga nát"), ("MnO2", "man gan đi ô xít"), ("ZnCl2", "kẽm clo rua"), ("CO2", "khí các bô níc"), ("H2O", "nước"), ("H2", "khí hi đrô"), ("O2", "khí ô xi"), ("Zn", "kẽm"), ("Fe", "sắt"), ("Cu", "đồng")]


def sim_spoken(text):
    s = str(text or "")
    for k, v in SIM_SAY:
        s = s.replace(k, v)
    return re.sub(r"\s+", " ", s.replace("->", " tạo thành ").replace("→", " tạo thành ").replace("↑", "").replace("↓", "").replace("°C", " độ xê").replace("mL", "mi li lít")).strip()


TTS_STYLE = "Hãy đọc to bằng giọng nữ giáo viên Việt Nam, ấm áp, tự nhiên, tốc độ vừa phải, giữ nguyên một giọng và một nhịp đọc đều đặn từ đầu đến cuối, ngắt nghỉ đúng dấu câu: "
_TTS_CACHE = {}


def _tts_pcm(client, text, model, voice):
    """Thu một đoạn bằng Gemini TTS (PCM 16-bit, 24 kHz). Có giới hạn thời gian, thử lại 1 lần, nhớ kết quả để khỏi thu lại."""
    key = (text, model, voice)
    if key in _TTS_CACHE:
        return _TTS_CACHE[key]
    last = None
    for _ in range(2):
        try:
            resp = client.models.generate_content(
                model=model, contents=TTS_STYLE + text,
                config=types.GenerateContentConfig(response_modalities=["AUDIO"], temperature=0, http_options=types.HttpOptions(timeout=40000),
                    speech_config=types.SpeechConfig(voice_config=types.VoiceConfig(prebuilt_voice_config=types.PrebuiltVoiceConfig(voice_name=voice)))))
            pcm = resp.candidates[0].content.parts[0].inline_data.data
            if pcm:
                if len(_TTS_CACHE) > 300:
                    _TTS_CACHE.clear()
                _TTS_CACHE[key] = pcm
                return pcm
        except Exception as e:
            last = e
    raise last or RuntimeError("TTS trống")


def _tts_chunks(text, limit=320):
    """Cắt theo câu: lời đọc bước (tối đa khoảng 320 ký tự) vẫn là một đoạn, chỉ phần kết luận dài mới bị cắt."""
    parts, out, cur = re.split(r"(?<=[.!?;])\s+", text.strip()), [], ""
    for p in parts:
        if cur and len(cur) + len(p) + 1 > limit:
            out.append(cur); cur = p
        else:
            cur = (cur + " " + p).strip()
    if cur:
        out.append(cur)
    return out or [text]


def _pcm_to_wav(pcm):
    import wave
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(24000); w.writeframes(pcm)
    return "data:audio/wav;base64," + base64.b64encode(buf.getvalue()).decode(), len(pcm) / 2 / 24000


def sim_attach_audio(sc):
    """Thu giọng đọc cho từng bước và màn kết luận, thu song song và dừng sau tối đa khoảng 100 giây.
    Đoạn nào lỗi hoặc quá hạn thì bỏ qua (trình duyệt đọc thay). Trả về (kịch bản, số đoạn thu được)."""
    from concurrent.futures import ThreadPoolExecutor, wait
    model, voice = st.secrets.get("GEMINI_TTS_MODEL", "gemini-2.5-flash-preview-tts"), st.secrets.get("GEMINI_TTS_VOICE", "Kore")
    client = get_client()
    out = {**sc, "buoc": [dict(s) for s in sc["buoc"]]}
    out.pop("am_ket", None); out.pop("am_ket_dur", None)
    texts = {}
    for i, s in enumerate(out["buoc"]):
        s.pop("am", None); s.pop("am_dur", None)
        if s.get("mo_ta"):
            texts[i] = sim_spoken(s["mo_ta"])
    texts["end"] = sim_spoken("Kết luận. " + ". ".join(x.rstrip(". ") for x in sc.get("hien_tuong", [])[:5]) + ". " + sc.get("giai_thich", ""))
    pool = ThreadPoolExecutor(max_workers=4)
    futs = {k: [pool.submit(_tts_pcm, client, c, model, voice) for c in _tts_chunks(t)] for k, t in texts.items()}
    wait([f for fl in futs.values() for f in fl], timeout=100)
    pool.shutdown(wait=False, cancel_futures=True)
    gap, ok = b"\x00\x00" * int(24000 * .22), 0
    for k, fl in futs.items():
        try:
            wav, dur = _pcm_to_wav(gap.join(f.result(timeout=0) for f in fl))
        except Exception:
            if k == "end":
                st.toast("Chưa thu được giọng AI cho phần kết luận, phần này sẽ dùng giọng của trình duyệt.", icon="🔈")
            continue
        ok += 1
        if k == "end":
            out["am_ket"], out["am_ket_dur"] = wav, dur
        else:
            out["buoc"][k]["am"], out["buoc"][k]["am_dur"] = wav, dur
    return out, ok


def sim_auto_dur(s):
    """Thời lượng tự động của một bước, giống cách bộ dựng tính (hành động, lời đọc)."""
    tt, last = .35, .35
    for a in s.get("hanh_dong", []):
        d = max(.4, float(a.get("thoi_gian") or SIM_BLOCK.get(a["loai"], .5)))
        off = last if a.get("dong_thoi") else tt
        last = off
        if a["loai"] in SIM_BLOCK and not a.get("dong_thoi"):
            tt = off + d + .15
        elif a["loai"] not in SIM_BLOCK:
            tt = max(tt, off + .4)
    speech = s["am_dur"] + .6 if s.get("am_dur") else len(s.get("mo_ta", "")) / 15 + 1
    return round(max(float(s.get("thoi_gian") or 4), tt + .9, speech) * 2) / 2

_HEXC = re.compile(r"^#[0-9a-fA-F]{6}$")
_SIM_TYPES = {"them_long", "them_ran", "doi_mau", "doi_mau_ran", "ket_tua", "sui_bot", "tan", "dun_nong", "thu_khi", "que_dom"}


def _sim_col(c, default):
    c = str(c or "").strip()
    return c.lower() if _HEXC.match(c) else default


def _sim_num(v, lo, hi, default):
    try:
        return max(lo, min(hi, float(v)))
    except Exception:
        return default


def _sim_species(lst):
    out = []
    for s in (lst or [])[:10]:
        if isinstance(s, dict) and (s.get("cong_thuc") or s.get("smiles")):
            out.append({"he_so": to_int(s.get("he_so")), "cong_thuc": str(s.get("cong_thuc") or "")[:30], "smiles": str(s.get("smiles") or "")[:80]})
    return out


def sim_clean(raw):
    """Chuẩn hóa kịch bản do AI soạn: loại giá trị lạ, giới hạn số lượng, bảo đảm máy dựng được. Trả về (kịch bản, lỗi)."""
    if not isinstance(raw, dict):
        return None, "Chưa lập được kịch bản, hãy mô tả thí nghiệm rõ hơn."
    if raw.get("tu_choi"):
        return None, "Không mô phỏng thí nghiệm này: " + str(raw["tu_choi"])[:200]
    names = {"ong_nghiem": "Ống nghiệm", "coc": "Cốc thủy tinh", "binh_tam_giac": "Bình tam giác"}
    vessels = []
    for v in (raw.get("dung_cu") or [])[:3]:
        if isinstance(v, dict):
            loai = v.get("loai") if v.get("loai") in names else "ong_nghiem"
            vid = re.sub(r"\W", "", str(v.get("id") or chr(65 + len(vessels))))[:3] or chr(65 + len(vessels))
            if vid in [x["id"] for x in vessels]:
                vid = chr(65 + len(vessels))
            vessels.append({"id": vid, "loai": loai, "nhan": str(v.get("nhan") or names[loai])[:24]})
    if not vessels:
        vessels = [{"id": "A", "loai": "ong_nghiem", "nhan": "Ống nghiệm"}]
    ids = [v["id"] for v in vessels]
    kit = raw.get("thu_khi") if isinstance(raw.get("thu_khi"), dict) else None
    if kit and kit.get("tu") not in ids:
        kit = {**kit, "tu": ids[0]}
    steps, nact = [], 0
    for s in (raw.get("buoc") or [])[:8]:
        if not isinstance(s, dict):
            continue
        acts = []
        for a in (s.get("hanh_dong") or [])[:6]:
            if not isinstance(a, dict) or a.get("loai") not in _SIM_TYPES:
                continue
            t = a["loai"]
            e = {"loai": t, "dung_cu": a.get("dung_cu") if a.get("dung_cu") in ids else ids[0], "chat": str(a.get("chat") or "")[:30], "dong_thoi": bool(a.get("dong_thoi"))}
            if a.get("thoi_gian") is not None:
                e["thoi_gian"] = _sim_num(a.get("thoi_gian"), .5, 10, 3)
            if t == "them_long":
                e.update(mau=_sim_col(a.get("mau"), "#f3f7ff"), muc=_sim_num(a.get("muc"), .05, .8, .3))
            elif t == "them_ran":
                e.update(mau=_sim_col(a.get("mau"), "#9aa4ad"), kieu=a.get("kieu") if a.get("kieu") in ("vien", "bot", "dinh", "mieng") else "vien")
            elif t in ("doi_mau", "doi_mau_ran"):
                e["mau"] = _sim_col(a.get("mau"), "#f3f7ff")
            elif t == "ket_tua":
                e.update(mau=_sim_col(a.get("mau"), "#fbfbf7"), luong=_sim_num(a.get("luong"), .2, 1, .7))
            elif t == "sui_bot":
                e["cuong_do"] = _sim_num(a.get("cuong_do"), 0, 1, .6)
            elif t == "tan":
                e["con_lai"] = _sim_num(a.get("con_lai"), 0, 1, .2)
            elif t == "dun_nong":
                e["bat"] = bool(a.get("bat", True))
            elif t == "thu_khi":
                e["luong"] = _sim_num(a.get("luong"), .2, .95, .7)
                if not kit:
                    kit = {"tu": ids[0], "khi": "Khí", "mau": "#dfe6f2"}
            elif t == "que_dom":
                e["ket_qua"] = a.get("ket_qua") if a.get("ket_qua") in ("bung_chay", "no_nho", "tat") else "bung_chay"
            acts.append(e)
        nact += len(acts)
        steps.append({"tieu_de": str(s.get("tieu_de") or "")[:40], "mo_ta": str(s.get("mo_ta") or "")[:320], "thoi_gian": _sim_num(s.get("thoi_gian"), 3, 10, 5), "hanh_dong": acts})
    if not steps or not nact:
        return None, "Chưa lập được kịch bản có hiện tượng quan sát được, hãy mô tả thí nghiệm cụ thể hơn."
    if not kit:  # que đóm chỉ có nghĩa khi có ống thu khí
        for s in steps:
            s["hanh_dong"] = [a for a in s["hanh_dong"] if a["loai"] != "que_dom"]
    else:
        kit = {"tu": kit.get("tu"), "khi": str(kit.get("khi") or "Khí")[:8], "mau": _sim_col(kit.get("mau"), "#dfe6f2")}
    out = {"ten": str(raw.get("ten") or "Thí nghiệm")[:90], "dung_cu": vessels, "thu_khi": kit, "buoc": steps,
           "hien_tuong": [str(x)[:140] for x in (raw.get("hien_tuong") or [])[:6]], "phuong_trinh": str(raw.get("phuong_trinh") or "")[:140],
           "giai_thich": str(raw.get("giai_thich") or "")[:300], "an_toan": [str(x)[:140] for x in (raw.get("an_toan") or [])[:4]],
           "chat_tham_gia": _sim_species(raw.get("chat_tham_gia")), "san_pham": _sim_species(raw.get("san_pham"))}
    return out, None


def sim_balance(sc):
    """True/False nếu kiểm tra được cân bằng bằng RDKit, None nếu thiếu SMILES hoặc SMILES lỗi."""
    L, R = sc.get("chat_tham_gia") or [], sc.get("san_pham") or []
    if not L or not R or not all(s.get("smiles") for s in L + R):
        return None
    bl, br = atom_balance([(s["he_so"], s["smiles"]) for s in L]), atom_balance([(s["he_so"], s["smiles"]) for s in R])
    return None if bl is None or br is None else bl == br


def sim_html(sc, custom=None):
    if custom and custom.get("on"):
        sc = {**sc, "buoc": [{**s, "thoi_gian_cung": custom["d"][i]} for i, s in enumerate(sc["buoc"]) if i < len(custom["d"])] + sc["buoc"][len(custom["d"]):], "thoi_gian_ket": custom.get("end")}
    data = {k: v for k, v in sc.items() if k in ("ten", "dung_cu", "thu_khi", "buoc", "hien_tuong", "phuong_trinh", "giai_thich", "am_ket", "am_ket_dur", "thoi_gian_ket")}
    return SIM_TEMPLATE.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))


def _sim_set(txt):
    st.session_state.sim_q = txt


def _sim_finish(sc):
    """Lưu kịch bản vào phiên và tự thu giọng đọc AI cho lời dẫn."""
    for k in [k for k in st.session_state if str(k).startswith("sim_d_")]:
        del st.session_state[k]
    st.session_state.sim_custom = None
    if st.session_state.n < MAX_PER_SESSION:   # giọng AI tự nhiên luôn bật; lỗi hoặc quá hạn thì trình duyệt đọc thay
        with st.spinner("Đang thu giọng đọc tự nhiên…"):
            sc2, ok = sim_attach_audio(sc)
        if ok:
            sc = sc2
            st.session_state.n += 1
    st.session_state.sim = sc
    add_history(sc["ten"][:60], "Mô phỏng")


def _sim_make(top, q):
    key = sim_match_builtin(q)
    if key:
        _sim_finish({k: v for k, v in SIM_BUILTINS[key].items() if k != "kw"})
        return
    if st.session_state.n >= MAX_PER_SESSION:
        show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        return
    with st.spinner("Đang lập kịch bản và dựng phim…"):
        try:
            sc, err = sim_clean(ask_ai([SIM_PROMPT.replace("__Q__", q.replace('"', "'"))], temperature=0.2))
            st.session_state.n += 1
        except Exception as e:
            show_top_error(top, "Có lỗi khi lập kịch bản, hãy thử lại.", e)
            return
    if err:
        st.session_state.sim = None
        show_top_error(top, err)
    else:
        _sim_finish(sc)


def _sim_ui(top):
    st.text_input("Thí nghiệm cần mô phỏng", placeholder="Ví dụ: cho kẽm vào dung dịch HCl và thu khí hiđro", key="sim_q")
    for row in range(0, len(SIM_EXAMPLES), 3):
        for c, e in zip(st.columns(3), SIM_EXAMPLES[row:row + 3]):
            c.button(e, key=f"sim_ex_{row}_{e[:6]}", on_click=_sim_set, args=(e,), width="stretch")
    q = (st.session_state.get("sim_q") or "").strip()
    if st.button("Tạo mô phỏng", type="primary", disabled=not q, key="sim_go"):
        _sim_make(top, q)
    sc = st.session_state.get("sim")
    if not sc:
        return
    st.divider()
    custom = st.session_state.get("sim_custom")
    page = sim_html(sc, custom)
    if hasattr(st, "iframe"):
        st.iframe(page, height=600)
    else:
        components.html(page, height=600)
    with st.expander("⏱️ Chỉnh thời gian từng bước tùy ý"):
        with st.form("sim_time"):
            cur = custom or {}
            auto = st.checkbox("Tự động theo lời đọc (khuyên dùng)", value=not cur.get("on"))
            vals = []
            for i, s in enumerate(sc["buoc"]):
                dv = float(cur["d"][i]) if cur.get("on") and i < len(cur.get("d", [])) else sim_auto_dur(s)
                vals.append(st.number_input(f"Bước {i + 1}: {s.get('tieu_de', '')[:34]} (giây)", 2.0, 90.0, value=dv, step=0.5, key=f"sim_d_{i}"))
            endv = st.number_input("Màn kết luận (giây)", 3.0, 90.0, value=float(cur.get("end") or max(8.0, (sc.get("am_ket_dur") or 0) + 1)), step=0.5, key="sim_d_end")
            if st.form_submit_button("Áp dụng thời gian"):
                st.session_state.sim_custom = {"on": not auto, "d": vals, "end": endv}
                st.rerun()
        st.caption("Bỏ chọn “Tự động” rồi nhập số giây cho từng bước. Giọng đọc sẽ tự nhanh hoặc chậm lại cho vừa thời lượng. Muốn nhanh chậm cả phim, kéo thanh tốc độ ⏱ bên dưới phim.")
    bal = sim_balance(sc)
    if sc.get("phuong_trinh"):
        st.markdown("**Phương trình phản ứng:** " + html.escape(_sub(sc["phuong_trinh"].replace("->", "→"))))
        if bal is False:
            st.warning("Phương trình trong kịch bản **chưa cân bằng** theo hệ số đã nêu, hãy kiểm tra lại.")
    if sc.get("hien_tuong"):
        st.markdown("**Hiện tượng quan sát được:**\n" + "\n".join("- " + html.escape(x) for x in sc["hien_tuong"]))
    if sc.get("giai_thich"):
        st.markdown("**Giải thích:** " + html.escape(sc["giai_thich"]))
    if sc.get("an_toan"):
        st.warning("**An toàn khi làm thí nghiệm thật:** " + " • ".join(sc["an_toan"]))



# ---------------- Điều hướng ----------------
HOME = "Trang chủ"
ASSIST = "Trợ lý AI"
SECTIONS = [
    ("Công thức cấu tạo", "Công thức cấu tạo viết tay", "Chụp ảnh công thức vẽ tay kiểu SGK → mã SMILES, tên gọi, khối lượng mol, tính chất.", structure_ui, "structure"),
    ("Phương trình phản ứng", "Phương trình phản ứng", "Đọc phương trình, cho biết loại phản ứng và kiểm tra đã cân bằng chưa.", reaction_ui, "flask"),
    ("Đồng phân", "Đồng phân", "Hỏi bằng tiếng Việt hoặc gõ công thức, máy sinh và đếm đồng phân chính xác.", isomer_ui, "isomer"),
    ("Dụng cụ hóa học", "Dụng cụ hóa học", "Chụp ảnh dụng cụ để biết tên, công dụng, cách dùng, hoặc gõ tên thí nghiệm để xem phim mô phỏng.", equipment_ui, "equipment"),
    ("Sơ đồ tư duy", "Sơ đồ tư duy", "Nhập chủ đề hoặc tên chất, máy vẽ sơ đồ để ôn tập, sửa được và tải về PowerPoint, draw.io.", mindmap_ui, "mind"),
]
LABELS = [x[0] for x in SECTIONS]
NAV_OPTS = [HOME] + LABELS
STRUCT, REACT, ISO, EQUIP, MIND = LABELS[0], LABELS[1], LABELS[2], LABELS[3], LABELS[4]

st.session_state.setdefault("nav", HOME)
st.session_state.setdefault("nav_side", HOME)


def go(to):
    st.session_state.nav = to
    st.session_state.nav_side = to if to in NAV_OPTS else None


def on_side():
    if st.session_state.nav_side:
        go(st.session_state.nav_side)


def _fold(s):
    s = unicodedata.normalize("NFD", (s or "").lower().replace("đ", "d"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip()


def quick_lookup(text):
    """Đưa chuỗi người dùng nhập tới công cụ phù hợp. Trả về False nếu chuỗi rỗng."""
    t = (text or "").strip()
    if not t:
        return False
    if smiles_like(t):  # giống SMILES: phân tích cấu tạo ngay
        st.session_state.st_mode = "⌨️ Nhập SMILES"
        st.session_state.st_smi = t
        st.session_state.pending_smi = t
        go(STRUCT)
        add_history(t, "Công thức")
    elif is_simple_formula(t) and "C" in t:  # công thức phân tử hữu cơ: đếm đồng phân
        st.session_state.iso_q = t
        go(ISO)
        add_history(t, "Đồng phân")
    else:  # còn lại: hỏi trợ lý AI
        st.session_state.pending_chat = t
        go(ASSIST)
        add_history(t, "Hỏi đáp")
    return True


def on_quick():
    quick_lookup(st.session_state.get("quick_in", ""))
    st.session_state.quick_in = ""


SEARCH_KEYS = {
    STRUCT: "cong thuc cau tao smiles viet tay ctct ten goi khoi luong mol",
    REACT: "phuong trinh phan ung can bang loai phan ung",
    ISO: "dong phan dem dong phan cis trans quang hoc",
    EQUIP: "dung cu hoa hoc thi nghiem ong nghiem binh cau",
    MIND: "so do tu duy mindmap on tap",
    ASSIST: "tro ly ai hoi dap chat",
}


def on_search():
    q = (st.session_state.get("gsearch") or "").strip()
    st.session_state.gsearch = ""
    if not q:
        return
    f = _fold(q)
    for lab, keys in SEARCH_KEYS.items():
        if f in _fold(lab) or (len(f) >= 3 and f in keys):
            go(lab)
            return
    quick_lookup(q)


def on_home_upload():
    f = st.session_state.get("home_up")
    if f is not None:
        st.session_state.pending_img = f.getvalue()
        go(STRUCT)


# ---------------- Tài liệu học tập (dữ liệu tra cứu có sẵn, không gọi AI) ----------------
_PT_NAMES = ("Hiđro,Heli,Liti,Beri,Bo,Cacbon,Nitơ,Oxi,Flo,Neon,Natri,Magie,Nhôm,Silic,Photpho,Lưu huỳnh,Clo,Agon,Kali,Canxi,"
             "Scandi,Titan,Vanađi,Crom,Mangan,Sắt,Coban,Niken,Đồng,Kẽm,Gali,Gecmani,Asen,Selen,Brom,Kripton,Rubiđi,Stronti,Ytri,Zirconi,"
             "Niobi,Molipđen,Tecneti,Rutheni,Rodi,Palađi,Bạc,Cađimi,Inđi,Thiếc,Antimon,Telu,Iot,Xenon,Xesi,Bari,Lantan,Xeri,"
             "Prazeođim,Neođim,Prometi,Samari,Europi,Gađolini,Tecbi,Đisprozi,Honmi,Eribi,Tuli,Ytecbi,Luteti,Hafni,Tantan,Vonfram,"
             "Reni,Osmi,Iriđi,Platin,Vàng,Thủy ngân,Tali,Chì,Bitmut,Poloni,Atatin,Radon,Franxi,Radi,Actini,Thori,Protactini,Urani,"
             "Neptuni,Plutoni,Amerixi,Curi,Berkeli,Califoni,Ensteni,Fecmi,Menđelevi,Nobeli,Lorenxi,Rutherfordi,Đubni,Seaborgi,Bohri,"
             "Hassi,Meitneri,Đacmstati,Rơnghen,Copernixi,Nihoni,Flerovi,Moscovi,Livermori,Tennessin,Oganesson").split(",")
assert len(_PT_NAMES) == 118

_PT_CATS = {
    "kl-kiem": ("Kim loại kiềm", {3, 11, 19, 37, 55, 87}),
    "kl-kt": ("Kim loại kiềm thổ", {4, 12, 20, 38, 56, 88}),
    "chuyen": ("Kim loại chuyển tiếp", set(range(21, 31)) | set(range(39, 49)) | set(range(72, 81)) | set(range(104, 113))),
    "sau-chuyen": ("Kim loại khác", {13, 31, 49, 50, 81, 82, 83, 84, 113, 114, 115, 116}),
    "a-kim-loai": ("Bán kim loại", {5, 14, 32, 33, 51, 52}),
    "phi-kim": ("Phi kim", {1, 6, 7, 8, 15, 16, 34}),
    "halogen": ("Halogen", {9, 17, 35, 53, 85, 117}),
    "khi-hiem": ("Khí hiếm", {2, 10, 18, 36, 54, 86, 118}),
    "lantan": ("Họ lantan", set(range(57, 72))),
    "actini": ("Họ actini", set(range(89, 104))),
}
assert set().union(*[v[1] for v in _PT_CATS.values()]) == set(range(1, 119))
_NO_STABLE = {43, 61} | set(range(84, 90)) | set(range(93, 119))  # không có đồng vị bền: ghi số khối trong ngoặc vuông


def _pt_cat(z):
    return next(k for k, (_, s) in _PT_CATS.items() if z in s)


def _pt_pos(z):
    """Trả về (hàng, cột) trên lưới bảng tuần hoàn dạng 18 cột."""
    if z == 1: return 1, 1
    if z == 2: return 1, 18
    if 3 <= z <= 4: return 2, z - 2
    if 5 <= z <= 10: return 2, z + 8
    if 11 <= z <= 12: return 3, z - 10
    if 13 <= z <= 18: return 3, z
    if 19 <= z <= 36: return 4, z - 18
    if 37 <= z <= 54: return 5, z - 36
    if z in (55, 56): return 6, z - 54
    if 57 <= z <= 71: return 9, z - 54
    if 72 <= z <= 86: return 6, z - 68
    if z in (87, 88): return 7, z - 86
    if 89 <= z <= 103: return 10, z - 86
    return 7, z - 100


def _pt_mass(z):
    m = Chem.GetPeriodicTable().GetAtomicWeight(z)
    return f"[{round(m)}]" if z in _NO_STABLE else f"{m:.2f}"


@st.dialog("Bảng tuần hoàn các nguyên tố hóa học", width="large")
def periodic_dialog():
    pt = Chem.GetPeriodicTable()
    cells = ['<div class="ptc ptx" style="grid-row:6;grid-column:3"><b>57–71</b></div>',
             '<div class="ptc ptx" style="grid-row:7;grid-column:3"><b>89–103</b></div>']
    for z in range(1, 119):
        r, c = _pt_pos(z)
        sym, nm = pt.GetElementSymbol(z), _PT_NAMES[z - 1]
        cells.append(f'<div class="ptc {_pt_cat(z)}" style="grid-row:{r};grid-column:{c}" title="{z} · {nm}">'
                     f'<i>{z}</i><b>{sym}</b><span>{nm}</span><em>{_pt_mass(z)}</em></div>')
    legend = "".join(f'<span><i class="{k}"></i>{v[0]}</span>' for k, v in _PT_CATS.items())
    st.markdown(f'<div class="ptwrap"><div class="ptgrid">{"".join(cells)}</div></div><div class="ptleg">{legend}</div>', unsafe_allow_html=True)
    st.caption("Khối lượng nguyên tử tính theo RDKit (đơn vị amu). Số trong ngoặc vuông [ ] là số khối của đồng vị bền nhất, "
               "dùng cho nguyên tố không có đồng vị bền.")
    opts = [f"{z} · {pt.GetElementSymbol(z)} · {_PT_NAMES[z - 1]}" for z in range(1, 119)]
    pick = st.selectbox("Tra cứu một nguyên tố", opts, index=5, key="pt_pick")
    z = opts.index(pick) + 1
    r, c = _pt_pos(z)
    a, b, d, e = st.columns(4)
    a.metric("Số hiệu (Z)", z)
    b.metric("Ký hiệu", pt.GetElementSymbol(z))
    d.metric("Khối lượng", _pt_mass(z))
    e.metric("Nhóm", _PT_CATS[_pt_cat(z)][0])
    if r >= 9:
        st.markdown(f"**Chu kì:** {6 if r == 9 else 7} (hàng riêng bên dưới bảng)")
    else:
        st.markdown(f"**Chu kì:** {r} · **Cột:** {c}")


def _sub(s):
    """Chuyển chữ số sau chữ cái hoặc dấu ngoặc đóng thành chỉ số dưới."""
    return re.sub(r"(?<=[A-Za-z\)])(\d+)", lambda m: m.group(1).translate(SUB), s).replace(")n", ")ₙ")


ORG_ROWS = [  # (loại hợp chất, công thức chung, nhóm chức hoặc đặc điểm, ví dụ)
    ("Ankan", "C<sub>n</sub>H<sub>2n+2</sub> (n ≥ 1)", "Chỉ có liên kết đơn", "CH4 (metan), C2H6 (etan)"),
    ("Xicloankan (1 vòng)", "C<sub>n</sub>H<sub>2n</sub> (n ≥ 3)", "Mạch vòng, liên kết đơn", "C3H6 (xiclopropan)"),
    ("Anken", "C<sub>n</sub>H<sub>2n</sub> (n ≥ 2)", "1 liên kết đôi C=C", "C2H4 (eten)"),
    ("Ankađien", "C<sub>n</sub>H<sub>2n–2</sub> (n ≥ 3)", "2 liên kết đôi C=C", "C4H6 (buta-1,3-đien)"),
    ("Ankin", "C<sub>n</sub>H<sub>2n–2</sub> (n ≥ 2)", "1 liên kết ba C≡C", "C2H2 (etin)"),
    ("Aren (đồng đẳng benzen)", "C<sub>n</sub>H<sub>2n–6</sub> (n ≥ 6)", "Có vòng benzen", "C6H6 (benzen), C7H8 (toluen)"),
    ("Dẫn xuất halogen no, đơn chức", "C<sub>n</sub>H<sub>2n+1</sub>X", "–X (F, Cl, Br, I)", "CH3Cl, C2H5Br"),
    ("Ancol no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>OH (C<sub>n</sub>H<sub>2n+2</sub>O)", "–OH gắn vào C no", "CH3OH, C2H5OH"),
    ("Phenol", "C<sub>6</sub>H<sub>5</sub>OH", "–OH gắn trực tiếp vào vòng benzen", "C6H5OH"),
    ("Anđehit no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>CHO (C<sub>n</sub>H<sub>2n</sub>O)", "–CHO", "HCHO, CH3CHO"),
    ("Xeton no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n</sub>O (n ≥ 3)", ">C=O", "CH3COCH3 (propanon)"),
    ("Axit cacboxylic no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>COOH (C<sub>n</sub>H<sub>2n</sub>O<sub>2</sub>)", "–COOH", "HCOOH, CH3COOH"),
    ("Este no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n</sub>O<sub>2</sub> (n ≥ 2)", "–COO–", "HCOOCH3, CH3COOC2H5"),
    ("Amin no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+3</sub>N (n ≥ 1)", "–NH2 (amin bậc 1)", "CH3NH2, C2H5NH2"),
    ("Amino axit", "H<sub>2</sub>N–R–COOH", "–NH2 và –COOH", "H2NCH2COOH (glyxin)"),
    ("Glucozơ, fructozơ", "C<sub>6</sub>H<sub>12</sub>O<sub>6</sub>", "Monosaccarit", "C6H12O6"),
    ("Saccarozơ, mantozơ", "C<sub>12</sub>H<sub>22</sub>O<sub>11</sub>", "Đisaccarit", "C12H22O11"),
    ("Tinh bột, xenlulozơ", "(C<sub>6</sub>H<sub>10</sub>O<sub>5</sub>)<sub>n</sub>", "Polisaccarit", "(C6H10O5)n"),
    ("Chất béo", "(RCOO)<sub>3</sub>C<sub>3</sub>H<sub>5</sub>", "Trieste của glixerol với axit béo", "(C17H35COO)3C3H5 (tristearin)"),
]


@st.dialog("Công thức hợp chất hữu cơ", width="large")
def organic_dialog():
    rows = "".join(f"<tr><td><b>{a}</b></td><td>{b}</td><td>{_sub(c)}</td><td>{_sub(d)}</td></tr>" for a, b, c, d in ORG_ROWS)
    st.markdown('<div class="reftbl"><table><thead><tr><th>Loại hợp chất</th><th>Công thức chung</th><th>Nhóm chức / đặc điểm</th><th>Ví dụ</th></tr></thead>'
                f'<tbody>{rows}</tbody></table></div>', unsafe_allow_html=True)
    st.markdown("**Độ bất bão hòa:** k = (2C + 2 + N − H − X) / 2, trong đó X là tổng số nguyên tử halogen. "
                "k bằng tổng số vòng và số liên kết π (một liên kết đôi tính 1, một liên kết ba tính 2).")
    st.caption("Các chất đồng đẳng hơn kém nhau một hoặc nhiều nhóm CH2.")
    st.button("Mở công cụ Công thức cấu tạo", key="org_open", on_click=go, args=(STRUCT,))


# Mỗi phương trình: (loại, vế trái, vế phải, điều kiện, mũi tên). Vế là danh sách (hệ số, công thức).
RX_ORG = [
    ("Thế", [(1, "CH4"), (1, "Cl2")], [(1, "CH3Cl"), (1, "HCl")], "ánh sáng", "→"),
    ("Cộng H2", [(1, "CH2=CH2"), (1, "H2")], [(1, "CH3CH3")], "Ni, t°", "→"),
    ("Cộng Br2", [(1, "CH2=CH2"), (1, "Br2")], [(1, "CH2BrCH2Br")], "", "→"),
    ("Cộng HX", [(1, "CH2=CH2"), (1, "HCl")], [(1, "CH3CH2Cl")], "", "→"),
    ("Cộng H2O", [(1, "CH2=CH2"), (1, "H2O")], [(1, "CH3CH2OH")], "H⁺, t°", "→"),
    ("Cộng Br2 (ankin)", [(1, "CH≡CH"), (2, "Br2")], [(1, "CHBr2CHBr2")], "", "→"),
    ("Tách nước", [(1, "CH3CH2OH")], [(1, "CH2=CH2"), (1, "H2O")], "H2SO4 đặc, 170°C", "→"),
    ("Oxi hóa ancol", [(2, "CH3CH2OH"), (1, "O2")], [(2, "CH3CHO"), (2, "H2O")], "Cu, t°", "→"),
    ("Đốt cháy", [(1, "C2H4"), (3, "O2")], [(2, "CO2"), (2, "H2O")], "t°", "→"),
    ("Đốt cháy", [(2, "C2H6"), (7, "O2")], [(4, "CO2"), (6, "H2O")], "t°", "→"),
    ("Thế (benzen)", [(1, "C6H6"), (1, "Br2")], [(1, "C6H5Br"), (1, "HBr")], "Fe, t°", "→"),
    ("Nitro hóa", [(1, "C6H6"), (1, "HNO3")], [(1, "C6H5NO2"), (1, "H2O")], "H2SO4 đặc", "→"),
    ("Thế (phenol)", [(1, "C6H5OH"), (3, "Br2")], [(1, "C6H2Br3OH↓"), (3, "HBr")], "", "→"),
    ("Este hóa", [(1, "CH3COOH"), (1, "C2H5OH")], [(1, "CH3COOC2H5"), (1, "H2O")], "H2SO4 đặc, t°", "⇌"),
    ("Xà phòng hóa", [(1, "CH3COOC2H5"), (1, "NaOH")], [(1, "CH3COONa"), (1, "C2H5OH")], "t°", "→"),
    ("Tráng bạc", [(1, "CH3CHO"), (2, "AgNO3"), (3, "NH3"), (1, "H2O")], [(1, "CH3COONH4"), (2, "Ag↓"), (2, "NH4NO3")], "t°", "→"),
    ("Lên men", [(1, "C6H12O6")], [(2, "C2H5OH"), (2, "CO2↑")], "enzim, 30–35°C", "→"),
    ("Axit – bazơ", [(1, "CH3COOH"), (1, "NaOH")], [(1, "CH3COONa"), (1, "H2O")], "", "→"),
    ("Amin + axit", [(1, "C6H5NH2"), (1, "HCl")], [(1, "C6H5NH3Cl")], "", "→"),
]
RX_INORG = [
    ("Kim loại + axit", [(1, "Zn"), (2, "HCl")], [(1, "ZnCl2"), (1, "H2↑")], "", "→"),
    ("Kim loại + axit", [(2, "Al"), (6, "HCl")], [(2, "AlCl3"), (3, "H2↑")], "", "→"),
    ("Kim loại + muối", [(1, "Fe"), (1, "CuSO4")], [(1, "FeSO4"), (1, "Cu")], "", "→"),
    ("Kim loại + nước", [(2, "Na"), (2, "H2O")], [(2, "NaOH"), (1, "H2↑")], "", "→"),
    ("Trung hòa", [(1, "NaOH"), (1, "HCl")], [(1, "NaCl"), (1, "H2O")], "", "→"),
    ("Trao đổi (kết tủa)", [(1, "AgNO3"), (1, "NaCl")], [(1, "AgCl↓"), (1, "NaNO3")], "", "→"),
    ("Trao đổi (kết tủa)", [(1, "BaCl2"), (1, "H2SO4")], [(1, "BaSO4↓"), (2, "HCl")], "", "→"),
    ("Muối cacbonat + axit", [(1, "CaCO3"), (2, "HCl")], [(1, "CaCl2"), (1, "CO2↑"), (1, "H2O")], "", "→"),
    ("Oxit axit + bazơ", [(1, "Ca(OH)2"), (1, "CO2")], [(1, "CaCO3↓"), (1, "H2O")], "", "→"),
    ("Phân hủy", [(1, "CaCO3")], [(1, "CaO"), (1, "CO2↑")], "t°", "→"),
    ("Phân hủy", [(2, "KClO3")], [(2, "KCl"), (3, "O2↑")], "MnO2, t°", "→"),
    ("Hóa hợp", [(2, "H2"), (1, "O2")], [(2, "H2O")], "t°", "→"),
    ("Hóa hợp", [(2, "Mg"), (1, "O2")], [(2, "MgO")], "t°", "→"),
    ("Oxit kim loại + H2", [(1, "CuO"), (1, "H2")], [(1, "Cu"), (1, "H2O")], "t°", "→"),
    ("Luyện gang", [(1, "Fe2O3"), (3, "CO")], [(2, "Fe"), (3, "CO2")], "t°", "→"),
    ("Axit đặc", [(1, "Cu"), (2, "H2SO4")], [(1, "CuSO4"), (1, "SO2↑"), (2, "H2O")], "H2SO4 đặc, t°", "→"),
]


def _atoms(formula):
    """Đếm nguyên tử của một công thức (hỗ trợ ngoặc), bỏ ký hiệu ↑ ↓ = ≡."""
    f = re.sub(r"[↑↓=≡]", "", formula)
    stack = [{}]
    for tok in re.findall(r"([A-Z][a-z]?)(\d*)|(\()|\)(\d*)", f):
        el, n, op, cn = tok
        if el:
            stack[-1][el] = stack[-1].get(el, 0) + (int(n) if n else 1)
        elif op:
            stack.append({})
        else:
            grp = stack.pop()
            for k, v in grp.items():
                stack[-1][k] = stack[-1].get(k, 0) + v * (int(cn) if cn else 1)
    return stack[0]


def rx_balanced(left, right):
    def tot(side):
        c = Counter()
        for k, f in side:
            for el, v in _atoms(f).items():
                c[el] += k * v
        return c
    return tot(left) == tot(right)


def _rx_html(rows):
    out = []
    for typ, L, R, cond, arrow in rows:
        side = lambda s: " + ".join((f"{k} " if k > 1 else "") + html.escape(_sub(f)) for k, f in s)
        arr = f'<span class="rxa"><small>{html.escape(_sub(cond))}</small><i>{arrow}</i></span>'
        out.append(f'<div class="rxr"><span class="rxt">{html.escape(_sub(typ))}</span><div class="rxe">{side(L)}{arr}{side(R)}</div></div>')
    return "".join(out)


@st.dialog("Phương trình phản ứng thường gặp", width="large")
def reactions_dialog():
    t1, t2 = st.tabs(["Hóa hữu cơ", "Hóa vô cơ"])
    with t1:
        st.markdown(_rx_html(RX_ORG), unsafe_allow_html=True)
        st.caption("Trùng hợp: nCH₂=CH₂ → (–CH₂–CH₂–)ₙ (điều kiện: t°, p, xúc tác).")
    with t2:
        st.markdown(_rx_html(RX_INORG), unsafe_allow_html=True)
    st.caption("Các phương trình đã được máy kiểm tra cân bằng nguyên tử. ↑ là chất khí, ↓ là chất kết tủa. Điều kiện ghi trên mũi tên.")
    st.button("Mở công cụ Phương trình phản ứng", key="rx_open", on_click=go, args=(REACT,))


@st.dialog("Mẹo & lưu ý", width="large")
def tips_dialog():
    st.markdown("**Chụp ảnh để máy đọc chính xác**\n"
                "- Viết công thức rõ nét, nền sáng, mỗi ảnh chỉ một chất.\n"
                "- Chụp thẳng góc, đủ sáng, tránh bóng đổ và nếp gấp giấy.\n"
                "- Công thức sai hóa trị sẽ được máy chỉ rõ nguyên tử sai, hãy sửa lại hình vẽ rồi chụp lại.\n"
                "- Có thể sửa mã SMILES rồi phân tích lại ngay dưới kết quả.\n\n"
                "**Gọi tên và quy tắc hay dùng (theo SGK)**\n"
                "- Danh pháp thay thế: chọn mạch chính dài nhất có nhóm chức, đánh số từ phía gần nhóm chức, "
                "gọi tên nhánh theo thứ tự chữ cái.\n"
                "- Quy tắc Mác-côp-nhi-côp: khi cộng HX hoặc H₂O vào anken bất đối xứng, H cộng vào C có nhiều H hơn, "
                "X hoặc OH cộng vào C có ít H hơn.\n"
                "- Quy tắc Zai-xép: khi tách HX hoặc H₂O, ưu tiên tách H ở C bên cạnh có bậc cao hơn.\n"
                "- Thế vào vòng benzen: –OH, –NH₂, –CH₃ định hướng vị trí o- và p-; –NO₂, –COOH, –CHO định hướng vị trí m-.\n"
                "- Khi cân bằng, kiểm tra số nguyên tử từng nguyên tố hai vế bằng nhau; C luôn có hóa trị 4.\n\n"
                "")


# ---------------- Bảng tính tan (axit, bazơ, muối) ----------------
# Cation: (nhãn hiển thị, ký hiệu để ghép công thức, điện tích)
SOL_CATIONS = [("H⁺", "H", 1), ("Na⁺", "Na", 1), ("K⁺", "K", 1), ("NH₄⁺", "NH4", 1), ("Ag⁺", "Ag", 1), ("Mg²⁺", "Mg", 2),
               ("Ca²⁺", "Ca", 2), ("Ba²⁺", "Ba", 2), ("Zn²⁺", "Zn", 2), ("Hg²⁺", "Hg", 2), ("Pb²⁺", "Pb", 2), ("Cu²⁺", "Cu", 2),
               ("Fe²⁺", "Fe", 2), ("Fe³⁺", "Fe", 3), ("Al³⁺", "Al", 3)]
# Anion (cột): (nhãn, gốc, điện tích, gốc nhiều nguyên tử cần ngoặc khi có chỉ số)
SOL_ANIONS = [("OH⁻", "OH", 1, True), ("Cl⁻", "Cl", 1, False), ("NO₃⁻", "NO3", 1, True), ("CH₃COO⁻", "CH3COO", 1, True),
              ("S²⁻", "S", 2, False), ("SO₃²⁻", "SO3", 2, True), ("SO₄²⁻", "SO4", 2, True), ("CO₃²⁻", "CO3", 2, True),
              ("SiO₃²⁻", "SiO3", 2, True), ("PO₄³⁻", "PO4", 3, True)]
# Mã: t = tan, i = ít tan, k = không tan (kết tủa), b = tan nhưng bay hơi hoặc dễ phân hủy, x = không tồn tại hoặc bị nước phân hủy, ? = chưa có số liệu chắc chắn
SOL_DATA = {
    #        OH   Cl   NO3  Ac   S    SO3  SO4  CO3  SiO3 PO4
    "H":    "x    b    b    b    b    b    t    b    k    t",
    "Na":   "t    t    t    t    t    t    t    t    t    t",
    "K":    "t    t    t    t    t    t    t    t    t    t",
    "NH4":  "b    t    t    t    t    t    t    t    ?    t",
    "Ag":   "x    k    t    i    k    k    i    k    ?    k",
    "Mg":   "k    t    t    t    x    i    t    k    k    k",
    "Ca":   "i    t    t    t    x    k    i    k    k    k",
    "Ba":   "t    t    t    t    ?    k    k    k    k    k",
    "Zn":   "k    t    t    t    k    k    t    k    k    k",
    "Hg":   "x    t    t    t    k    ?    x    x    ?    k",
    "Pb":   "k    i    t    t    k    k    k    k    k    k",
    "Cu":   "k    t    t    t    k    x    t    x    k    k",
    "Fe2":  "k    t    t    t    k    k    t    k    k    k",
    "Fe3":  "k    t    t    t    x    x    t    x    k    k",
    "Al":   "k    t    t    t    x    x    t    x    k    k",
}
SOL_KEY = {("Fe", 2): "Fe2", ("Fe", 3): "Fe3"}
SOL_TEXT = {"t": ("Tan", "sol-t"), "i": ("Ít tan", "sol-i"), "k": ("Không tan (kết tủa)", "sol-k"),
            "b": ("Tan, bay hơi hoặc dễ phân hủy", "sol-b"), "x": ("Không tồn tại hoặc bị nước phân hủy", "sol-x"),
            "?": ("Chưa có số liệu chắc chắn", "sol-q")}


def _sol_code(ci, ai):
    sym, ch = SOL_CATIONS[ci][1], SOL_CATIONS[ci][2]
    return SOL_DATA[SOL_KEY.get((sym, ch), sym)].split()[ai]


def _sol_formula(ci, ai):
    from math import gcd
    _, csym, cch = SOL_CATIONS[ci]
    _, asym, ach, apoly = SOL_ANIONS[ai]
    g = gcd(cch, ach)
    nc, na = ach // g, cch // g
    if csym == "H" and asym == "CH3COO":
        return "CH₃COOH"
    if csym == "H" and asym == "OH":
        return "H₂O"
    c = csym if nc == 1 else (f"({csym})" if csym == "NH4" else csym) + str(nc)
    a = asym if na == 1 else (f"({asym})" if apoly else asym) + str(na)
    return _sub(c + a)


def _sol_table(sel=None):
    head = "".join(f"<th>{a[0]}</th>" for a in SOL_ANIONS)
    rows = []
    for ci, c in enumerate(SOL_CATIONS):
        cells = []
        for ai in range(len(SOL_ANIONS)):
            code = _sol_code(ci, ai)
            cls = SOL_TEXT[code][1] + (" sol-sel" if sel == (ci, ai) else "")
            cells.append(f'<td class="{cls}" title="{html.escape(_sol_formula(ci, ai))}: {SOL_TEXT[code][0]}">{ "–" if code == "x" else code}</td>')
        rows.append(f"<tr><th>{c[0]}</th>{''.join(cells)}</tr>")
    return f'<div class="soltbl"><table><thead><tr><th></th>{head}</tr></thead><tbody>{"".join(rows)}</tbody></table></div>'


@st.dialog("Bảng tính tan trong nước", width="large")
def solubility_dialog():
    st.markdown("Chọn một cation và một anion để tra tính tan của hợp chất tạo thành. Ô đang chọn được viền sáng trong bảng.")
    c1, c2 = st.columns(2)
    ci = c1.selectbox("Cation (kim loại, H, NH₄)", range(len(SOL_CATIONS)), format_func=lambda i: SOL_CATIONS[i][0], index=1, key="sol_cat")
    ai = c2.selectbox("Anion (gốc axit, OH)", range(len(SOL_ANIONS)), format_func=lambda i: SOL_ANIONS[i][0], index=1, key="sol_an")
    code = _sol_code(ci, ai)
    txt, cls = SOL_TEXT[code]
    st.markdown(f'<div class="solres"><b>{_sol_formula(ci, ai)}</b><span class="solbadge {cls}">{txt}</span></div>', unsafe_allow_html=True)
    st.markdown(_sol_table((ci, ai)), unsafe_allow_html=True)
    st.markdown('<div class="sollg"><span class="sol-t">t</span> tan <span class="sol-i">i</span> ít tan <span class="sol-k">k</span> không tan (kết tủa) '
                '<span class="sol-b">b</span> tan, bay hơi hoặc dễ phân hủy <span class="sol-x">–</span> không tồn tại hoặc bị nước phân hủy '
                '<span class="sol-q">?</span> chưa có số liệu chắc chắn</div>', unsafe_allow_html=True)


# ---------------- Các khối giao diện ----------------
def topbar():
    with st.container(key="topbar"):
        a, b = st.columns([5, 2], vertical_alignment="center")
        with a:
            st.text_input("Tìm kiếm", key="gsearch", placeholder="Tìm công cụ, công thức, tên chất…", on_change=on_search, label_visibility="collapsed")
        with b:
            st.markdown(f'<div class="tb-right">{icon_svg("sun", "#5b6293", 22)}{icon_svg("bell", "#5b6293", 22)}<div class="av">T</div></div>',
                        unsafe_allow_html=True)


def hero():
    st.markdown(
        '<div class="hero3"><div class="hl"><span class="pill">Chào bạn ! 👋</span>'
        f'<h1>Khám phá thế giới Hóa học<br>cùng <em>{html.escape(APP_NAME)}</em></h1>'
        '<div class="sub"><span>Nhanh chóng</span><i>•</i><span>Chính xác</span><i>•</i><span>Thông minh</span></div>'
        '<div class="chips"><span>⚡ Nhanh chóng</span><span>🎯 Chính xác</span><span>💡 Tiện lợi</span><span>💜 Học mọi lúc</span></div></div>'
        f'<div class="hr"><img class="art" src="{_uri(HERO_ART)}"><div class="slogan">Mỗi phản ứng<br>là một khám phá!<u></u></div></div></div>',
        unsafe_allow_html=True,
    )


def page_banner(icon, title, desc):
    st.markdown(f'<div class="pbanner"><div class="pb-ic">{icon_svg(icon, "#fff", 28)}</div><div><b>{html.escape(title)}</b>'
                f'<span>{html.escape(desc)}</span></div></div>', unsafe_allow_html=True)


TILES = [("Công thức\ncấu tạo", "structure", "t-pink", STRUCT), ("Phương trình\nphản ứng", "flask", "t-blue", REACT),
         ("Đồng phân", "isomer", "t-green", ISO), ("Dụng cụ\nhóa học", "equipment", "t-orange", EQUIP), ("Sơ đồ tư duy", "mind", "t-violet", MIND)]
FEATS = [("⚡", "Nhanh chóng"), ("🎯", "Chính xác"), ("💡", "Tiện lợi"), ("💜", "Học mọi lúc")]
HKIND = {"Công thức": ("structure", "c1"), "Phản ứng": ("flask", "c3"), "Đồng phân": ("isomer", "c0"), "Dụng cụ": ("equipment", "c2"),
         "Sơ đồ": ("mind", "c3"), "Mô phỏng": ("equipment", "c2"), "Luyện tập": ("practice", "c0"), "Hỏi đáp": ("chat", "c1")}


def home_ui():
    left, right = st.columns([2.15, 1], gap="large")
    with left:
        with st.container(key="upcard"):
            st.markdown(
                f'<div class="uphead"><div class="hex">{icon_svg("pen", "#fff", 34, 2)}</div><div><h3>Công thức cấu tạo viết tay</h3>'
                '<p>Chụp ảnh công thức vẽ tay kiểu SGK → mã SMILES, tên gọi, khối lượng mol, tính chất.</p></div>'
                f'<div class="hw">Chụp nhanh<br>Nhận kết quả ngay!</div></div>', unsafe_allow_html=True)
            cu, cd = st.columns([1, 1.5], gap="medium", vertical_alignment="center")
            with cu:
                st.file_uploader("Tải ảnh công thức", type=["jpg", "jpeg", "png", "webp"], key="home_up",
                                 label_visibility="collapsed", on_change=on_home_upload)
            with cd:
                st.markdown(
                    f'<div class="demo"><img class="paper" src="{_uri(PAPER_SVG)}"><span class="arrow">{icon_svg("arrow", "#6d4df0", 28, 2)}</span>'
                    '<div class="res"><span class="badge">✦ Kết quả ngay lập tức</span><h4>C₇H₆O₂<i>●</i></h4><b>Axit benzoic</b>'
                    '<div>SMILES: O=C(O)c1ccccc1</div><div>Khối lượng mol: 122.12 g/mol</div><div>Tính chất: Axit yếu, thể rắn</div><div>Tính tan: Ít tan trong nước lạnh, tan nhiều trong nước nóng và etanol</div>'
                    '<small>Ví dụ minh họa</small></div></div>', unsafe_allow_html=True)

        st.markdown(f'<div class="sec-title">{icon_svg("home", "currentColor", 22)}Khám phá nhanh</div>', unsafe_allow_html=True)
        with st.container(key="tilesrow"):
            for c, (txt, ic, cls, target), i in zip(st.columns(5, gap="small"), TILES, range(5)):
                with c, st.container(key=f"qk_{i}"):
                    st.markdown(f'<div class="tile {cls}"><div class="tt"><div class="ib">{icon_svg(ic, "currentColor", 18, 2)}</div><b>{txt}</b></div>'
                                f'<span class="ar">{icon_svg("arrow", "currentColor", 18, 2.2)}</span></div>', unsafe_allow_html=True)
                    st.button("Mở", key=f"qkb_{i}", on_click=go, args=(target,))

        st.markdown('<div class="feat">' + '<i></i>'.join(f"<span>{e} {t}</span>" for e, t in FEATS) + '</div>', unsafe_allow_html=True)

    with right:
        with st.container(key="quick"):
            st.markdown(f'<div class="qhead"><div class="qi">{icon_svg("pen", "#fff", 24, 2)}</div><div><b>Nhập nhanh</b>'
                        '<span>Nhập công thức để tra cứu ngay</span></div></div>', unsafe_allow_html=True)
            with st.form("quick_form", border=False, clear_on_submit=False):
                qa, qb = st.columns([5, 1.2], vertical_alignment="center", gap="small")
                with qa:
                    st.text_input("Công thức", key="quick_in", placeholder="VD: C4H10O, CCO, axit axetic…", label_visibility="collapsed")
                with qb:
                    st.form_submit_button("➜", on_click=on_quick)

        with st.container(key="res"):
            st.markdown(f'<div class="rsh">{icon_svg("equipment", "currentColor", 24)}'
                        '<div><b>Tài liệu học tập</b><span>Kho tài liệu &amp; công thức quan trọng</span></div></div>', unsafe_allow_html=True)
            if st.button("Bảng tuần hoàn", icon=":material/grid_view:", key="r_pt"):
                periodic_dialog()
            if st.button("Bảng tính tan", icon=":material/water_drop:", key="r_sol"):
                solubility_dialog()
            if st.button("Công thức hữu cơ", icon=":material/hub:", key="r_org"):
                organic_dialog()
            if st.button("Phương trình phản ứng", icon=":material/science:", key="r_rx"):
                reactions_dialog()
            if st.button("Mẹo & lưu ý", icon=":material/lightbulb:", key="r_tip"):
                tips_dialog()

        hist = st.session_state.hist
        rows = "".join(
            f'<div class="hrow"><div class="hc {HKIND.get(h.get("k"), HKIND["Công thức"])[1]}">{icon_svg(HKIND.get(h.get("k"), HKIND["Công thức"])[0], "currentColor", 20, 2)}</div>'
            f'<div><b>{html.escape(h["q"][:40] + ("…" if len(h["q"]) > 40 else ""))}</b><small>{html.escape(h.get("k", ""))} · Hôm nay · {h["t"]}</small></div></div>'
            for h in hist[:6]) or '<div class="hempty">Chưa có lượt tra cứu nào. Hãy thử Nhập nhanh hoặc dùng một công cụ bên trái.</div>'
        st.markdown(f'<div class="rcard"><div class="rh">{icon_svg("search", "currentColor", 22)}<b>Lịch sử tra cứu</b></div>{rows}</div>', unsafe_allow_html=True)
        if hist:
            with st.container(key="hclear"):
                st.button("Xóa lịch sử", key="hist_clear", on_click=lambda: st.session_state.hist.clear())
        st.markdown(f'<div class="quote"><div class="lf">🌿</div><p>“Hóa học không khó, chỉ cần bạn có công cụ đúng!”<em>— {html.escape(APP_NAME)}</em></p></div>',
                    unsafe_allow_html=True)


with st.sidebar:
    st.markdown(f'<div class="sb-brand"><img src="{_uri(LOGO_SVG)}"><div><b>{html.escape(APP_NAME[:3])}<i>{html.escape(APP_NAME[3:])}</i></b>'
                '<span>AI trong tay – Hóa học dễ dàng</span></div></div>', unsafe_allow_html=True)
    with st.container(key="sidenav"):
        st.radio("Điều hướng", NAV_OPTS, index=None, key="nav_side", on_change=on_side, label_visibility="collapsed")
    with st.container(key="aibtn"):
        st.markdown(f'<div class="aicard"><img src="{_uri(AI_FACE)}"><div><b>Trợ lý AI</b><span>Hỏi bất cứ điều gì về<br>hóa học, mình luôn ở đây!</span></div>'
                    '<div class="go">›</div></div>', unsafe_allow_html=True)
        st.button("Mở trợ lý AI", key="aibtn_b", on_click=go, args=(ASSIST,))
    st.markdown(f'<div class="mascot"><img src="{_uri(MASCOT_SVG)}"><div class="hw">Học hóa<br>thật dễ!</div></div>', unsafe_allow_html=True)

# Trên điện thoại: tự đóng thanh bên sau khi chọn một mục (mã chạy trong trang chính để không bị hủy khi làm mới)
components.html("""<script>
(function(){const w=window.parent,d=w.document;if(w.__molsnapNav)return;w.__molsnapNav=true;
const s=d.createElement('script');
s.textContent="document.addEventListener('click',function(e){if(window.innerWidth>640)return;"
+"var o=e.target.closest('[data-testid=\\"stRadioOption\\"],.st-key-aibtn');"
+"if(o&&o.closest('[data-testid=\\"stSidebar\\"]')){setTimeout(function(){"
+"var b=document.querySelector('[data-testid=\\"stSidebarCollapseButton\\"] button');if(b)b.click();},350);}},true);";
d.head.appendChild(s);})();
</script>""", height=0)

topbar()

nav = st.session_state.nav
if nav == HOME or (nav not in LABELS and nav != ASSIST):
    hero()
    home_ui()
elif nav == ASSIST:
    page_banner("chat", "Trợ lý AI", "Hỏi bất cứ điều gì về hóa học, mình luôn ở đây!")
    with st.container(key="panel"):
        assistant_ui()
else:
    label, title, desc, fn, ic = SECTIONS[LABELS.index(nav)]
    page_banner(ic, title, desc)
    with st.container(key="panel"):
        fn()
